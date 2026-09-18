"""`write_document` — real .docx written from Markdown in the backend process.

What these are actually defending, in order of how much damage they prevent:

1. **The file opens, in Word.** A .docx is a ZIP of XML whose element ORDER is fixed by
   ECMA-376; get it wrong and Word shows "unreadable content" and offers to repair — while
   `zipfile` and python-docx both read the file back happily. So every part is round-tripped
   AND `test_every_property_element_keeps_its_schema_child_order` walks the raw XML.
2. **Path parity with `write_file`.** The shared layer lives in `tools/office.py` and is
   cross-checked there; `test_path_parity_*` repeats the matrix with a .docx payload so a
   change on the .docx side cannot quietly reach a folder `write_file` refuses.
3. **The office look is really applied.** 宋体 body, 黑体 headings, a 2-character first line
   on body text and NOT on headings, lists or table cells, 1.5 spacing, A4, page numbers.
   Each of those is one attribute, and each is invisible in a read-back that only counts
   paragraphs.
4. **The merge markers mean what the schema says.** `<` and `^` are this tool's one
   extension to GFM; a wrong rectangle silently reshapes a report's header.
"""

from __future__ import annotations

import errno
import json
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

import aisuite as ai
from coworker.roots import RootDir, normalize_roots
from coworker.tools import document as document_module
from coworker.tools import office as office_module
from coworker.tools.document import (
    _PBDR_SEQ,
    _PPR_SEQ,
    _SCHEMA,
    _TBLPR_SEQ,
    _TCPR_SEQ,
    document_tools,
)
from coworker.tools.office import office_tools

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
# By code point: both are invisible in an editor, and a test that hinges on one
# must not hide it in a string literal.
_BOM = chr(0xFEFF)
_PUA = chr(0xE000)


def _tool(workspace, roots=None):
    tools = document_tools(str(workspace), roots=roots)
    assert [getattr(t, "__name__", "") for t in tools] == ["write_document"]
    return tools[0]


def _open(path: Path):
    from docx import Document

    assert zipfile.ZipFile(path).testzip() is None  # a readable ZIP container
    return Document(str(path))


def _parts(path: Path) -> dict[str, str]:
    """Every XML part of the package, decoded. Some of what matters here is invisible to a
    python-docx read-back — a theme-backed `w:color` that overrides the explicit one, an
    `xml:space` that went missing — so the bytes themselves get inspected."""
    with zipfile.ZipFile(path) as archive:
        return {
            name: archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.endswith((".xml", ".rels"))
        }


def _style_xml(styles: str, style_id: str) -> str:
    match = re.search(
        rf'<w:style [^>]*w:styleId="{style_id}">.*?</w:style>', styles, re.S
    )
    assert match, f"style {style_id} missing from styles.xml"
    return match.group(0)


# -- the typical report ----------------------------------------------------------

_REPORT = """# 季度运营报告

首段正文，**加粗**、*斜体*、`代码`、~~删除~~ 与[链接](https://example.invalid/r)。

## 一、总体情况

正文第二段，说明数据来源与口径。

### 1. 明细表

| 项目 | 金额  | <      | 备注 |
|:-----|------:|:------:|------|
| ^    | 上半年 | 下半年 | ^    |
| 甲类 | 1,200 | 1,350  | 稳定 |

### 2. 要点

1. 第一条要点
2. 第二条要点
   - 子项甲
   - 子项乙
3. 第三条要点

- 无序甲
- 无序乙

> 引用一句话。

```text
def f():
    return 1
```

---

<!-- pagebreak -->

## 二、下一步

1. 新列表从一开始
2. 第二项
"""


def test_a_typical_report_round_trips_and_carries_the_office_look(tmp_path):
    out = _tool(tmp_path)(path="季度报告.docx", markdown=_REPORT)
    target = tmp_path / "季度报告.docx"
    assert target.is_file()
    assert out.startswith(f"Wrote {target}")

    document = _open(target)
    styles_used = [p.style.name for p in document.paragraphs]
    # the leading H1 became the document Title; the rest are real Heading styles, so Word's
    # navigation pane and a generated table of contents both see the outline
    assert styles_used[0] == "Title"
    assert styles_used.count("Title") == 1
    assert "Heading 2" in styles_used and "Heading 3" in styles_used
    assert [s for s in styles_used if s.startswith("List Number")] == [
        "List Number",
        "List Number",
        "List Number",
        "List Number",
        "List Number",
    ]
    assert [s for s in styles_used if s.startswith("List Bullet")] == [
        "List Bullet 2",
        "List Bullet 2",
        "List Bullet",
        "List Bullet",
    ]
    assert "Quote" in styles_used
    assert len(document.tables) == 1

    text = "\n".join(p.text for p in document.paragraphs)
    assert "季度运营报告" in text
    assert "首段正文，加粗、斜体、代码、删除 与链接。" in text  # marks are formatting, not text
    assert "    return 1" in text  # the code block's indentation survived

    parts = _parts(target)
    body, styles = parts["word/document.xml"], parts["word/styles.xml"]

    # -- fonts: Chinese and Latin are separate faces, set once on the styles
    assert 'w:eastAsia="宋体"' in styles
    assert styles.count('w:eastAsia="黑体"') == 5  # Title + Heading 1..4
    assert 'w:ascii="Times New Roman"' in _style_xml(styles, "Normal")
    assert 'w:ascii="Arial"' in _style_xml(styles, "Heading2")
    # the built-in Heading/Title styles ship theme-backed fonts and colours that WIN over
    # an explicit value, so setting ours has to have removed them. Matched case-INSENSITIVELY:
    # the attributes are `w:asciiTheme` but also `w:themeShade`, and a capital-T check
    # silently missed the second spelling while `w:themeShade="BF"` sat on Title and
    # Heading 1 for a whole review cycle.
    for style_id in ("Title", "Heading1", "Heading2", "Heading3", "Heading4"):
        block = _style_xml(styles, style_id)
        assert "theme" not in block.lower(), style_id
        assert 'w:val="000000"' in block, style_id

    # -- the 2-character first line is on body text ONLY
    assert styles.count('w:firstLineChars="200"') == 1
    assert 'w:firstLineChars="200"' in _style_xml(styles, "Normal")
    assert 'w:firstLine="480"' in _style_xml(styles, "Normal")  # for readers that ignore Chars
    assert 'w:firstLineChars="200"' not in body  # nothing sets it per paragraph
    for style_id in (
        "Title",
        "Heading1",
        "Heading2",
        "Heading3",
        "Heading4",
        "ListNumber",
        "ListNumber2",
        "ListNumber3",
        "ListBullet",
        "ListBullet2",
        "ListBullet3",
        "Quote",
        "Footer",
    ):
        assert 'w:firstLineChars="0"' in _style_xml(styles, style_id), style_id

    # -- 1.5 line spacing, A4, and the title
    assert 'w:line="360" w:lineRule="auto"' in _style_xml(styles, "Normal")
    assert '<w:pgSz w:w="11906" w:h="16838"/>' in body  # A4 portrait, in twips
    assert '<w:pgMar w:top="1440"' in body and 'w:left="1797"' in body
    title = _style_xml(styles, "Title")
    assert '<w:jc w:val="center"/>' in title
    assert '<w:sz w:val="44"/>' in title  # 22pt = 二号
    assert "<w:b/>" in title
    assert "w:pBdr" not in title  # the built-in blue rule is not an office look
    assert '<w:sz w:val="36"/>' in _style_xml(styles, "Heading1")  # 18pt

    # -- centred page number, as a real PAGE field
    footer = next(v for k, v in parts.items() if re.match(r"word/footer\d*\.xml", k))
    assert "PAGE" in footer and 'w:fldCharType="begin"' in footer
    assert '<w:jc w:val="center"/>' in footer

    # -- the merged header: 金额 spans two columns, 项目 and 备注 span two rows
    table = document.tables[0]
    assert (len(table.rows), len(table.columns)) == (3, 4)
    assert table.cell(0, 1)._tc is table.cell(0, 2)._tc  # horizontal
    assert table.cell(0, 0)._tc is table.cell(1, 0)._tc  # vertical
    assert table.cell(0, 3)._tc is table.cell(1, 3)._tc
    assert table.cell(0, 0)._tc is not table.cell(0, 1)._tc
    assert body.count("<w:gridSpan") == 1  # one horizontally merged cell
    assert body.count("<w:vMerge") == 4  # two vertical merges: a restart and a continue each
    assert [c.text for c in table.rows[0].cells] == ["项目", "金额", "金额", "备注"]
    assert [c.text for c in table.rows[2].cells] == ["甲类", "1,200", "1,350", "稳定"]
    assert '<w:tblW w:type="pct" w:w="5000"/>' in body  # fills the text column
    assert body.count('w:fill="DDEBF7"') == 3  # the three anchored header cells
    assert table.cell(0, 0).paragraphs[0].runs[0].bold is True
    # the header repeats on every page the table spills onto — and only the header does
    rows_xml = re.findall(r"<w:tr\b.*?</w:tr>", body, re.S)
    assert [("<w:tblHeader/>" in r) for r in rows_xml] == [True, False, False]

    # -- code: shaded, and its leading spaces are preserved rather than collapsed
    assert body.count('w:fill="F2F2F2"') == 2  # one paragraph per code line
    assert '<w:t xml:space="preserve">    return 1</w:t>' in body

    # -- the rule and the page break
    assert "<w:pBdr>" in body and '<w:bottom w:val="single"' in body
    assert body.count('<w:br w:type="page"/>') == 1

    assert out.endswith(
        "Document: 12 paragraph(s), 5 heading(s), 1 table(s), 3 list(s), "
        "1 page break(s); style office."
    )


def test_no_theme_attribute_survives_anywhere_we_set_an_explicit_value(tmp_path):
    """The twin attribute names only LOOK regular, and deriving them got it wrong.

    `w:color`'s modifiers are `w:themeTint`/`w:themeShade` — NOT `w:themeColorTint` — so a
    helper that appended "Tint"/"Shade" to the stem it was given left `w:themeShade="BF"`
    on Title and Heading 1. Word applies the theme shade to whatever `w:val` says, so the
    explicit black was being tinted by a value nobody had asked for. `w:shd` is the trap in
    the other direction: there the twins really are `w:themeFillTint`/`w:themeFillShade`.
    """
    _tool(tmp_path)(
        path="th.docx",
        markdown="# 标题\n\n# 一级\n\n## 二级\n\n### 三级\n\n#### 四级\n\n> 引用\n\n"
        "| A |\n|---|\n| 甲 |\n\n```\nx\n```\n",
    )
    parts = _parts(tmp_path / "th.docx")

    for style_id in (
        "Normal",
        "Title",
        "Heading1",
        "Heading2",
        "Heading3",
        "Heading4",
        "Quote",
    ):
        block = _style_xml(parts["word/styles.xml"], style_id)
        assert "theme" not in block.lower(), f"{style_id}: {block}"

    # every shading this tool writes carries only the explicit fill
    shades = re.findall(r"<w:shd[^/]*/>", parts["word/document.xml"])
    assert shades, "no shading was written at all"
    for shade in shades:
        assert "theme" not in shade.lower(), shade
        assert 'w:val="clear"' in shade and 'w:color="auto"' in shade


def test_every_property_element_keeps_its_schema_child_order(tmp_path):
    """The failure this catches is the expensive one: ECMA-376 fixes the child order of
    `w:pPr`, `w:tblPr`, `w:tcPr` and `w:pBdr`, Word refuses a document that gets it wrong
    ("unreadable content", offer to repair) — and `zipfile`, python-docx and this test file's
    other assertions all pass right through it. Anything built with `OxmlElement` and
    appended instead of inserted lands here.
    """
    import lxml.etree as etree

    tool = _tool(tmp_path)
    tool(path="a.docx", markdown=_REPORT)
    tool(path="b.docx", markdown=_REPORT, style="plain")

    sequences = {
        f"{_W}pPr": _PPR_SEQ,
        f"{_W}tblPr": _TBLPR_SEQ,
        f"{_W}tcPr": _TCPR_SEQ,
        f"{_W}pBdr": _PBDR_SEQ,
    }
    checked = 0
    for name in ("a.docx", "b.docx"):
        with zipfile.ZipFile(tmp_path / name) as archive:
            for part in [n for n in archive.namelist() if n.endswith(".xml")]:
                root = etree.fromstring(archive.read(part))
                for tag, sequence in sequences.items():
                    rank = {f"{_W}{t.split(':', 1)[1]}": i for i, t in enumerate(sequence)}
                    for node in root.iter(tag):
                        seen = [rank[c.tag] for c in node if c.tag in rank]
                        assert seen == sorted(seen), (
                            f"{name}/{part}: {tag} children out of schema order: "
                            f"{[c.tag for c in node]}"
                        )
                        checked += 1
    assert checked > 50, checked  # the walk actually found the elements


# -- ordered lists restart -------------------------------------------------------


def test_each_ordered_list_restarts_at_one(tmp_path):
    """`List Number` names ONE shared `w:num`, so two numbered lists in a document would
    run 1,2,3,4,5,6 — the single most obvious way a generated report looks wrong. Each list
    gets a fresh `w:num` over the same `abstractNum`, with `startOverride`."""
    markdown = "1. 甲\n2. 乙\n\n段落。\n\n1. 丙\n2. 丁\n"
    _tool(tmp_path)(path="n.docx", markdown=markdown)
    target = tmp_path / "n.docx"

    parts = _parts(target)
    numbering = parts["word/numbering.xml"]
    overrides = re.findall(
        r'<w:num w:numId="(\d+)">\s*<w:abstractNumId w:val="(\d+)"/>\s*'
        r'<w:lvlOverride w:ilvl="0">\s*<w:startOverride w:val="1"/>',
        numbering,
    )
    assert len(overrides) == 2, numbering  # one per list, and no more
    assert overrides[0][1] == overrides[1][1]  # both reuse List Number's own abstractNum
    assert overrides[0][0] != overrides[1][0]  # but they are separate numbering instances

    document = _open(target)
    ids = [
        p._p.pPr.numPr.numId.val
        for p in document.paragraphs
        if p.style.name == "List Number"
    ]
    assert len(ids) == 4
    assert ids[0] == ids[1] and ids[2] == ids[3] and ids[0] != ids[2]
    assert set(ids) == {int(overrides[0][0]), int(overrides[1][0])}


def test_a_bullet_list_leaves_the_styles_own_numbering_alone(tmp_path):
    """Bullets do not count, so there is nothing to restart — and injecting a `w:numPr` a
    bullet style already supplies is how you end up with two glyphs per item."""
    _tool(tmp_path)(path="b.docx", markdown="- 甲\n- 乙\n")
    document = _open(tmp_path / "b.docx")
    for paragraph in document.paragraphs:
        if paragraph.style.name == "List Bullet":
            assert paragraph._p.pPr.numPr is None
    assert "<w:startOverride" not in _parts(tmp_path / "b.docx")["word/numbering.xml"]


def test_nesting_past_three_levels_lands_on_the_third_list_style(tmp_path):
    """Word ships three levels of list style. A fourth level of nesting is clamped rather
    than refused: the content still reads, which is more than a rejected call gives."""
    markdown = "- 一\n  - 二\n    - 三\n      - 四\n"
    _tool(tmp_path)(path="d.docx", markdown=markdown)
    document = _open(tmp_path / "d.docx")
    assert [p.style.name for p in document.paragraphs] == [
        "List Bullet",
        "List Bullet 2",
        "List Bullet 3",
        "List Bullet 3",
    ]


def test_a_list_item_that_opens_with_a_nested_list_still_gets_its_own_marker(tmp_path):
    """`- ` with only a sub-list under it gives a `list_item` containing NO paragraph, so
    nothing consumed the item's own turn and the outer bullet vanished — the sub-items
    appeared to hang off nothing. The item gets an empty marker line instead, and it has to
    come BEFORE the nested content."""
    _tool(tmp_path)(path="nf.docx", markdown="- \n  - 甲\n  - 乙\n- 正常\n")
    document = _open(tmp_path / "nf.docx")
    assert [(p.style.name, p.text) for p in document.paragraphs] == [
        ("List Bullet", ""),
        ("List Bullet 2", "甲"),
        ("List Bullet 2", "乙"),
        ("List Bullet", "正常"),
    ]


def test_an_ordered_item_that_opens_with_a_nested_list_keeps_the_numbering(tmp_path):
    """The empty marker line has to take the item's `w:numId`, or the outer list would
    number 1, 1 instead of 1, 2."""
    _tool(tmp_path)(path="no.docx", markdown="1. \n   1. 甲\n2. 乙\n")
    document = _open(tmp_path / "no.docx")
    outer = [p for p in document.paragraphs if p.style.name == "List Number"]
    assert [p.text for p in outer] == ["", "乙"]
    assert outer[0]._p.pPr.numPr.numId.val == outer[1]._p.pPr.numPr.numId.val
    nested = [p for p in document.paragraphs if p.style.name == "List Number 2"]
    assert [p.text for p in nested] == ["甲"]
    # the nested list is its own numbering instance, so it starts at 1 too
    assert nested[0]._p.pPr.numPr.numId.val != outer[0]._p.pPr.numPr.numId.val


def test_a_completely_empty_list_item_still_occupies_a_line(tmp_path):
    _tool(tmp_path)(path="ne.docx", markdown="- \n- 乙\n")
    assert [(p.style.name, p.text) for p in _open(tmp_path / "ne.docx").paragraphs] == [
        ("List Bullet", ""),
        ("List Bullet", "乙"),
    ]


def test_a_second_paragraph_in_one_list_item_does_not_draw_a_second_bullet(tmp_path):
    _tool(tmp_path)(path="l.docx", markdown="1. 第一段\n\n   第二段\n\n2. 下一项\n")
    document = _open(tmp_path / "l.docx")
    items = [p for p in document.paragraphs if p.style.name == "List Number"]
    assert [p.text for p in items] == ["第一段", "第二段", "下一项"]
    # numId 0 is OOXML for "no numbering": the indent is kept, the number is not repeated
    assert items[1]._p.pPr.numPr.numId.val == 0
    assert items[0]._p.pPr.numPr.numId.val == items[2]._p.pPr.numPr.numId.val != 0


# -- table merge markers ---------------------------------------------------------

_HEAD = "| A | B | C |\n|---|---|---|\n"


def test_merge_markers_build_rectangles(tmp_path):
    """Both directions, and a 2x2 block — the shape a two-level report header needs."""
    markdown = (
        "| 甲 | 乙 | < | 丙 |\n"
        "|----|----|---|----|\n"
        "| ^  | 丁 | 戊 | ^  |\n"
        "| 己 | 庚 | <  | 辛 |\n"
    )
    _tool(tmp_path)(path="m.docx", markdown=markdown)
    table = _open(tmp_path / "m.docx").tables[0]
    assert table.cell(0, 1)._tc is table.cell(0, 2)._tc  # 乙 over two columns
    assert table.cell(0, 0)._tc is table.cell(1, 0)._tc  # 甲 over two rows
    assert table.cell(0, 3)._tc is table.cell(1, 3)._tc  # 丙 over two rows
    assert table.cell(2, 1)._tc is table.cell(2, 2)._tc  # 庚 over two columns
    assert [c.text for c in table.rows[0].cells] == ["甲", "乙", "乙", "丙"]
    assert [c.text for c in table.rows[1].cells] == ["甲", "丁", "戊", "丙"]


def test_a_two_by_two_merge_block_is_one_cell(tmp_path):
    markdown = (
        "| 甲 | < | 丙 |\n|---|---|---|\n| ^ | ^ | 丁 |\n| 戊 | 己 | 庚 |\n"
    )
    _tool(tmp_path)(path="m2.docx", markdown=markdown)
    table = _open(tmp_path / "m2.docx").tables[0]
    anchor = table.cell(0, 0)._tc
    assert all(
        table.cell(r, c)._tc is anchor for r in (0, 1) for c in (0, 1)
    )
    assert table.cell(0, 0).text == "甲"
    assert table.cell(2, 0).text == "戊"


@pytest.mark.parametrize(
    "markdown,message",
    [
        (
            f"{_HEAD}| < | b | c |\n",
            "row 2, column 1: < has no cell to its left",
        ),
        (
            "| ^ | B | C |\n|---|---|---|\n| a | b | c |\n",
            "row 1, column 1: ^ has no cell above it",
        ),
        (
            # the cell above this ^ is itself a < merged leftwards, and THAT block does not
            # extend down (row 3 column 1 is content), so the ^ has nothing to join
            f"{_HEAD}| a | < | c |\n| d | ^ | f |\n",
            "row 3, column 2: the ^ marker is not part of any merged block",
        ),
        (
            # the cell left of this < is a ^ claimed by the header above it, so the < is
            # asking to extend a block it is not part of
            f"{_HEAD}| a | ^ | < |\n",
            "row 2, column 3: the < marker is not part of any merged block",
        ),
    ],
    ids=["left-edge", "top-edge", "orphan-caret", "orphan-pair"],
)
def test_a_marker_that_belongs_to_no_block_is_refused_by_name(
    tmp_path, markdown, message
):
    """Refusing costs one turn and the model resends the table; accepting silently reshapes
    the header of a document the user is about to forward."""
    with pytest.raises(ValueError) as exc:
        _tool(tmp_path)(path="x.docx", markdown=markdown)
    assert message in str(exc.value), str(exc.value)
    assert "table 1" in str(exc.value)
    assert "merges into the cell on its left" in str(exc.value)  # the rule, restated
    assert list(tmp_path.iterdir()) == []  # nothing half-written


def test_the_refusal_names_which_table_went_wrong(tmp_path):
    markdown = f"{_HEAD}| a | b | c |\n\n段落。\n\n{_HEAD}| < | b | c |\n"
    with pytest.raises(ValueError, match="table 2, row 2, column 1"):
        _tool(tmp_path)(path="x.docx", markdown=markdown)


def test_a_caret_merges_whatever_is_above_it_body_rows_included(tmp_path):
    """`^` is not a header feature: it merges the cell above it wherever it appears, which
    is how a repeated category down the left of a data table is written."""
    markdown = f"{_HEAD}| 甲 | 一 | x |\n| ^ | 二 | y |\n| 乙 | 三 | z |\n"
    _tool(tmp_path)(path="v.docx", markdown=markdown)
    table = _open(tmp_path / "v.docx").tables[0]
    assert table.cell(1, 0)._tc is table.cell(2, 0)._tc
    assert table.cell(1, 0).text == "甲"
    assert table.cell(3, 0).text == "乙"
    assert [c.text for c in table.rows[2].cells] == ["甲", "二", "y"]


def test_a_lone_angle_bracket_is_only_a_marker_when_it_is_the_whole_cell(tmp_path):
    """`<` next to anything else is content. A cell reading "< 100" is a comparison, and a
    report full of those must not turn into a merged mess."""
    markdown = f"{_HEAD}| < 100 | b<c | a ^ b |\n"
    _tool(tmp_path)(path="ok.docx", markdown=markdown)
    table = _open(tmp_path / "ok.docx").tables[0]
    assert [c.text for c in table.rows[1].cells] == ["< 100", "b<c", "a ^ b"]
    assert table.cell(1, 0)._tc is not table.cell(1, 1)._tc


# -- GFM table details -----------------------------------------------------------


def test_ragged_rows_are_squared_off_against_the_header(tmp_path):
    markdown = "| A | B | C |\n|---|---|---|\n| 1 |\n| 1 | 2 | 3 | 4 |\n"
    _tool(tmp_path)(path="r.docx", markdown=markdown)
    table = _open(tmp_path / "r.docx").tables[0]
    assert len(table.columns) == 3
    assert [c.text for c in table.rows[1].cells] == ["1", "", ""]  # short row padded
    assert [c.text for c in table.rows[2].cells] == ["1", "2", "3"]  # long row cut


def test_an_escaped_pipe_stays_inside_its_cell(tmp_path):
    _tool(tmp_path)(path="p.docx", markdown="| A | B |\n|---|---|\n| a \\| b | c |\n")
    table = _open(tmp_path / "p.docx").tables[0]
    assert [c.text for c in table.rows[1].cells] == ["a | b", "c"]


def test_the_delimiter_rows_colons_set_the_column_alignment(tmp_path):
    from docx.enum.text import WD_ALIGN_PARAGRAPH as ALIGN

    markdown = (
        "| 左 | 中 | 右 | 默认 |\n|:---|:--:|---:|----|\n| a | b | c | d |\n"
    )
    _tool(tmp_path)(path="a.docx", markdown=markdown)
    table = _open(tmp_path / "a.docx").tables[0]
    assert [c.paragraphs[0].alignment for c in table.rows[1].cells] == [
        ALIGN.LEFT,
        ALIGN.CENTER,
        ALIGN.RIGHT,
        None,  # no colon: whatever the style says, which for text is left
    ]
    # the header row is centred whatever the column says
    assert {c.paragraphs[0].alignment for c in table.rows[0].cells} == {ALIGN.CENTER}


def test_a_table_over_the_row_limit_is_refused_with_somewhere_to_put_the_data(tmp_path):
    rows = "".join(f"| {i} | x |\n" for i in range(5_001))
    with pytest.raises(ValueError) as exc:
        _tool(tmp_path)(path="big.docx", markdown=f"| A | B |\n|---|---|\n{rows}")
    assert "over the 5000 limit" in str(exc.value)
    assert "write_spreadsheet" in str(exc.value)
    assert list(tmp_path.iterdir()) == []


def test_a_long_table_repeats_its_header_row_across_pages(tmp_path):
    """`w:tblHeader` on the header row's `w:trPr`. Word will not infer it, and a 200-row
    report whose second page has unlabelled columns is the sort of thing the reader
    notices and the author does not."""
    rows = "".join(f"| 项目{i} | {i} |\n" for i in range(60))
    _tool(tmp_path)(path="long.docx", markdown=f"| 项目 | 金额 |\n|---|---|\n{rows}")
    body = _parts(tmp_path / "long.docx")["word/document.xml"]
    assert body.count("<w:tblHeader/>") == 1
    first_row = re.search(r"<w:tr\b.*?</w:tr>", body, re.S).group(0)
    assert "<w:tblHeader/>" in first_row


def test_cells_carry_the_office_cell_size_and_no_first_line_indent(tmp_path):
    _tool(tmp_path)(path="c.docx", markdown="| A |\n|---|\n| 甲 |\n")
    table = _open(tmp_path / "c.docx").tables[0]
    from docx.shared import Pt

    assert table.cell(1, 0).paragraphs[0].runs[0].font.size == Pt(10.5)  # 五号
    body = _parts(tmp_path / "c.docx")["word/document.xml"]
    assert 'w:firstLineChars="0"' in body  # set per cell paragraph, not inherited


# -- inline ----------------------------------------------------------------------


def test_a_web_link_becomes_a_real_hyperlink_with_a_relationship(tmp_path):
    url = "https://example.invalid/report?q=1"
    _tool(tmp_path)(path="h.docx", markdown=f"见[季度报告]({url})。")
    parts = _parts(tmp_path / "h.docx")
    body, rels = parts["word/document.xml"], parts["word/_rels/document.xml.rels"]
    match = re.search(r'<w:hyperlink r:id="(rId\d+)">(.*?)</w:hyperlink>', body, re.S)
    assert match, body
    assert "季度报告" in match.group(2)
    assert '<w:u w:val="single"/>' in match.group(2)
    assert 'w:val="0563C1"' in match.group(2)
    # the id resolves to an EXTERNAL relationship: an internal one is a dead click
    assert re.search(
        rf'Id="{match.group(1)}"[^>]*Target="{re.escape(url)}"[^>]*'
        r'TargetMode="External"',
        rels,
    ), rels
    assert _open(tmp_path / "h.docx").paragraphs[0].text == "见季度报告。"


@pytest.mark.parametrize(
    "target", ["./report.docx", "#section-2", "ftp://host/file", "tel:+8610000"]
)
def test_a_link_that_is_not_a_web_address_degrades_to_text(tmp_path, target):
    """A relative path or an anchor cannot be clicked out of a document that has been
    emailed somewhere else, and a `w:hyperlink` to one is a broken promise. The reader gets
    the target spelled out instead."""
    _tool(tmp_path)(path="x.docx", markdown=f"见[附件]({target})。")
    body = _parts(tmp_path / "x.docx")["word/document.xml"]
    assert "<w:hyperlink" not in body
    assert _open(tmp_path / "x.docx").paragraphs[0].text == f"见附件 ({target})。"


def test_a_javascript_url_never_reaches_the_document_at_all(tmp_path):
    """markdown-it refuses to build a link for javascript:/vbscript:/data:, so the markdown
    stays literal text — pinned here because it is a security property this tool leans on
    rather than re-implements."""
    _tool(tmp_path)(path="j.docx", markdown="点[这里](javascript:alert(1))。")
    body = _parts(tmp_path / "j.docx")["word/document.xml"]
    assert "<w:hyperlink" not in body and "javascript" in body
    assert _open(tmp_path / "j.docx").paragraphs[0].text == "点[这里](javascript:alert(1))。"


def test_marks_become_run_formatting_rather_than_characters(tmp_path):
    _tool(tmp_path)(
        path="f.docx",
        markdown="纯文本 **粗** *斜* `码` ~~删~~ 结束",
    )
    paragraph = _open(tmp_path / "f.docx").paragraphs[0]
    assert paragraph.text == "纯文本 粗 斜 码 删 结束"
    found = {}
    for run in paragraph.runs:
        if run.text == "粗":
            found["bold"] = run.bold
        if run.text == "斜":
            found["italic"] = run.italic
        if run.text == "删":
            found["strike"] = run.font.strike
        if run.text == "码":
            found["code"] = run.font.name
    assert found == {"bold": True, "italic": True, "strike": True, "code": "Consolas"}


def test_a_hard_break_is_a_line_break_and_a_soft_one_is_not(tmp_path):
    # two trailing spaces and a trailing backslash are CommonMark's two hard breaks
    markdown = "甲" + "  \n" + "乙\n\n丙\\\n丁\n"
    _tool(tmp_path)(path="br.docx", markdown=markdown)
    document = _open(tmp_path / "br.docx")
    assert [p.text for p in document.paragraphs] == ["甲\n乙", "丙\n丁"]


def test_a_soft_wrap_between_chinese_characters_does_not_become_a_space(tmp_path):
    """CommonMark says a soft break is a space, and for Latin text it has to be. Chinese
    has no inter-word space, so a report whose author wrapped the paragraph in their editor
    would otherwise come out with a visible gap at every wrapped line."""
    _tool(tmp_path)(
        path="s.docx", markdown="中文第一行\n中文第二行\n\nEnglish first\nline two\n"
    )
    document = _open(tmp_path / "s.docx")
    assert document.paragraphs[0].text == "中文第一行中文第二行"
    assert document.paragraphs[1].text == "English first line two"


def test_an_image_becomes_a_placeholder_and_the_receipt_counts_it(tmp_path):
    out = _tool(tmp_path)(
        path="i.docx", markdown="![示意图](a.png) 与 ![](b.png) 与 ![说明](c.svg)\n"
    )
    assert "Images are not embedded in this version: 3 omitted (alt text kept)." in out
    assert _open(tmp_path / "i.docx").paragraphs[0].text == (
        "[图片: 示意图] 与 [图片] 与 [图片: 说明]"
    )


def test_a_document_without_images_says_nothing_about_them(tmp_path):
    assert "Images" not in _tool(tmp_path)(path="n.docx", markdown="只有文字。\n")


# -- page breaks, HTML, headings -------------------------------------------------


@pytest.mark.parametrize("eol", ["\n", "\r\n", "\r"], ids=["lf", "crlf", "cr"])
@pytest.mark.parametrize(
    "line",
    [
        "<!-- pagebreak -->",
        "<!--pagebreak-->",
        "<!-- PAGEBREAK -->",
        "<!--   PageBreak   -->",
        "   <!-- pagebreak -->   ",
        "\t<!-- pagebreak -->",
    ],
)
def test_the_pagebreak_comment_is_recognised_however_it_is_spaced(tmp_path, line, eol):
    """The line-ending axis is not hypothetical. `_preprocess` runs BEFORE markdown-it does
    its own `\\r\\n` normalisation, splits on `\\n`, and matches a pattern anchored with
    `$` — so a CRLF document left a `\\r` on the marker line, missed the match, and printed
    `<!-- pagebreak -->` into the report as visible text with no break at all."""
    markdown = eol.join(["前一页。", line, "后一页。", ""])
    out = _tool(tmp_path)(path="pb.docx", markdown=markdown)
    assert "1 page break(s)" in out
    body = _parts(tmp_path / "pb.docx")["word/document.xml"]
    assert body.count('<w:br w:type="page"/>') == 1
    assert "pagebreak" not in body  # neither the comment nor the sentinel reaches the file
    assert [p.text for p in _open(tmp_path / "pb.docx").paragraphs] == [
        "前一页。",
        "",
        "后一页。",
    ]


@pytest.mark.parametrize("eol", ["\r\n", "\r"], ids=["crlf", "cr"])
def test_windows_line_endings_leave_every_other_block_intact(tmp_path, eol):
    """Line endings are normalised for the whole document, not just the pagebreak scan, so
    the fence tracking and the table parser see what they expect too."""
    markdown = eol.join(
        [
            "# 标题",
            "",
            "| A | B |",
            "|---|---|",
            "| 1 | 2 |",
            "",
            "```",
            "<!-- pagebreak -->",
            "```",
            "",
            "- 甲",
            "- 乙",
            "",
        ]
    )
    out = _tool(tmp_path)(path="w.docx", markdown=markdown)
    assert "1 table(s)" in out and "1 list(s)" in out
    assert "0 page break(s)" in out  # the one inside the fence is still just code
    document = _open(tmp_path / "w.docx")
    assert document.paragraphs[0].style.name == "Title"
    assert [c.text for c in document.tables[0].rows[1].cells] == ["1", "2"]
    assert "<!-- pagebreak -->" in "\n".join(p.text for p in document.paragraphs)
    assert "\r" not in _parts(tmp_path / "w.docx")["word/document.xml"]


def test_a_leading_byte_order_mark_does_not_swallow_the_title(tmp_path):
    """`\\ufeff# 标题` is not a heading — it is a paragraph that happens to start with a
    hash. A BOM arrives whenever the Markdown came from a file Notepad or Excel wrote."""
    _tool(tmp_path)(path="bom.docx", markdown=_BOM + "# 季度报告\n\n正文。\n")
    document = _open(tmp_path / "bom.docx")
    assert [(p.style.name, p.text) for p in document.paragraphs] == [
        ("Title", "季度报告"),
        ("Normal", "正文。"),
    ]
    assert _BOM not in _parts(tmp_path / "bom.docx")["word/document.xml"]


def test_a_pagebreak_comment_inside_a_code_fence_stays_code(tmp_path):
    """A Markdown cheat sheet that SHOWS the syntax must not break its own page."""
    markdown = "说明：\n\n```\n<!-- pagebreak -->\n```\n"
    out = _tool(tmp_path)(path="pf.docx", markdown=markdown)
    assert "0 page break(s)" in out
    assert "<!-- pagebreak -->" in "\n".join(
        p.text for p in _open(tmp_path / "pf.docx").paragraphs
    )


def test_a_sentinel_the_author_typed_cannot_forge_a_page_break(tmp_path):
    """The private-use sentinel is stripped from the input before it can be substituted in,
    so the only way one exists in the token stream is that this module put it there."""
    out = _tool(tmp_path)(
        path="sn.docx", markdown=f"文字 {document_module._PAGEBREAK} 文字\n"
    )
    assert "0 page break(s)" in out
    assert _PUA not in _parts(tmp_path / "sn.docx")["word/document.xml"]


def test_html_arrives_as_text_rather_than_markup(tmp_path):
    markdown = "<div>块级 HTML</div>\n\n段落含 <b>行内</b> 标记。\n"
    _tool(tmp_path)(path="ht.docx", markdown=markdown)
    document = _open(tmp_path / "ht.docx")
    assert [p.text for p in document.paragraphs] == [
        "<div>块级 HTML</div>",
        "段落含 <b>行内</b> 标记。",
    ]


def test_headings_past_the_fourth_level_land_on_heading_four(tmp_path):
    markdown = "## 二\n\n### 三\n\n#### 四\n\n##### 五\n\n###### 六\n"
    _tool(tmp_path)(path="h.docx", markdown=markdown)
    assert [p.style.name for p in _open(tmp_path / "h.docx").paragraphs] == [
        "Heading 2",
        "Heading 3",
        "Heading 4",
        "Heading 4",
        "Heading 4",
    ]


def test_only_a_leading_h1_becomes_the_title(tmp_path):
    tool = _tool(tmp_path)
    tool(path="t1.docx", markdown="# 标题\n\n正文。\n\n# 后面的一级标题\n")
    assert [p.style.name for p in _open(tmp_path / "t1.docx").paragraphs] == [
        "Title",
        "Normal",
        "Heading 1",
    ]
    # an H1 that is not the first block is an ordinary Heading 1
    tool(path="t2.docx", markdown="正文在前。\n\n# 一级标题\n")
    assert [p.style.name for p in _open(tmp_path / "t2.docx").paragraphs] == [
        "Normal",
        "Heading 1",
    ]


def test_a_blockquote_keeps_its_own_style_and_indent(tmp_path):
    _tool(tmp_path)(path="q.docx", markdown="> 引用第一段。\n>\n> 引用第二段。\n\n正文。\n")
    document = _open(tmp_path / "q.docx")
    assert [p.style.name for p in document.paragraphs] == ["Quote", "Quote", "Normal"]
    quote = _style_xml(_parts(tmp_path / "q.docx")["word/styles.xml"], "Quote")
    assert 'w:val="595959"' in quote
    assert 'w:left="419"' in quote or 'w:left="420"' in quote  # 0.74 cm


# -- the plain style -------------------------------------------------------------


def test_plain_leaves_the_stylesheet_and_the_page_alone(tmp_path):
    out = _tool(tmp_path)(path="p.docx", markdown=_REPORT, style="plain")
    assert out.endswith("style plain.")
    parts = _parts(tmp_path / "p.docx")
    styles, body = parts["word/styles.xml"], parts["word/document.xml"]
    assert "firstLineChars" not in styles and "firstLineChars" not in body
    assert "宋体" not in styles and "黑体" not in styles
    assert 'w:line="360"' not in styles
    assert '<w:pgSz w:w="11906"' not in body  # the template's own page size is kept
    # but the structure is still a document: styles, a real table, page numbers
    document = _open(tmp_path / "p.docx")
    assert [p.style.name for p in document.paragraphs][0] == "Title"
    assert len(document.tables) == 1
    assert '<w:tblW w:type="pct" w:w="5000"/>' in body
    assert body.count('w:fill="DDEBF7"') == 3
    footer = next(v for k, v in parts.items() if re.match(r"word/footer\d*\.xml", k))
    assert "PAGE" in footer


@pytest.mark.parametrize("style", ["Office", "  plain  ", None, ""])
def test_the_style_name_is_read_forgivingly(tmp_path, style):
    out = _tool(tmp_path)(path="s.docx", markdown="正文。\n", style=style)
    assert out.endswith(f"style {str(style or 'office').strip().casefold()}.")


def test_an_unknown_style_is_refused_with_both_names(tmp_path):
    with pytest.raises(ValueError) as exc:
        _tool(tmp_path)(path="s.docx", markdown="正文。\n", style="公文")
    assert '"office" or "plain"' in str(exc.value)
    assert list(tmp_path.iterdir()) == []


# -- arguments and limits --------------------------------------------------------


def test_empty_markdown_is_refused_with_a_worked_example(tmp_path):
    tool = _tool(tmp_path)
    for blank in ("", "   ", "\n\n\t\n", "\x00\x07"):
        with pytest.raises(ValueError, match="markdown is empty"):
            tool(path="e.docx", markdown=blank)
    assert "季度报告" in _err(tool, "")  # the message shows the shape that works
    assert list(tmp_path.iterdir()) == []


def test_markdown_over_the_character_limit_is_refused(tmp_path):
    with pytest.raises(ValueError) as exc:
        _tool(tmp_path)(path="l.docx", markdown="甲" * 2_000_001)
    assert "over the 2000000 limit" in str(exc.value)
    assert "write_spreadsheet" in str(exc.value)
    assert list(tmp_path.iterdir()) == []


def test_a_non_string_markdown_is_refused_not_stringified(tmp_path):
    tool = _tool(tmp_path)
    for bad in (["# 标题"], {"markdown": "x"}, 42):
        with pytest.raises(ValueError) as exc:
            tool(path="x.docx", markdown=bad)
        assert "markdown must be a string" in str(exc.value), bad
        assert type(bad).__name__ in str(exc.value), bad
    assert list(tmp_path.iterdir()) == []


def test_xml_illegal_control_characters_are_dropped_rather_than_thrown(tmp_path):
    """lxml refuses \\x07 outright. A model that pasted terminal output into a report would
    otherwise lose the whole document over one stray byte."""
    _tool(tmp_path)(path="ctl.docx", markdown="正\x07文\x00带\x1f控制符\n\n- 列\x08表\n")
    document = _open(tmp_path / "ctl.docx")
    text = "\n".join(p.text for p in document.paragraphs)
    assert text == "正文带控制符\n列表"
    assert not re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", _parts(tmp_path / "ctl.docx")["word/document.xml"])


def _err(tool, markdown: str) -> str:
    with pytest.raises(ValueError) as exc:
        tool(path="e.docx", markdown=markdown)
    return str(exc.value)


# -- paths -----------------------------------------------------------------------


def test_the_extension_is_checked_and_never_appended(tmp_path):
    tool = _tool(tmp_path)
    with pytest.raises(ValueError) as exc:
        tool(path="报告.docm", markdown="正文。")
    assert "macros" in str(exc.value) and ".docx" in str(exc.value)
    with pytest.raises(ValueError) as exc:
        tool(path="报告.doc", markdown="正文。")
    assert "old binary Word format" in str(exc.value)
    for bad in ("报告", "报告.md", "报告.DOCX.txt", "报告.odt"):
        with pytest.raises(ValueError, match=r"must end in \.docx"):
            tool(path=bad, markdown="正文。")
    assert not any(tmp_path.iterdir())  # no auto-renamed file appeared anywhere

    tool(path="报告.DOCX", markdown="正文。")  # upper case is fine
    assert [p.name for p in tmp_path.iterdir()] == ["报告.DOCX"]


def test_chinese_path_and_missing_parent_directories(tmp_path):
    out = _tool(tmp_path)(path="子目录/第二层/季度 报告 (最终).docx", markdown=_REPORT)
    target = tmp_path / "子目录" / "第二层" / "季度 报告 (最终).docx"
    assert target.is_file() and str(target) in out
    assert _open(target).paragraphs[0].text == "季度运营报告"


def test_writing_over_a_directory_is_refused(tmp_path):
    (tmp_path / "报告.docx").mkdir()
    with pytest.raises(ValueError, match="Path is a directory"):
        _tool(tmp_path)(path="报告.docx", markdown="正文。")


def _parity_roots(tmp_path):
    ws = tmp_path / "工作区"
    scratch = tmp_path / "scratch"
    ro = tmp_path / "readonly"
    for d in (ws, scratch, ro):
        d.mkdir()
    roots = normalize_roots(
        [
            RootDir(path=ws, writable=True, label="workspace"),
            RootDir(path=scratch, writable=True, label="scratch"),
            RootDir(path=ro, writable=False, label="readonly"),
        ]
    )
    return ws, scratch, ro, roots


def _outcome(call, tmp_path: Path, suffix: str):
    """('error', type, message-with-the-extension-masked) or ('ok', landing directory)."""
    try:
        call()
    except Exception as exc:  # noqa: BLE001 - the comparison is the point
        return ("error", type(exc).__name__, str(exc).replace(suffix, "<ext>"))
    found = list(tmp_path.rglob(f"probe{suffix}"))
    assert len(found) == 1, found
    return ("ok", str(found[0].parent))


_PARITY_CASES = [
    "relative",
    "relative-subdir",
    "absolute-primary",
    "absolute-secondary-root",
    "absolute-outside-every-root",
    "dot-dot-escape",
    "dot-dot-reentry",
    "read-only-root",
    "tilde-outside",
    "tilde-inside",
]


@pytest.mark.parametrize("case", _PARITY_CASES)
def test_path_parity_with_aisuite_write_file(tmp_path, monkeypatch, case):
    """One root set, two tools, the same ten inputs: accept/reject, exception type, message
    and landing directory must match. The layer under test is `office.py`'s and is checked
    there too — repeated here because `write_document` is the tool a user's approval prompt
    will be naming, and a divergence on this side is the same security answer."""
    ws, scratch, ro, roots = _parity_roots(tmp_path)
    write_file = next(
        t
        for t in ai.toolkits.files(roots=roots)
        if getattr(t, "__name__", "") == "write_file"
    )
    write_document = _tool(ws, roots)

    def path_for(suffix: str) -> str:
        if case == "relative":
            return f"probe{suffix}"
        if case == "relative-subdir":
            return f"新建文件夹/深/probe{suffix}"
        if case == "absolute-primary":
            return str(ws / f"probe{suffix}")
        if case == "absolute-secondary-root":
            return str(scratch / f"probe{suffix}")
        if case == "absolute-outside-every-root":
            return str(tmp_path / f"probe{suffix}")
        if case == "dot-dot-escape":
            return f"../probe{suffix}"
        if case == "dot-dot-reentry":
            # judged AFTER resolution, so a path that walks out and back in is fine
            return str(scratch / "子" / ".." / f"probe{suffix}")
        if case == "read-only-root":
            return str(ro / f"probe{suffix}")
        return f"~/probe{suffix}"  # both tilde cases

    if case == "tilde-inside":
        for var in ("USERPROFILE", "HOME"):
            monkeypatch.setenv(var, str(scratch))
        monkeypatch.delenv("HOMEPATH", raising=False)

    mine = _outcome(
        lambda: write_document(path=path_for(".docx"), markdown="正文。"),
        tmp_path,
        ".docx",
    )
    theirs = _outcome(
        lambda: write_file(path=path_for(".csv"), content="a\n"), tmp_path, ".csv"
    )
    assert mine == theirs, case


def test_no_writable_root_is_refused_by_name(tmp_path):
    ro = tmp_path / "ro"
    ro.mkdir()
    roots = normalize_roots([RootDir(path=ro, writable=False, label="ro")])
    with pytest.raises(PermissionError, match="write_document is disabled"):
        _tool(ro, roots)(path="x.docx", markdown="正文。")


def test_a_folder_granted_mid_session_is_reachable_without_rebuilding(tmp_path):
    """`roots` is held by reference, like every other roots consumer — a folder the user
    grants mid-turn must work on that same turn."""
    ws = tmp_path / "ws"
    granted = tmp_path / "granted"
    for d in (ws, granted):
        d.mkdir()
    roots = normalize_roots([RootDir(path=ws, writable=True, label="workspace")])
    tool = _tool(ws, roots)
    target = str(granted / "报告.docx")

    with pytest.raises(PermissionError, match="escapes allowed roots"):
        tool(path=target, markdown="正文。")
    roots.append(RootDir(path=granted, writable=True, label="granted"))
    assert "(root: granted)" in tool(path=target, markdown="正文。")


# -- atomic write ----------------------------------------------------------------


def _boom(*_a, **_k):
    raise RuntimeError("save blew up")


def test_a_failed_save_leaves_neither_a_target_nor_a_temp_file(tmp_path, monkeypatch):
    monkeypatch.setattr(document_module, "_save_document", _boom)
    with pytest.raises(RuntimeError, match="save blew up"):
        _tool(tmp_path)(path="报告.docx", markdown=_REPORT)
    assert list(tmp_path.iterdir()) == []  # no half .docx, no .part left behind


def test_a_failed_save_keeps_the_previous_file_intact(tmp_path, monkeypatch):
    tool = _tool(tmp_path)
    tool(path="报告.docx", markdown="旧数据。")
    target = tmp_path / "报告.docx"
    before = target.read_bytes()

    monkeypatch.setattr(document_module, "_save_document", _boom)
    with pytest.raises(RuntimeError):
        tool(path="报告.docx", markdown="新数据。")

    assert target.read_bytes() == before
    assert _open(target).paragraphs[0].text == "旧数据。"
    assert [p.name for p in tmp_path.iterdir()] == ["报告.docx"]


def test_a_locked_target_blames_word_rather_than_excel(tmp_path, monkeypatch):
    """The swap seam is `office.py`'s and is SHARED — which is the point of this test: the
    guess about what has the file open has to name the right program for the format."""

    def locked(_src, _dst):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(office_module, "_replace", locked)
    with pytest.raises(PermissionError) as exc:
        _tool(tmp_path)(path="报告.docx", markdown="正文。")
    assert "is open in another program (Word?)" in str(exc.value)
    assert "Excel" not in str(exc.value)
    assert str(tmp_path / "报告.docx") in str(exc.value)
    assert list(tmp_path.iterdir()) == []  # the temp file is cleaned up regardless


def test_a_file_name_the_os_rejects_is_a_value_error_not_a_bare_oserror(
    tmp_path, monkeypatch
):
    """`os.replace` answers "the file is locked" and "that name is not a file name" with
    the same exception class, and only the first is something the USER can fix. The second
    has to reach the model as a ValueError naming the path, or it gets `OSError: [WinError
    87] 参数错误` and no idea which argument was wrong.

    Driven through the seam so the conversion is covered on every platform; the Windows
    test below proves what actually triggers it.
    """

    def invalid(_src, _dst):
        raise OSError(errno.EINVAL, "Invalid argument")

    monkeypatch.setattr(office_module, "_replace", invalid)
    with pytest.raises(ValueError) as exc:
        _tool(tmp_path)(path="报告.docx", markdown="正文。")
    message = str(exc.value)
    assert "refused this file name" in message
    assert str(tmp_path / "报告.docx") in message
    assert "Invalid argument" in message
    assert list(tmp_path.iterdir()) == []  # and the .part is still cleaned up


@pytest.mark.skipif(os.name != "nt", reason="NTFS alternate data streams are Windows-only")
def test_an_alternate_data_stream_path_is_refused_by_name(tmp_path):
    """The real trigger, end to end: `a.txt:b.docx` is NTFS alternate-data-stream syntax.
    `Path.resolve` accepts it, the roots check accepts it, the temp file even writes — and
    then `os.replace` fails with WinError 87."""
    with pytest.raises(ValueError) as exc:
        _tool(tmp_path)(path="a.txt:b.docx", markdown="正文。")
    assert "refused this file name" in str(exc.value)
    assert 'no : * ? " < > | characters' in str(exc.value)
    assert not (tmp_path / "a.txt:b.docx").exists()
    # no half-written payload under a confusing name. (The zero-byte `.a.txt` host file the
    # stream hangs off is left alone on purpose: it is an ordinary name we did not choose
    # and may already have been the user's.)
    assert [p.name for p in tmp_path.iterdir() if p.name.endswith(".part")] == []


@pytest.mark.skipif(
    os.name != "nt", reason="only Windows refuses to replace a read-only file"
)
def test_a_read_only_target_is_not_blamed_on_word(tmp_path):
    """End to end, with a genuinely read-only file and the real `os.replace`."""
    tool = _tool(tmp_path)
    tool(path="报告.docx", markdown="旧数据。")
    target = tmp_path / "报告.docx"
    target.chmod(0o444)
    try:
        if os.access(target, os.W_OK):  # pragma: no cover - filesystem without the bit
            pytest.skip("this filesystem ignores the read-only bit")
        with pytest.raises(PermissionError) as exc:
            tool(path="报告.docx", markdown="新数据。")
        message = str(exc.value)
        assert "is read-only / write-protected" in message
        assert "open in another program" not in message
        assert [p.name for p in tmp_path.iterdir()] == ["报告.docx"]  # no .part left
        assert _open(target).paragraphs[0].text == "旧数据。"  # and it is untouched
    finally:
        target.chmod(0o666)


def test_overwriting_an_existing_file_is_the_documented_behaviour(tmp_path):
    tool = _tool(tmp_path)
    tool(path="报告.docx", markdown="旧。")
    tool(path="报告.docx", markdown="新。")
    assert _open(tmp_path / "报告.docx").paragraphs[0].text == "新。"


# -- the receipt -----------------------------------------------------------------


def test_the_receipt_names_the_path_the_root_and_what_was_written(tmp_path):
    ws = tmp_path / "工作区"
    scratch = tmp_path / "scratch"
    for d in (ws, scratch):
        d.mkdir()
    roots = normalize_roots(
        [
            RootDir(path=ws, writable=True, label="workspace"),
            RootDir(path=scratch, writable=True, label="scratch"),
        ]
    )
    out = _tool(ws, roots)(path=str(scratch / "报告.docx"), markdown=_REPORT)
    assert out.startswith(f"Wrote {(scratch / '报告.docx').resolve()} (root: scratch). ")
    assert (
        "Document: 12 paragraph(s), 5 heading(s), 1 table(s), 3 list(s), "
        "1 page break(s); style office." in out
    )


def test_the_deepest_root_names_the_file_like_write_file_does(tmp_path):
    outer = tmp_path / "outer"
    inner = outer / "inner"
    inner.mkdir(parents=True)
    roots = normalize_roots(
        [
            RootDir(path=outer, writable=True, label="outer"),
            RootDir(path=inner, writable=True, label="inner"),
        ]
    )
    out = _tool(outer, roots)(path=str(inner / "报告.docx"), markdown="正文。")
    assert "(root: inner)" in out


def test_a_successful_write_never_reports_failure_when_formatting_breaks(
    tmp_path, monkeypatch
):
    """Same rule as `write_file`'s wrapper: once the bytes are on disk, a receipt bug must
    not become an error the model reads as "the file does not exist"."""
    monkeypatch.setattr(
        document_module,
        "_receipt",
        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("formatting blew up")),
    )
    out = _tool(tmp_path)(path="报告.docx", markdown="正文。")
    target = tmp_path / "报告.docx"
    assert out == f"Wrote {target}"
    assert _open(target).paragraphs[0].text == "正文。"


def test_the_counts_are_zero_rather_than_missing_for_a_bare_paragraph(tmp_path):
    out = _tool(tmp_path)(path="x.docx", markdown="只有一段话。\n")
    assert out.endswith(
        "Document: 1 paragraph(s), 0 heading(s), 0 table(s), 0 list(s), "
        "0 page break(s); style office."
    )


# -- schema / import cost --------------------------------------------------------


def test_the_tool_advertises_the_schema_and_its_write_metadata(tmp_path):
    tool = _tool(tmp_path)
    assert tool.__coworker_schema__ is _SCHEMA
    meta = tool.__aisuite_tool_metadata__
    assert meta.name == "write_document"
    assert meta.category == "filesystem"
    assert meta.risk_level == "medium" and meta.requires_approval is True
    assert meta.capabilities == ["write_file"]


def test_the_schema_documents_the_merge_markers_and_stays_small():
    """The portability subset itself is checked alongside `write_spreadsheet`'s in
    `test_office_tools.py`; this is the part that is specific to this tool."""
    fn = _SCHEMA["function"]
    assert fn["parameters"]["properties"]["style"]["enum"] == ["office", "plain"]
    assert "must end with .docx" in fn["parameters"]["properties"]["path"]["description"]
    # the one extension to GFM has to be discoverable from the schema alone
    assert "merges into the cell on its left" in fn["description"]
    assert "<!-- pagebreak -->" in fn["description"]
    assert len(json.dumps(fn)) < 3000


def test_the_office_toolkit_hands_out_both_writers(tmp_path):
    tools = office_tools(str(tmp_path))
    assert [t.__name__ for t in tools] == ["write_spreadsheet", "write_document"]
    assert "Wrote " in tools[1](path="x.docx", markdown="正文。")


def test_a_build_without_python_docx_says_what_to_deliver_instead(tmp_path, monkeypatch):
    """The dependency is bundled, but a trimmed build must fail with an instruction, not a
    traceback: the model needs to hear "deliver a .md" to salvage the turn."""
    monkeypatch.setitem(sys.modules, "docx", None)
    with pytest.raises(RuntimeError) as exc:
        _tool(tmp_path)(path="报告.docx", markdown="正文。")
    assert "missing python-docx" in str(exc.value)
    assert "write_file" in str(exc.value) and ".md" in str(exc.value)
    assert list(tmp_path.iterdir()) == []


def test_importing_the_module_does_not_import_python_docx():
    """Lazy by design: a session that never writes a document must not pay the import (lxml
    alone is tens of milliseconds), and a build without it must still start."""
    code = (
        "import sys; import coworker.tools.document as m; "
        "print('docx' in sys.modules, 'markdown_it' in sys.modules, bool(m.document_tools))"
    )
    proc = subprocess.run(
        [sys.executable, "-c", code],
        cwd=str(Path(__file__).resolve().parents[1]),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    assert proc.returncode == 0, proc.stderr
    assert proc.stdout.strip() == "False False True"
