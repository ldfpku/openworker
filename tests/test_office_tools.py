"""`write_spreadsheet` — real .xlsx written in the backend process (tools/office.py).

Three things these cover, in order of how much damage they prevent:

1. **Path parity with `write_file`.** The tool re-implements aisuite's `_resolve` because a
   binary payload cannot ride through `write_file(content: str)`. `test_path_parity_*`
   drives both tools over one root set and compares accept/reject, exception type, message
   wording and the directory the bytes actually land in — the two must never drift.
2. **The file opens.** Every write is round-tripped with openpyxl and checked as a ZIP:
   a corrupt .xlsx looks like success to the agent and like a broken attachment to the user.
3. **Values mean what the user typed.** An 18-digit order number that comes back as
   1.23457E+17 is data loss the user discovers days later.
"""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import zipfile
from datetime import datetime
from pathlib import Path

import pytest

import aisuite as ai
from coworker.roots import RootDir, normalize_roots
from coworker.tools import office as office_module
from coworker.tools.document import _SCHEMA as _DOC_SCHEMA
from coworker.tools.office import _SCHEMA, office_tools


def _tool(workspace, roots=None):
    """`write_spreadsheet`, plus the standing check that it is still the FIRST tool the
    toolkit hands out — `tests/test_document_tool.py` reaches for `[1]` by the same token,
    and App.tsx's FILE_WRITE_TOOLS list is written against both names."""
    tools = office_tools(str(workspace), roots=roots)
    assert [getattr(t, "__name__", "") for t in tools] == [
        "write_spreadsheet",
        "write_document",
    ]
    return tools[0]


def _load(path: Path):
    from openpyxl import load_workbook

    assert zipfile.ZipFile(path).testzip() is None  # a readable ZIP container
    return load_workbook(path)


def _sheet_xml(path: Path, index: int = 1) -> str:
    """The raw worksheet part. Some corruption is invisible to openpyxl on read-back —
    an inline string cell with no <is> child loads fine here and makes Excel offer to
    repair the file — so the bytes themselves have to be inspected."""
    with zipfile.ZipFile(path) as archive:
        return archive.read(f"xl/worksheets/sheet{index}.xml").decode("utf-8")


def _assert_no_empty_inline_strings(path: Path, index: int = 1) -> None:
    xml = _sheet_xml(path, index)
    assert re.search(r'<c[^>]*t="inlineStr"[^>]*/>', xml) is None, xml
    assert xml.count('t="inlineStr"') == xml.count("<is>"), xml


def _assert_framed(sheet, *, max_row: int, max_col: int) -> None:
    """Every used cell is boxed in, and every merged range is boxed in as a whole.

    The tool sets all four sides on all four sides of each merge member; openpyxl then
    strips the INTERIOR sides on save (a merged range is drawn as one box), which is the
    behaviour being pinned here — the visible result is a frame, never a merged title open
    on three sides.
    """
    boxes = [
        (r.min_col, r.min_row, r.max_col, r.max_row) for r in sheet.merged_cells.ranges
    ]
    for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            c, r = cell.column, cell.row
            home = next(
                (b for b in boxes if b[0] <= c <= b[2] and b[1] <= r <= b[3]), None
            )
            sides = cell.border
            if home is None:
                assert (
                    sides.left.style
                    == sides.right.style
                    == sides.top.style
                    == sides.bottom.style
                    == "thin"
                ), cell.coordinate
                continue
            min_col, min_row, max_c, max_r = home
            if c == min_col:
                assert sides.left.style == "thin", cell.coordinate
            if c == max_c:
                assert sides.right.style == "thin", cell.coordinate
            if r == min_row:
                assert sides.top.style == "thin", cell.coordinate
            if r == max_r:
                assert sides.bottom.style == "thin", cell.coordinate


# -- the typical report template -------------------------------------------------

_TEMPLATE = {
    "name": "季度汇总",
    "rows": [
        ["2026 年第三季度部门报告"],
        ["数据截止 2026-09-18"],
        ["项目名称", "负责部门", "金额", "占比", "开始日期", "备注"],
        ["项目A", "部门一", "1,234.50", "12.5%", "2026-09-18", "00123"],
        ["项目B", "部门二", "987", "7%", "2026-09-19", "换行\n备注"],
        ["合计", "", "=SUM(C4:C5)", "", "", ""],
    ],
    "merges": ["A1:F1", "A2:F2"],
    "header_rows": 1,
}


def test_report_template_round_trips_with_its_default_look(tmp_path):
    out = _tool(tmp_path)(path="季度报告.xlsx", sheets=[dict(_TEMPLATE)])
    target = tmp_path / "季度报告.xlsx"
    assert target.is_file()
    assert out.startswith(f"Wrote {target}")

    sheet = _load(target)["季度汇总"]  # the Chinese sheet name survived
    assert {str(r) for r in sheet.merged_cells.ranges} == {"A1:F1", "A2:F2"}

    # title: bold 14, centred both ways, and NO fill (the shading belongs to the header)
    assert sheet["A1"].value == "2026 年第三季度部门报告"
    assert sheet["A1"].font.b is True and sheet["A1"].font.sz == 14
    assert sheet["A1"].alignment.horizontal == "center"
    assert sheet["A1"].alignment.vertical == "center"
    assert sheet["A1"].fill.fill_type in (None, "none")

    # subtitle: regular weight, left aligned
    assert not sheet["A2"].font.b
    assert sheet["A2"].alignment.horizontal == "left"

    # header: bold, shaded, centred, wrapping — and the freeze sits right below it
    for ref in ("A3", "F3"):
        assert sheet[ref].font.b is True
        assert sheet[ref].fill.start_color.rgb == "FFDDEBF7"
        assert sheet[ref].alignment.horizontal == "center"
        assert sheet[ref].alignment.wrap_text is True
    assert sheet.freeze_panes == "A4"

    # data: vertically centred; a cell with a newline wraps
    assert sheet["A4"].alignment.vertical == "center"
    assert sheet["F5"].alignment.wrap_text is True
    assert sheet["A4"].alignment.wrap_text in (None, False)

    _assert_framed(sheet, max_row=6, max_col=6)

    # auto width counts a CJK glyph as two: "项目名称" is 8, not 4
    assert sheet.column_dimensions["A"].width == 10


def test_receipt_names_the_absolute_path_its_root_and_each_sheet(tmp_path):
    ws = tmp_path / "工作区"
    scratch = tmp_path / "scratch"
    for d in (ws, scratch):
        d.mkdir()
    roots = normalize_roots(
        [
            RootDir(path=ws, writable=True, label="workspace"),
            RootDir(path=scratch, writable=True, label="scratch"),
        ]
    )
    out = _tool(ws, roots)(
        path=str(scratch / "报告.xlsx"),
        sheets=[dict(_TEMPLATE), {"name": "明细", "rows": [["a", "b"], ["1", "2"]]}],
    )

    assert out.startswith(f"Wrote {(scratch / '报告.xlsx').resolve()} (root: scratch).")
    assert "季度汇总 6 rows x 6 cols, 2 merges" in out
    assert "明细 2 rows x 2 cols" in out
    # a formula was written, so the model is told the values are computed on open
    assert "Formulas are calculated when the file is opened in Excel/WPS." in out


def test_receipt_omits_the_formula_note_when_there_are_no_formulas(tmp_path):
    out = _tool(tmp_path)(path="x.xlsx", sheets=[{"rows": [["a"], ["1"]]}])
    assert "Formulas" not in out
    assert "Sheets: Sheet1 2 rows x 1 cols" in out


def test_a_successful_write_never_reports_failure_when_formatting_breaks(
    tmp_path, monkeypatch
):
    """Same rule as `write_file`'s wrapper: once the bytes are on disk, a receipt bug must
    not become an error the model reads as "the file does not exist"."""
    monkeypatch.setattr(
        office_module,
        "_receipt",
        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("formatting blew up")),
    )
    out = _tool(tmp_path)(path="报告.xlsx", sheets=[{"rows": [["a"]]}])
    target = tmp_path / "报告.xlsx"
    assert out == f"Wrote {target}"
    assert _load(target)["Sheet1"]["A1"].value == "a"


# -- cell values -----------------------------------------------------------------

# (input, expected value, expected data_type, expected number_format or None = any)
_VALUES = [
    ("", None, "n", None),
    ("hello 你好", "hello 你好", "s", None),
    ("00123", "00123", "s", None),  # leading zero: stays text, no data lost
    ("123456789012345678", "123456789012345678", "s", None),  # 18 digits > float64
    ("'00456", "00456", "s", None),  # the explicit escape hatch
    ("'=SUM(A1:A2)", "=SUM(A1:A2)", "s", None),  # quoted formula is TEXT
    ("=SUM(A1:A2)", "=SUM(A1:A2)", "f", None),
    ("=", "=", "s", None),  # a lone "=" is not a formula
    ("0", 0, "n", "General"),
    ("42", 42, "n", "General"),
    ("-0.5", -0.5, "n", "General"),
    ("0.5", 0.5, "n", "General"),
    ("1,234", 1234, "n", "#,##0"),
    ("1,234.50", 1234.5, "n", "#,##0.00"),
    ("-1,234.5", -1234.5, "n", "#,##0.0"),
    ("12%", 0.12, "n", "0%"),
    ("12.5%", 0.125, "n", "0.0%"),
    ("2026-09-18", datetime(2026, 9, 18), "d", "yyyy-mm-dd"),
    ("2026-09-18 14:30", datetime(2026, 9, 18, 14, 30), "d", "yyyy-mm-dd hh:mm"),
    (
        "2026-09-18 14:30:05",
        datetime(2026, 9, 18, 14, 30, 5),
        "d",
        "yyyy-mm-dd hh:mm:ss",
    ),
    ("2026-02-30", "2026-02-30", "s", None),  # not a real date: keep the string
    ("1.2.3", "1.2.3", "s", None),
    ("+5", "+5", "s", None),
    (" 12 ", " 12 ", "s", None),
    ("1e5", "1e5", "s", None),
    ("0012.5", "0012.5", "s", None),
    ("abc\x07def", "abcdef", "s", None),  # openpyxl refuses \x07 outright
    (42, 42, "n", None),
    (1234567890123456789, "1234567890123456789", "s", None),  # 19-digit int → text
    (1.5, 1.5, "n", None),
    (True, True, "b", None),
    (None, None, "n", None),
]


def test_every_cell_value_rule_has_a_worked_example(tmp_path):
    rows = [[raw] for raw, _v, _t, _f in _VALUES]
    out = _tool(tmp_path)(
        path="值.xlsx", sheets=[{"rows": rows, "header_rows": 0, "borders": False}]
    )
    assert out.startswith("Wrote ")
    sheet = _load(tmp_path / "值.xlsx")["Sheet1"]

    for i, (raw, expected, data_type, number_format) in enumerate(_VALUES, 1):
        cell = sheet.cell(row=i, column=1)
        label = f"row {i} (input {raw!r})"
        if isinstance(expected, float) and not isinstance(expected, bool):
            assert cell.value == pytest.approx(expected), label
        else:
            assert cell.value == expected, label
            assert type(cell.value) is type(expected), label
        assert cell.data_type == data_type, label
        if number_format is not None:
            assert cell.number_format == number_format, label


def test_a_lone_apostrophe_is_an_empty_cell_not_a_corrupt_one(tmp_path):
    """The escape hatch with nothing behind it. Forced text plus an empty string made
    openpyxl write `<c t="inlineStr">` with no `<is>` child, and Excel condemns the whole
    workbook over one such cell — so it has to be an empty cell instead."""
    out = _tool(tmp_path)(
        path="撇号.xlsx",
        sheets=[
            {
                "rows": [
                    ["'", "'\x07", "'x", "", "'0"],  # "'\x07" is a lone quote too
                    ["正常", "", "'", "x", ""],
                ],
                "header_rows": 0,
            }
        ],
    )
    assert out.startswith("Wrote ")
    target = tmp_path / "撇号.xlsx"
    _assert_no_empty_inline_strings(target)
    sheet = _load(target)["Sheet1"]
    assert sheet["A1"].value is None and sheet["B1"].value is None
    assert sheet["C1"].value == "x" and sheet["E1"].value == "0"
    assert sheet["C1"].data_type == "s" and sheet["E1"].data_type == "s"
    assert sheet["C2"].value is None
    assert sheet["A2"].value == "正常"


def test_a_truncated_formula_is_refused_instead_of_written(tmp_path):
    """A model that runs out of output budget mid-cell sends `=SUM(B2:B`. openpyxl copies
    it into <f> verbatim and Excel then refuses the FILE, so refusing the call — and
    naming the cell — is the cheap outcome."""
    tool = _tool(tmp_path)
    for formula in ("=SUM(", "=SUM(B2:B", '=IF(A1>1,"y","n"', '=CONCAT("a)', "=A1)"):
        with pytest.raises(ValueError) as exc:
            tool(
                path="f.xlsx",
                sheets=[{"name": "汇总", "rows": [["项目", "金额"], ["项目A", formula]]}],
            )
        message = str(exc.value)
        assert "unbalanced parentheses or quotes" in message, formula
        assert "cell B2" in message and "汇总" in message, formula
        assert "prefix ' to write it as text" in message, formula
    with pytest.raises(ValueError, match="over Excel's 8192 limit"):
        tool(path="f.xlsx", sheets=[{"rows": [["=SUM(" + "A1," * 3000 + "A2)"]]}])
    assert not any(tmp_path.iterdir())

    # a formula whose parentheses live inside a STRING literal is fine, and so is the
    # escape hatch for a formula the model wants to show rather than compute
    tool(
        path="f.xlsx",
        sheets=[
            {
                "rows": [
                    ['=IF(A2>1,"yes (maybe)","no")'],
                    ["'=SUM("],
                    ["=SUM(A1:A2)"],
                ],
                "header_rows": 0,
            }
        ],
    )
    sheet = _load(tmp_path / "f.xlsx")["Sheet1"]
    assert sheet["A1"].data_type == "f"
    assert sheet["A2"].value == "=SUM(" and sheet["A2"].data_type == "s"
    assert sheet["A3"].data_type == "f"


def test_an_unterminated_number_format_is_refused(tmp_path):
    """Same failure one layer down: a half-typed format string lands in styles.xml, which
    Excel cannot report per-cell — it declares the file unreadable."""
    tool = _tool(tmp_path)
    for fmt in ("0.00[unterminated", '0.00"abc', "0.00]", "0.00\x07"):
        with pytest.raises(ValueError) as exc:
            tool(
                path="nf.xlsx",
                sheets=[{"rows": [["a"], ["1"]], "styles": [{"range": "A2", "format": fmt}]}],
            )
        message = str(exc.value)
        assert "styles[0] format" in message, fmt
        assert ("unbalanced [] brackets or quotes" in message) or (
            "control characters" in message
        ), fmt
    assert not any(tmp_path.iterdir())

    # the formats a report actually uses, brackets and all
    tool(
        path="nf.xlsx",
        sheets=[
            {
                "rows": [["金额"], ["-1234.5"]],
                "styles": [
                    {"range": "A2", "format": "#,##0.00;[Red]-#,##0.00"},
                    {"range": "A1", "format": 'yyyy"年"mm"月"'},
                ],
            }
        ],
    )
    sheet = _load(tmp_path / "nf.xlsx")["Sheet1"]
    assert sheet["A2"].number_format == "#,##0.00;[Red]-#,##0.00"
    assert sheet["A1"].number_format == 'yyyy"年"mm"月"'


def _format_writer(tool):
    """`tool(...)` with one `styles[0].format` under test and everything else fixed."""

    def write(number_format):
        return tool(
            path="nf.xlsx",
            sheets=[
                {
                    "rows": [["金额"], ["1234.5"]],
                    "styles": [{"range": "A2", "format": number_format}],
                }
            ],
        )

    return write


def test_a_number_format_over_excels_255_character_ceiling_is_refused(tmp_path):
    """openpyxl writes a format code of any length; Excel's own limit is 255, and a longer
    one is another whole-file repair prompt."""
    write = _format_writer(_tool(tmp_path))
    with pytest.raises(ValueError) as exc:
        write("0." + "0" * 300)
    assert "over Excel's 255 limit" in str(exc.value)
    assert "styles[0] format" in str(exc.value)
    assert not any(tmp_path.iterdir())

    # right at the ceiling it still goes through
    at_limit = "0." + "0" * 253
    assert len(at_limit) == 255
    write(at_limit)
    assert _load(tmp_path / "nf.xlsx")["Sheet1"]["A2"].number_format == at_limit


def test_a_non_string_number_format_is_refused_not_stringified(tmp_path):
    """`str(["0.00"])` is the literal format code `['0.00']`: it passes every structural
    check and the user gets a column showing `['0.00']`."""
    write = _format_writer(_tool(tmp_path))
    for bad in (["0.00"], {"format": "0.00"}, 0.25):
        with pytest.raises(ValueError) as exc:
            write(bad)
        assert "styles[0] format must be a string" in str(exc.value), bad
        assert type(bad).__name__ in str(exc.value), bad
    assert not any(tmp_path.iterdir())

    write("0.00")
    assert _load(tmp_path / "nf.xlsx")["Sheet1"]["A2"].number_format == "0.00"


def test_an_over_long_cell_is_truncated_and_the_receipt_says_so(tmp_path):
    """openpyxl cuts at 32,767 characters without a word. A silently shortened cell is a
    claim the agent would go on to make on our behalf, so it goes in Adjusted."""
    tool = _tool(tmp_path)
    out = tool(
        path="长.xlsx",
        sheets=[
            {"rows": [["标题"], ["很长" * 20_000], ["也很长" * 11_000]], "header_rows": 0}
        ],
    )
    assert "Adjusted: 2 cell(s) truncated to 32,767 characters." in out
    sheet = _load(tmp_path / "长.xlsx")["Sheet1"]
    assert len(sheet["A2"].value) == 32_767
    assert len(sheet["A3"].value) == 32_767

    # exactly at the ceiling: nothing is cut and nothing is reported
    out = tool(path="刚好.xlsx", sheets=[{"rows": [["x" * 32_767]], "header_rows": 0}])
    assert "truncated" not in out
    assert len(_load(tmp_path / "刚好.xlsx")["Sheet1"]["A1"].value) == 32_767


def test_a_nested_array_in_a_cell_is_reported_not_stringified(tmp_path):
    with pytest.raises(ValueError, match="a cell cannot hold list"):
        _tool(tmp_path)(path="x.xlsx", sheets=[{"rows": [[["a", "b"]]]}])


# -- merges, styles, names, limits ----------------------------------------------


def test_merges_are_validated_and_overlaps_named(tmp_path):
    tool = _tool(tmp_path)
    rows = [["a", "b", "c"], ["1", "2", "3"]]

    with pytest.raises(ValueError) as exc:
        tool(path="x.xlsx", sheets=[{"rows": rows, "merges": ["A1:C1", "B1:B2"]}])
    assert '"A1:C1"' in str(exc.value) and '"B1:B2"' in str(exc.value)
    assert "overlap" in str(exc.value)

    with pytest.raises(ValueError, match="use A1 notation"):
        tool(path="x.xlsx", sheets=[{"rows": rows, "merges": ["A1-C1"]}])
    with pytest.raises(ValueError, match="use A1 notation"):
        tool(path="x.xlsx", sheets=[{"rows": rows, "merges": ["1:3"]}])
    with pytest.raises(ValueError, match="outside the sheet"):
        tool(path="x.xlsx", sheets=[{"rows": rows, "merges": ["A1:ZZZ9999999"]}])
    assert not any(tmp_path.iterdir())  # nothing half-written along the way

    # a one-cell "merge" is a no-op, and a reversed range is normalised
    tool(path="ok.xlsx", sheets=[{"rows": rows, "merges": ["B2", "C1:A1"]}])
    assert {str(r) for r in _load(tmp_path / "ok.xlsx")["Sheet1"].merged_cells.ranges} == {
        "A1:C1"
    }


def test_styles_apply_in_order_on_the_merge_anchor(tmp_path):
    tool = _tool(tmp_path)
    tool(
        path="s.xlsx",
        sheets=[
            {
                "rows": [["标题"], ["项目", "金额"], ["项目A", "1200"]],
                "merges": ["A1:B1"],
                "styles": [
                    {"range": "A1:B1", "fill": "#4472C4", "color": "FFFFFF", "height": 30},
                    {"range": "B3", "format": "#,##0.00", "align": "right", "bold": True},
                    {"range": "A1", "size": 20},  # later entry wins
                ],
            }
        ],
    )
    sheet = _load(tmp_path / "s.xlsx")["Sheet1"]
    assert sheet["A1"].fill.start_color.rgb == "FF4472C4"
    assert sheet["A1"].font.color.rgb == "FFFFFFFF"
    assert sheet["A1"].font.sz == 20 and sheet["A1"].font.b is True
    assert sheet.row_dimensions[1].height == 30
    assert sheet["B3"].number_format == "#,##0.00"
    assert sheet["B3"].alignment.horizontal == "right" and sheet["B3"].font.b is True
    # the non-anchor half of the merge never carries the font (openpyxl drops it silently)
    assert sheet["B1"].value is None


def test_style_arguments_are_checked_with_a_worked_example(tmp_path):
    tool = _tool(tmp_path)
    rows = [["a"]]
    for style, message in (
        ({"range": "A1", "fill": "blue"}, "hex colour"),
        ({"range": "A1", "color": "#12345"}, "hex colour"),
        ({"range": "A1", "align": "middle"}, "left, center or right"),
        ({"range": "A1", "size": "big"}, "must be a number"),
        ({"range": "A1", "height": 0}, "positive number"),
        ({"range": "A1", "bold": "maybe"}, "must be true or false"),
        ({"bold": True}, "must be an object with a `range`"),
        ({"range": "A1:XFD1048576"}, "over the"),
    ):
        with pytest.raises(ValueError) as exc:
            tool(path="x.xlsx", sheets=[{"rows": rows, "styles": [style]}])
        assert message in str(exc.value), style
    assert not any(tmp_path.iterdir())


def test_sheet_names_are_repaired_deduplicated_and_reported(tmp_path):
    long_name = "部" * 40
    repaired = "汇总_表" + "_" * 6 + "1"  # [ ] : * ? / \ all become _
    out = _tool(tmp_path)(
        path="n.xlsx",
        sheets=[
            {"name": "汇总[表]:*?/\\1", "rows": [["a"]]},
            {"name": repaired, "rows": [["a"]]},  # collides after repair
            {"name": long_name, "rows": [["a"]]},
            {"rows": [["a"]]},  # no name at all
            {"name": "'引号'", "rows": [["a"]]},
            {"name": "Data", "rows": [["a"]]},
            {"name": "data", "rows": [["a"]]},  # Excel's duplicate check ignores case
            {"name": "History", "rows": [["a"]]},  # reserved by Excel, case regardless
        ],
    )
    names = _load(tmp_path / "n.xlsx").sheetnames
    assert names[0] == repaired
    assert names[1] == f"{repaired} (2)"
    assert names[5] == "Data" and names[6] == "data (2)"
    assert names[7] == "History (2)"
    assert 'sheet name "History" became "History (2)"' in out
    assert names[2] == long_name[:31] and len(names[2]) == 31
    assert names[3] == "Sheet4"
    assert names[4] == "引号"
    assert all(len(n) <= 31 for n in names)
    assert "Adjusted:" in out and f'became "{repaired} (2)"' in out


def test_structural_errors_point_at_the_shape_that_works(tmp_path):
    tool = _tool(tmp_path)
    for sheets, message in (
        ([], "non-empty array of sheet objects"),
        ({}, "has no `rows`"),  # a lone object is unwrapped, then missing its rows
        ("汇总", "non-empty array of sheet objects"),
        ([{"name": "x"}], "has no `rows`"),
        (["rows"], "must be an object with a `rows` array"),
        ([{"rows": "a,b"}], "rows must be an array of arrays"),
        ([{"rows": ["a", "b"]}], "rows[0] is str, not an array"),
        ([{"rows": [["a"]]}] * 51, "at most 50 sheets"),
        ([{"rows": [["a"]], "merges": 3}], "merges must be an array"),
        ([{"rows": [["a"]], "col_widths": 3}], "col_widths must be an array"),
        ([{"rows": [["a"]], "styles": 3}], "styles must be an array"),
        ([{"rows": [["a"]], "header_rows": "两行"}], "header_rows must be a whole number"),
    ):
        with pytest.raises(ValueError) as exc:
            tool(path="x.xlsx", sheets=sheets)
        assert message in str(exc.value), sheets
    assert not any(tmp_path.iterdir())


def test_a_single_sheet_object_is_accepted_unwrapped(tmp_path):
    """Models hand over `sheets={...}` often enough that refusing costs a turn for nothing."""
    _tool(tmp_path)(path="one.xlsx", sheets={"rows": [["a", "b"]]})
    assert _load(tmp_path / "one.xlsx")["Sheet1"]["B1"].value == "b"


def test_too_many_cells_is_refused_with_the_csv_route(tmp_path):
    rows = [["x"] * 100 for _ in range(1000)]  # 100k cells per sheet
    with pytest.raises(ValueError) as exc:
        _tool(tmp_path)(path="big.xlsx", sheets=[{"rows": rows} for _ in range(6)])
    assert "exceeds the 500000 limit" in str(exc.value)
    assert ".csv" in str(exc.value)
    assert not any(tmp_path.iterdir())


def test_header_rows_zero_and_explicit_widths(tmp_path):
    _tool(tmp_path)(
        path="w.xlsx",
        sheets=[
            {
                "rows": [["项目名称", "金额"], ["项目A", "1"]],
                "header_rows": 0,
                "col_widths": [30],
            }
        ],
    )
    sheet = _load(tmp_path / "w.xlsx")["Sheet1"]
    assert sheet.freeze_panes is None  # nothing to freeze without a header
    assert not sheet["A1"].font.b
    assert sheet.column_dimensions["A"].width == 30  # given
    assert sheet.column_dimensions["B"].width == 6  # auto, at the floor


def test_a_title_merged_wider_than_the_data_still_counts_as_a_title(tmp_path):
    """The used area covers the merges: a title across six columns when only three carry
    data is how a report template actually arrives."""
    _tool(tmp_path)(
        path="t.xlsx",
        sheets=[{"rows": [["标题"], ["a", "b", "c"], ["1", "2", "3"]], "merges": ["A1:F1"]}],
    )
    sheet = _load(tmp_path / "t.xlsx")["Sheet1"]
    assert sheet["A1"].font.sz == 14  # treated as the title, not as the header
    assert sheet["A2"].fill.start_color.rgb == "FFDDEBF7"
    assert sheet.freeze_panes == "A3"


# -- paths -----------------------------------------------------------------------


def test_the_extension_is_checked_and_never_appended(tmp_path):
    tool = _tool(tmp_path)
    sheets = [{"rows": [["a"]]}]
    with pytest.raises(ValueError) as exc:
        tool(path="报告.xlsm", sheets=sheets)
    assert "macros" in str(exc.value) and ".xlsx" in str(exc.value)
    with pytest.raises(ValueError) as exc:
        tool(path="报告.xls", sheets=sheets)
    assert "old binary Excel format" in str(exc.value)
    for bad in ("报告", "报告.csv", "报告.XLSX.txt"):
        with pytest.raises(ValueError, match="must end in .xlsx"):
            tool(path=bad, sheets=sheets)
    assert not any(tmp_path.iterdir())  # no auto-renamed file appeared anywhere

    # upper case is fine, and the file lands under the name that was asked for
    tool(path="报告.XLSX", sheets=sheets)
    assert [p.name for p in tmp_path.iterdir()] == ["报告.XLSX"]


def test_chinese_path_and_missing_parent_directories(tmp_path):
    tool = _tool(tmp_path)
    out = tool(path="子目录/第二层/季度 报告 (最终).xlsx", sheets=[dict(_TEMPLATE)])
    target = tmp_path / "子目录" / "第二层" / "季度 报告 (最终).xlsx"
    assert target.is_file() and str(target) in out
    assert _load(target).sheetnames == ["季度汇总"]


def test_writing_over_a_directory_is_refused(tmp_path):
    (tmp_path / "报告.xlsx").mkdir()
    with pytest.raises(ValueError, match="Path is a directory"):
        _tool(tmp_path)(path="报告.xlsx", sheets=[{"rows": [["a"]]}])


# -- path parity with aisuite's write_file --------------------------------------


def _parity_roots(tmp_path):
    ws = tmp_path / "工作区"
    scratch = tmp_path / "scratch"
    ro = tmp_path / "readonly"
    for d in (ws, scratch, ro):
        d.mkdir()
    roots = normalize_roots(
        [
            RootDir(path=ws, writable=True, label="workspace"),
            RootDir(path=scratch, writable=True, label="scratch"),
            RootDir(path=ro, writable=False, label="readonly"),
        ]
    )
    return ws, scratch, ro, roots


def _outcome(call, tmp_path: Path, suffix: str):
    """('error', type, message-with-the-extension-masked) or ('ok', landing directory)."""
    try:
        call()
    except Exception as exc:  # noqa: BLE001 - the comparison is the point
        return ("error", type(exc).__name__, str(exc).replace(suffix, "<ext>"))
    found = list(tmp_path.rglob(f"probe{suffix}"))
    assert len(found) == 1, found
    return ("ok", str(found[0].parent))


_PARITY_CASES = [
    "relative",
    "relative-subdir",
    "absolute-primary",
    "absolute-secondary-root",
    "absolute-outside-every-root",
    "dot-dot-escape",
    "dot-dot-reentry",
    "read-only-root",
    "tilde-outside",
    "tilde-inside",
]


@pytest.mark.parametrize("case", _PARITY_CASES)
def test_path_parity_with_aisuite_write_file(tmp_path, monkeypatch, case):
    """One root set, two tools, the same ten inputs: accept/reject, exception type, message
    and landing directory must match. `write_spreadsheet` re-implements `_resolve` (it
    cannot borrow one that only takes `content: str`), so this is the guard that keeps the
    copy honest."""
    ws, scratch, ro, roots = _parity_roots(tmp_path)
    write_file = next(
        t
        for t in ai.toolkits.files(roots=roots)
        if getattr(t, "__name__", "") == "write_file"
    )
    write_spreadsheet = _tool(ws, roots)

    def path_for(suffix: str) -> str:
        if case == "relative":
            return f"probe{suffix}"
        if case == "relative-subdir":
            return f"新建文件夹/深/probe{suffix}"
        if case == "absolute-primary":
            return str(ws / f"probe{suffix}")
        if case == "absolute-secondary-root":
            return str(scratch / f"probe{suffix}")
        if case == "absolute-outside-every-root":
            return str(tmp_path / f"probe{suffix}")
        if case == "dot-dot-escape":
            return f"../probe{suffix}"
        if case == "dot-dot-reentry":
            # judged AFTER resolution, so a path that walks out and back in is fine
            return str(scratch / "子" / ".." / f"probe{suffix}")
        if case == "read-only-root":
            return str(ro / f"probe{suffix}")
        return f"~/probe{suffix}"  # both tilde cases

    if case == "tilde-inside":
        # `~` must be expanded the same way by both. Point HOME at a root so the case also
        # covers acceptance, not just a shared refusal.
        for var in ("USERPROFILE", "HOME"):
            monkeypatch.setenv(var, str(scratch))
        monkeypatch.delenv("HOMEPATH", raising=False)

    mine = _outcome(
        lambda: write_spreadsheet(path=path_for(".xlsx"), sheets=[{"rows": [["a"]]}]),
        tmp_path,
        ".xlsx",
    )
    theirs = _outcome(
        lambda: write_file(path=path_for(".csv"), content="a\n"), tmp_path, ".csv"
    )
    assert mine == theirs, case


def test_a_symlink_out_of_the_roots_is_judged_by_its_real_path(tmp_path):
    """A junction or symlink inside a root pointing elsewhere must not become a way out —
    both tools resolve first and check the real path, so both refuse."""
    ws = tmp_path / "ws"
    outside = tmp_path / "outside"
    for d in (ws, outside):
        d.mkdir()
    link = ws / "链接"
    try:
        link.symlink_to(outside, target_is_directory=True)
    except (OSError, NotImplementedError) as exc:  # Windows without developer mode
        pytest.skip(f"symlinks unavailable here: {exc}")

    roots = normalize_roots([RootDir(path=ws, writable=True, label="workspace")])
    write_file = next(
        t
        for t in ai.toolkits.files(roots=roots)
        if getattr(t, "__name__", "") == "write_file"
    )
    with pytest.raises(PermissionError, match="escapes allowed roots"):
        _tool(ws, roots)(path=str(link / "probe.xlsx"), sheets=[{"rows": [["a"]]}])
    with pytest.raises(PermissionError, match="escapes allowed roots"):
        write_file(path=str(link / "probe.csv"), content="a\n")
    assert list(outside.iterdir()) == []


def test_no_writable_root_is_refused_like_the_toolkit(tmp_path):
    ro = tmp_path / "ro"
    ro.mkdir()
    roots = normalize_roots([RootDir(path=ro, writable=False, label="ro")])
    with pytest.raises(PermissionError, match="no writable directory"):
        _tool(ro, roots)(path="x.xlsx", sheets=[{"rows": [["a"]]}])


def test_a_folder_granted_mid_session_is_reachable_without_rebuilding(tmp_path):
    """`roots` is held by reference, like every other roots consumer — a folder the user
    grants mid-turn must work on that same turn (the 2026-08-31 grant loop)."""
    ws = tmp_path / "ws"
    granted = tmp_path / "granted"
    for d in (ws, granted):
        d.mkdir()
    roots = normalize_roots([RootDir(path=ws, writable=True, label="workspace")])
    tool = _tool(ws, roots)
    target = str(granted / "报告.xlsx")

    with pytest.raises(PermissionError, match="escapes allowed roots"):
        tool(path=target, sheets=[{"rows": [["a"]]}])
    roots.append(RootDir(path=granted, writable=True, label="granted"))
    assert "(root: granted)" in tool(path=target, sheets=[{"rows": [["a"]]}])


def test_the_deepest_root_names_the_file_like_write_file_does(tmp_path):
    outer = tmp_path / "outer"
    inner = outer / "inner"
    inner.mkdir(parents=True)
    roots = normalize_roots(
        [
            RootDir(path=outer, writable=True, label="outer"),
            RootDir(path=inner, writable=True, label="inner"),
        ]
    )
    out = _tool(outer, roots)(
        path=str(inner / "报告.xlsx"), sheets=[{"rows": [["a"]]}]
    )
    assert "(root: inner)" in out


# -- atomic write ----------------------------------------------------------------


def _boom(*_a, **_k):
    raise RuntimeError("save blew up")


def test_a_failed_save_leaves_neither_a_target_nor_a_temp_file(tmp_path, monkeypatch):
    monkeypatch.setattr(office_module, "_save_workbook", _boom)
    with pytest.raises(RuntimeError, match="save blew up"):
        _tool(tmp_path)(path="报告.xlsx", sheets=[{"rows": [["a"]]}])
    assert list(tmp_path.iterdir()) == []  # no half .xlsx, no .part left behind


def test_a_failed_save_keeps_the_previous_file_intact(tmp_path, monkeypatch):
    tool = _tool(tmp_path)
    tool(path="报告.xlsx", sheets=[{"rows": [["旧数据"]]}])
    target = tmp_path / "报告.xlsx"
    before = target.read_bytes()

    monkeypatch.setattr(office_module, "_save_workbook", _boom)
    with pytest.raises(RuntimeError):
        tool(path="报告.xlsx", sheets=[{"rows": [["新数据"]]}])

    assert target.read_bytes() == before
    assert _load(target)["Sheet1"]["A1"].value == "旧数据"
    assert [p.name for p in tmp_path.iterdir()] == ["报告.xlsx"]


def test_a_locked_target_says_the_file_is_open_elsewhere(tmp_path, monkeypatch):
    def locked(_src, _dst):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(office_module, "_replace", locked)
    with pytest.raises(PermissionError) as exc:
        _tool(tmp_path)(path="报告.xlsx", sheets=[{"rows": [["a"]]}])
    assert "is open in another program (Excel?)" in str(exc.value)
    assert str(tmp_path / "报告.xlsx") in str(exc.value)
    assert list(tmp_path.iterdir()) == []  # the temp file is cleaned up regardless


def test_the_refusal_distinguishes_read_only_from_a_file_someone_has_open(
    tmp_path, monkeypatch
):
    """"Close Excel" sends the user hunting for a window that does not exist when the real
    cause is a read-only file, a write-protected folder or Defender's controlled folder
    access — all of which arrive as the same PermissionError."""
    target = tmp_path / "报告.xlsx"
    target.write_bytes(b"PK\x03\x04")

    monkeypatch.setattr(office_module.os, "access", lambda *_a, **_k: False)
    read_only = office_module._replace_denied(target)
    assert "is read-only / write-protected" in read_only
    assert "open in another program" not in read_only
    assert "Clear the read-only flag" in read_only

    monkeypatch.setattr(office_module.os, "access", lambda *_a, **_k: True)
    locked = office_module._replace_denied(target)
    assert "is open in another program (Excel?)" in locked
    # a target that does not exist yet cannot be the read-only one, so both are named
    monkeypatch.setattr(office_module.os, "access", lambda *_a, **_k: False)
    absent = office_module._replace_denied(tmp_path / "缺失.xlsx")
    assert "is open in another program (Excel?)" in absent
    assert "read-only / write-protected" in absent


@pytest.mark.skipif(
    os.name != "nt", reason="only Windows refuses to replace a read-only file"
)
def test_a_read_only_target_is_not_blamed_on_excel(tmp_path):
    """End to end, with a genuinely read-only file and the real `os.replace`."""
    tool = _tool(tmp_path)
    tool(path="报告.xlsx", sheets=[{"rows": [["旧数据"]]}])
    target = tmp_path / "报告.xlsx"
    target.chmod(0o444)
    try:
        if os.access(target, os.W_OK):  # pragma: no cover - filesystem without the bit
            pytest.skip("this filesystem ignores the read-only bit")
        with pytest.raises(PermissionError) as exc:
            tool(path="报告.xlsx", sheets=[{"rows": [["新数据"]]}])
        message = str(exc.value)
        assert "is read-only / write-protected" in message
        assert "open in another program" not in message
        assert [p.name for p in tmp_path.iterdir()] == ["报告.xlsx"]  # no .part left
        assert _load(target)["Sheet1"]["A1"].value == "旧数据"  # and it is untouched
    finally:
        target.chmod(0o666)


def test_overwriting_an_existing_file_is_the_documented_behaviour(tmp_path):
    tool = _tool(tmp_path)
    tool(path="报告.xlsx", sheets=[{"rows": [["旧"]]}])
    tool(path="报告.xlsx", sheets=[{"rows": [["新"]]}])
    assert _load(tmp_path / "报告.xlsx")["Sheet1"]["A1"].value == "新"


# -- schema / import cost --------------------------------------------------------


@pytest.mark.parametrize(
    "schema,name,required",
    [
        (_SCHEMA, "write_spreadsheet", ["path", "sheets"]),
        (_DOC_SCHEMA, "write_document", ["path", "markdown"]),
    ],
    ids=["write_spreadsheet", "write_document"],
)
def test_schema_sticks_to_the_subset_every_provider_accepts(schema, name, required):
    """anyOf, type arrays, null, additionalProperties and default each get a request
    rejected by at least one hosted provider. Keep the schema boring.

    Both office tools ride in every session's prompt, so both are held to it."""
    banned = {"anyOf", "oneOf", "allOf", "not", "default", "additionalProperties", "$ref"}
    allowed_types = {
        "object",
        "array",
        "string",
        "integer",
        "number",
        "boolean",
    }

    def walk(node, where="root"):
        if isinstance(node, dict):
            for key, value in node.items():
                assert key not in banned, f"{where}.{key}"
                if key == "type":
                    assert isinstance(value, str), f"{where}.type is a list"
                    assert value in allowed_types, f"{where}.type = {value}"
                walk(value, f"{where}.{key}")
        elif isinstance(node, list):
            for i, value in enumerate(node):
                walk(value, f"{where}[{i}]")

    fn = schema["function"]
    walk(fn["parameters"], "parameters")  # the envelope's own "type": "function" is fine
    assert schema["type"] == "function"
    assert fn["name"] == name
    assert fn["parameters"]["required"] == required
    # The tool sits in every session's prompt; a runaway description is a per-turn bill.
    assert len(json.dumps(fn)) < 3000


def test_the_tool_advertises_the_schema_and_its_write_metadata(tmp_path):
    tool = _tool(tmp_path)
    assert tool.__coworker_schema__ is _SCHEMA
    meta = tool.__aisuite_tool_metadata__
    assert meta.risk_level == "medium" and meta.requires_approval is True


def test_a_build_without_openpyxl_says_what_to_deliver_instead(tmp_path, monkeypatch):
    """The dependency is bundled, but a trimmed build must fail with an instruction, not a
    traceback: the model needs to hear "deliver a .csv" to salvage the turn."""
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    with pytest.raises(RuntimeError) as exc:
        _tool(tmp_path)(path="报告.xlsx", sheets=[{"rows": [["a"]]}])
    assert "missing openpyxl" in str(exc.value)
    assert ".csv" in str(exc.value) and "U+FEFF" in str(exc.value)
    assert list(tmp_path.iterdir()) == []


def test_importing_the_module_does_not_import_openpyxl():
    """Lazy by design: a session that never writes a spreadsheet must not pay the import,
    and a build without openpyxl must still start."""
    code = (
        "import sys; import coworker.tools.office as m; "
        "print('openpyxl' in sys.modules, bool(m.office_tools))"
    )
    proc = subprocess.run(
        [sys.executable, "-c", code],
        cwd=str(Path(__file__).resolve().parents[1]),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    assert proc.returncode == 0, proc.stderr
    assert proc.stdout.strip() == "False True"
