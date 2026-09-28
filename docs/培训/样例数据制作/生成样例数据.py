#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""培训样例数据生成脚本。

生成 docs/培训/OpenWorker培训/ 下 01-质量周报 / 02-交期风险 / 03-设备点检 /
04-备件盘点 / 05-返厂维修 / 07-租赁台账 六个文件夹里的 Excel 和 Word，以及
docs/培训/样例数据说明.md。06-返厂检验 是图片，由同目录的 渲染检验记录.cjs 生成。
用法见同目录的 README.md。

设计要点：
- 所有日期都写成相对 --today（T）的偏移（日历日 T+n 或工作日 wd(-k)），05 用 2026
  年固定日期。这样换一个 --today，行数/顺序/非日期内容/统计结果都不变。
- 「点名」的行（埋的问题、边界情况）逐条写死在常量里；其余填充内容（型号、库位、
  故障描述文字等）用各数据集专属的 random.Random(种子) 生成，种子固定，结果可重复。
- 不写公式：应用的读取工具用 data_only=True 打开，openpyxl 写的公式没有缓存值。
- 写完文件后重新用 openpyxl.load_workbook(data_only=True) 读回来做断言（verify_*），
  断言不通过就非零退出、不生成说明文档；样例数据说明.md 里的数字和行号也都来自
  重新读文件（doc_*），不是生成时顺手记的。
"""

from __future__ import annotations

import argparse
import random
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Optional

sys.stdout.reconfigure(encoding="utf-8")

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


# ============================================================================
# 常量
# ============================================================================

WB_TITLE = "示例数据（培训用，非真实业务数据）"
WB_CREATOR = "OpenWorker 培训"

RESERVED_TOOL_IDS = {"YZJ-0417", "YZJ-0432", "YZJ-0502"}  # 保留给 06 文件夹
RESERVED_NUMBERS = {417, 432, 502}  # 数字本身也避开，不论前缀，减少混淆

TYPE_PREFIX = {
    "液压震击器": "YZJ",
    "机械震击器": "JZJ",
    "减震器": "JZQ",
    "螺杆钻具": "LG",
    "加速器": "JSQ",
}

SPEC_POOL = ["4-3/4", "6-1/4", "6-1/2", "6-3/4", "8"]

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
HEADER_FONT = Font(bold=True)
DATE_FMT = "yyyy-mm-dd"

# ============================================================================
# 日期 / 工作日 助手
# ============================================================================


def wd(t: date, k: int) -> date:
    """严格早于 t 的第 k 个工作日（周一至周五，不考虑节假日），k >= 1。"""
    if k < 1:
        raise ValueError("wd(k) 要求 k >= 1")
    d = t - timedelta(days=1)
    n = 0
    while True:
        if d.weekday() < 5:
            n += 1
            if n == k:
                return d
        d -= timedelta(days=1)


def cal(t: date, n: int) -> date:
    """T+n 日历日偏移（n 可为负）。"""
    return t + timedelta(days=n)


def friday_on_or_before(t: date) -> date:
    d = t
    while d.weekday() != 4:  # 周五
        d -= timedelta(days=1)
    return d


def iso_week(d: date) -> int:
    return d.isocalendar()[1]


# ============================================================================
# 宽度 / Excel 写入助手
# ============================================================================


def _disp_width(s: str) -> int:
    w = 0
    for ch in s:
        w += 2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1
    return w


def _cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


def style_workbook(wb: Workbook) -> None:
    wb.properties.title = WB_TITLE
    wb.properties.creator = WB_CREATOR


def write_table(
    ws: Worksheet,
    headers: list[str],
    rows: list[list[Any]],
    date_cols: tuple[int, ...] = (),
    header_row: int = 1,
    freeze: Optional[str] = None,
) -> None:
    """写表头（加粗+浅灰底）+ 数据行，日期列写 date 对象并设置格式，自动列宽，冻结表头。"""
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=j, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    for i, row in enumerate(rows, start=header_row + 1):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v)
            if (j - 1) in date_cols and isinstance(v, date):
                c.number_format = DATE_FMT
    widths = [_disp_width(h) for h in headers]
    for row in rows:
        for j, v in enumerate(row):
            widths[j] = max(widths[j], _disp_width(_cell_text(v)))
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w + 2
    ws.freeze_panes = freeze if freeze else f"A{header_row + 1}"


def new_workbook_with_sheet(title: str) -> tuple[Workbook, Worksheet]:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    style_workbook(wb)
    return wb, ws


# ============================================================================
# 清理输出目录
# ============================================================================

OWNED_SUBDIRS = [
    "01-质量周报",
    "02-交期风险",
    "03-设备点检",
    "04-备件盘点",
    "05-返厂维修",
    "07-租赁台账",
]


def clean_outputs(out_root: Path) -> None:
    for name in OWNED_SUBDIRS:
        d = out_root / name
        d.mkdir(parents=True, exist_ok=True)
        for f in list(d.glob("*.xlsx")) + list(d.glob("*.docx")):
            f.unlink()


# ============================================================================
# 01 · 质量周报
# ============================================================================

INSPECTORS = ["Q-011", "Q-012", "Q-017"]

# 零件决定「产品」，固定不变（机加/热处理两份检验记录共用）
PART_TO_PRODUCT = {
    "芯轴": "液压震击器", "活塞": "液压震击器", "外筒": "液压震击器", "上接头": "液压震击器",
    "下接头": "机械震击器", "花键套": "减震器", "转子": "螺杆钻具", "壳体": "螺杆钻具",
}

# (day 1-5, 设备, 工序, 零件, 送检, 不良, 不良类型, 处置)
JIAJIA_NAMED = [
    (1, "CK-03", "螺纹加工", "上接头", 20, 3, "螺纹中径超差", "返修"),
    (2, "CK-03", "螺纹加工", "下接头", 20, 2, "螺纹中径超差", "返修"),
    (3, "CK-03", "螺纹加工", "上接头", 20, 0, "", ""),
    (4, "CK-03", "螺纹加工", "上接头", 20, 3, "螺纹中径超差", "返修"),
    (5, "CK-03", "螺纹加工", "下接头", 20, 4, "螺纹中径超差", "返修"),
    (2, "CK-01", "精车", "芯轴", 30, 1, "外圆尺寸超差", "返修"),
    (3, "SK-01", "深孔钻", "外筒", 12, 1, "深孔直线度超差", "报废"),
    (4, "CK-05", "精车", "活塞", 25, 2, "表面粗糙度不合格", "返修"),
]

JIAJIA_FILLER_PROCESSES = ["粗车", "精车", "深孔钻", "铣花键"]
JIAJIA_FILLER_DEVICES = {
    "粗车": ["CK-01", "CK-02", "CK-04", "CK-05"],
    "精车": ["CK-01", "CK-02", "CK-04", "CK-05"],
    "深孔钻": ["SK-01", "SK-02"],
    "铣花键": ["XK-01"],
}
JIAJIA_FILLER_PARTS = {
    "粗车": ["芯轴", "外筒", "上接头", "下接头", "活塞", "转子", "壳体"],
    "精车": ["芯轴", "外筒", "上接头", "下接头", "活塞", "转子", "壳体"],
    "深孔钻": ["外筒", "芯轴"],
    "铣花键": ["花键套"],
}

RECHULI_NAMED = [
    # day, 设备, 零件, 送检, 不良, 不良类型, 处置, 工序
    (2, "RC-01", "芯轴", 24, 1, "硬度不合格", "返修", "回火"),
    (3, "RC-02", "花键套", 20, 1, "硬度不合格", "返修", "渗氮"),
    (4, "RC-02", "活塞", 22, 2, "硬度不合格", "返修", "调质"),
    (5, "RC-02", "芯轴", 24, 4, "硬度不合格", "报废", "调质"),
]
RECHULI_FILLER_DAYS_DEVICES = [
    (1, "RC-01"), (1, "RC-02"), (2, "RC-02"),
    (3, "RC-01"), (4, "RC-01"), (5, "RC-01"),
]
RECHULI_PROCS = ["调质", "回火", "渗氮"]
RECHULI_PARTS = ["芯轴", "活塞", "花键套", "外筒"]

ZHUANGPEI_NAMED = [
    # day, 设备, 工序, 送检, 不良, 不良类型, 处置
    (2, "ZP-01", "总装", 12, 1, "密封件装配划伤", "返修"),
    (3, "SY-01", "试压", 10, 1, "试压渗漏", "返修"),
    (5, "ZP-01", "总装", 12, 1, "密封件装配划伤", "返修"),
]
ZHUANGPEI_DEVICE_PROC = {"ZP-01": "总装", "SY-01": "试压", "ZJ-01": "震击测试"}
ZHUANGPEI_PART = "液压震击器总成"

NCR_ROWS = [
    # 序, day, 来源车间, 零件, 设备, 不良描述, 数量, 处置, 责任工序, 原因分析, 纠正措施, 状态
    (1, 1, "机加车间", "上接头", "CK-03", "螺纹中径超差", 3, "返修", "螺纹加工",
     "刀片磨损未及时更换", "更换刀片，恢复首件必检", "已关闭"),
    (2, 2, "机加车间", "下接头", "CK-03", "螺纹中径超差", 2, "返修", "螺纹加工",
     "换刀后刀补未更新", "重新对刀，修正刀补", "已关闭"),
    (3, 2, "热处理车间", "芯轴", "RC-01", "硬度不合格", 1, "返修", "回火",
     "回火温度设定偏低", "按工艺卡重新回火", "已关闭"),
    (4, 3, "机加车间", "外筒", "SK-01", "深孔直线度超差", 1, "报废", "深孔钻",
     "钻杆支撑松动", "紧固支撑，加做首件检查", "处理中"),
    (5, 4, "机加车间", "上接头", "CK-03", "螺纹中径超差", 3, "返修", "螺纹加工",
     "换刀后仍然复发，原因待查", "待定", "处理中"),
    (6, 4, "热处理车间", "活塞", "RC-02", "硬度不合格", 2, "返修", "调质",
     "炉温均匀性待核查", "待定", "处理中"),
    (7, 5, "机加车间", "下接头", "CK-03", "螺纹中径超差", 4, "返修", "螺纹加工",
     "疑似主轴或卡盘跳动", "待设备动力组检查", "处理中"),
    (8, 5, "热处理车间", "芯轴", "RC-02", "硬度不合格", 4, "报废", "调质",
     "RC-02 温控仪表疑似漂移", "待校准", "处理中"),
]


def build_01(out_root: Path, T: date, summary: list) -> None:
    d = out_root / "01-质量周报"
    F = friday_on_or_before(T)
    wk = iso_week(F)
    days = {i: cal(F, i - 5) for i in (1, 2, 3, 4, 5)}  # D1=F-4 ... D5=F

    rng_j = random.Random(20260101)
    rng_r = random.Random(20260102)
    rng_z = random.Random(20260103)
    rng_ins = random.Random(20260104)

    # ---- 机加车间 ----
    named_by_day: dict[int, list] = {i: [] for i in range(1, 6)}
    for row in JIAJIA_NAMED:
        named_by_day[row[0]].append(row)
    jiajia_rows_raw = []
    for day_i in range(1, 6):
        rows_today = list(named_by_day[day_i])
        need = 5 - len(rows_today)
        for _ in range(need):
            proc = rng_j.choice(JIAJIA_FILLER_PROCESSES)
            device = rng_j.choice(JIAJIA_FILLER_DEVICES[proc])
            part = rng_j.choice(JIAJIA_FILLER_PARTS[proc])
            sent = rng_j.randint(10, 40)
            rows_today.append((day_i, device, proc, part, sent, 0, "", ""))
        jiajia_rows_raw.extend(rows_today)

    batch_counter = 0
    jiajia_final = []
    for r in jiajia_rows_raw:
        day_i, device, proc, part, sent, bad, bad_type, disp = r
        batch_counter += 1
        batch_no = f"P{wk:02d}J{batch_counter:03d}"
        qualified = sent - bad
        inspector = rng_ins.choice(INSPECTORS)
        jiajia_final.append([
            days[day_i], batch_no, PART_TO_PRODUCT[part], part, proc, device, sent,
            qualified, bad, bad_type, disp, inspector,
        ])

    ws_j_headers = ["日期", "批次号", "产品", "零件名称", "工序", "设备编号",
                     "送检数", "合格数", "不良数", "不良类型", "处置", "检验员"]
    wb_j, ws_j = new_workbook_with_sheet("检验记录")
    write_table(ws_j, ws_j_headers, jiajia_final, date_cols=(0,))
    p = d / "检验记录-机加车间.xlsx"
    wb_j.save(p)
    summary.append((p, len(jiajia_final)))

    # ---- 热处理车间 ----
    named_r_by_day: dict[int, list] = {i: [] for i in range(1, 6)}
    for row in RECHULI_NAMED:
        named_r_by_day[row[0]].append(row)
    rechuli_rows = []
    filler_iter = iter(RECHULI_FILLER_DAYS_DEVICES)
    for day_i in range(1, 6):
        entries = []
        for row in named_r_by_day[day_i]:
            entries.append(row)
        rechuli_rows.append((day_i, entries))
    # 补齐 filler：每天缺的设备
    for day_i, device in RECHULI_FILLER_DAYS_DEVICES:
        part = rng_r.choice(RECHULI_PARTS)
        proc = rng_r.choice(RECHULI_PROCS)
        sent = rng_r.randint(15, 28)
        for entry in rechuli_rows:
            if entry[0] == day_i:
                entry[1].append((day_i, device, part, sent, 0, "", "", proc))

    rechuli_final = []
    for day_i, entries in rechuli_rows:
        # 每天固定 RC-01 先、RC-02 后
        entries_sorted = sorted(entries, key=lambda e: e[1])
        for e in entries_sorted:
            _, device, part, sent, bad, bad_type, disp, proc = e
            qualified = sent - bad
            inspector = rng_ins.choice(INSPECTORS)
            rechuli_final.append([
                days[day_i], f"P{wk:02d}R{len(rechuli_final)+1:03d}", PART_TO_PRODUCT[part],
                part, proc, device, sent, qualified, bad, bad_type, disp, inspector,
            ])

    wb_r, ws_r = new_workbook_with_sheet("检验记录")
    write_table(ws_r, ws_j_headers, rechuli_final, date_cols=(0,))
    p = d / "检验记录-热处理车间.xlsx"
    wb_r.save(p)
    summary.append((p, len(rechuli_final)))

    # ---- 装配车间 ----
    named_z_by_day: dict[int, list] = {i: [] for i in range(1, 6)}
    for row in ZHUANGPEI_NAMED:
        named_z_by_day[row[0]].append(row)
    zhuangpei_final = []
    for day_i in range(1, 6):
        devices_today = {"ZP-01", "SY-01", "ZJ-01"}
        named_devices_today = {r[1] for r in named_z_by_day[day_i]}
        rows_today = list(named_z_by_day[day_i])
        for device in sorted(devices_today - named_devices_today):
            proc = ZHUANGPEI_DEVICE_PROC[device]
            sent = rng_z.randint(8, 15)
            rows_today.append((day_i, device, proc, sent, 0, "", ""))
        rows_today.sort(key=lambda r: r[1])  # 设备编号升序，固定顺序
        for r in rows_today:
            _, device, proc, sent, bad, bad_type, disp = r
            qualified = sent - bad
            inspector = rng_ins.choice(INSPECTORS)
            zhuangpei_final.append([
                days[day_i], f"P{wk:02d}Z{len(zhuangpei_final)+1:03d}", "液压震击器",
                ZHUANGPEI_PART, proc, device, sent, qualified, bad, bad_type, disp,
                inspector,
            ])

    wb_z, ws_zh = new_workbook_with_sheet("检验记录")
    write_table(ws_zh, ws_j_headers, zhuangpei_final, date_cols=(0,))
    p = d / "检验记录-装配车间.xlsx"
    wb_z.save(p)
    summary.append((p, len(zhuangpei_final)))

    # ---- 不良品处理单 ----
    ncr_headers = ["单号", "日期", "来源车间", "零件名称", "设备编号", "不良描述",
                   "数量", "处置方式", "责任工序", "原因分析", "纠正措施", "状态"]
    ncr_rows = []
    for row in NCR_ROWS:
        (seq, day_i, ws_name, part, device, desc, qty, disp, proc, cause, action,
         status) = row
        ncr_rows.append([
            f"NCR-{wk:02d}-{seq:02d}", days[day_i], ws_name, part, device, desc,
            qty, disp, proc, cause, action, status,
        ])
    wb_n, ws_n = new_workbook_with_sheet("不良品处理单")
    write_table(ws_n, ncr_headers, ncr_rows, date_cols=(1,))
    p = d / "不良品处理单.xlsx"
    wb_n.save(p)
    summary.append((p, len(ncr_rows)))

    # ---- 上周质量周报.docx ----
    wk_prev = wk - 1
    f11 = cal(F, -11)
    f7 = cal(F, -7)
    doc = Document()
    title = doc.add_heading(f"质量周报（第 {wk_prev} 周）", level=1)
    sub = doc.add_paragraph(f"统计周期：{f11.isoformat()} 至 {f7.isoformat()}　编制：质量安全部")
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_heading("一、本周概况", level=2)
    t1 = doc.add_table(rows=7, cols=2)
    t1.style = "Table Grid"
    t1.cell(0, 0).text = "项目"
    t1.cell(0, 1).text = "数值"
    overview = [
        ("送检数", "1180"), ("合格数", "1162"), ("不良数", "18"),
        ("合格率", "98.5%"), ("返修", "15"), ("报废", "3"),
    ]
    for i, (k, v) in enumerate(overview, start=1):
        t1.cell(i, 0).text = k
        t1.cell(i, 1).text = v

    doc.add_heading("二、不良分布", level=2)
    t2 = doc.add_table(rows=4, cols=4)
    t2.style = "Table Grid"
    for j, h in enumerate(["车间", "送检数", "不良数", "主要不良类型"]):
        t2.cell(0, j).text = h
    dist = [
        ("机加车间", "640", "11", "螺纹中径超差、外圆尺寸超差"),
        ("热处理车间", "230", "3", "硬度不合格"),
        ("装配车间", "310", "4", "密封件装配划伤、试压渗漏"),
    ]
    for i, row in enumerate(dist, start=1):
        for j, v in enumerate(row):
            t2.cell(i, j).text = v

    doc.add_heading("三、重复出现的问题", level=2)
    doc.add_paragraph("CK-03 螺纹加工出现螺纹中径超差 5 件，已更换刀片，下周跟踪。", style="List Bullet")
    doc.add_paragraph("密封件装配划伤连续两周出现，装配车间已安排专项培训。", style="List Bullet")

    doc.add_heading("四、下周关注点", level=2)
    doc.add_paragraph("跟踪 CK-03 螺纹加工首件检验结果。", style="List Bullet")
    doc.add_paragraph("热处理车间 RC-02 炉温记录抽查。", style="List Bullet")

    section = doc.sections[0]
    footer = section.footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.text = "示例数据 · 培训用"
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in fp.runs:
        run.font.size = Pt(8)

    core = doc.core_properties
    core.title = WB_TITLE
    core.author = WB_CREATOR

    p = d / "上周质量周报.docx"
    doc.save(p)
    summary.append((p, None))


# ============================================================================
# 02 · 交期风险
# ============================================================================

ORDERS = [
    # 订单号, 客户, 产品, 规格, 数量, 下单偏移, 交期偏移, 状态
    ("SO-26-101", "客户A", "液压震击器", "6-1/2", 6, -60, -2, "生产中"),
    ("SO-26-102", "客户B", "减震器", "8", 4, -50, 3, "生产中"),
    ("SO-26-103", "客户C", "液压震击器", "4-3/4", 10, -45, 6, "生产中"),
    ("SO-26-104", "客户A", "螺杆钻具", "6-3/4", 8, -40, 9, "生产中"),
    ("SO-26-105", "客户D", "机械震击器", "6-1/4", 5, -35, 12, "生产中"),
    ("SO-26-106", "客户E", "加速器", "6-1/2", 3, -30, 13, "生产中"),
    ("SO-26-107", "客户B", "液压震击器", "8", 4, -28, 21, "生产中"),
    ("SO-26-108", "客户F", "减震器", "6-1/2", 6, -25, 30, "生产中"),
    ("SO-26-109", "客户C", "螺杆钻具", "4-3/4", 12, -20, 45, "生产中"),
    ("SO-26-110", "客户G", "减震器", "8", 20, -15, 60, "生产中"),
    ("SO-26-111", "客户D", "液压震击器", "6-1/2", 5, -70, -10, "已发货"),
    ("SO-26-112", "客户H", "机械震击器", "4-3/4", 4, -65, -5, "已发货"),
]

STAGES = ["下料", "粗车", "热处理（外协）", "精车与深孔", "螺纹加工", "装配试压"]

PROGRESS = [
    # 订单号, 当前工序, 已完成, 完成率, 最近更新偏移, 异常类型, 异常说明
    ("SO-26-101", "精车与深孔", 3, 55, -1, "设备故障", "深孔钻床 SK-02 主轴异响停机检修，深孔工序积压"),
    ("SO-26-102", "装配试压", 5, 95, -1, "无", ""),
    ("SO-26-103", "粗车", 1, 35, -1, "缺料", "芯轴锻件未到货，供应商答复预计 {T+4} 到厂"),
    ("SO-26-104", "热处理（外协）", 2, 45, -6, "工序拖期", "外协热处理已拖期 5 天，回厂时间未确认"),
    ("SO-26-105", "螺纹加工", 4, 65, -1, "缺料", "密封件套件缺货，采购在途，装配无法开始"),
    ("SO-26-106", "螺纹加工", 4, 75, -1, "无", ""),
    ("SO-26-107", "精车与深孔", 3, 45, -2, "无", ""),
    ("SO-26-108", "热处理（外协）", 2, 30, -2, "无", ""),
    ("SO-26-109", "粗车", 1, 15, -1, "无", ""),
    ("SO-26-110", "下料", 0, 5, -1, "无", ""),
    ("SO-26-111", "已完工", 6, 100, -10, "无", ""),
    ("SO-26-112", "已完工", 6, 100, -5, "无", ""),
]


def build_02(out_root: Path, T: date, summary: list) -> None:
    d = out_root / "02-交期风险"

    order_headers = ["订单号", "客户", "产品", "规格(OD,in)", "数量", "下单日期",
                      "合同交期", "订单状态"]
    order_rows = []
    for o in ORDERS:
        order_id, cust, prod, spec, qty, order_off, due_off, status = o
        order_rows.append([
            order_id, cust, prod, spec, qty, cal(T, order_off), cal(T, due_off), status,
        ])
    wb_o, ws_o = new_workbook_with_sheet("订单")
    write_table(ws_o, order_headers, order_rows, date_cols=(5, 6))
    p = d / "订单清单.xlsx"
    wb_o.save(p)
    summary.append((p, len(order_rows)))

    sched_headers = ["订单号", "工序序号", "工序", "计划开始", "计划完成", "设备/班组"]
    sched_rows = []
    for idx, o in enumerate(ORDERS):
        order_id, cust, prod, spec, qty, order_off, due_off, status = o
        start_off = order_off + 3
        end_off = due_off - 2
        span = end_off - start_off
        cuts = [start_off + round(i * span / 6) for i in range(7)]
        ck_device = f"CK-0{(idx % 5) + 1}"
        sk_device = "SK-02" if order_id == "SO-26-101" else ("SK-01" if idx % 2 == 0 else "SK-02")
        devices = ["下料班", ck_device, "外协", sk_device, "CK-03", "装配班"]
        for i in range(6):
            sched_rows.append([
                order_id, i + 1, STAGES[i], cal(T, cuts[i]), cal(T, cuts[i + 1]),
                devices[i],
            ])
    wb_s, ws_s = new_workbook_with_sheet("排产")
    write_table(ws_s, sched_headers, sched_rows, date_cols=(3, 4))
    p = d / "排产计划.xlsx"
    wb_s.save(p)
    summary.append((p, len(sched_rows)))

    prog_headers = ["订单号", "当前工序", "已完成工序数", "总工序数", "完成率(%)",
                     "最近更新", "异常类型", "异常说明"]
    prog_rows = []
    for row in PROGRESS:
        order_id, cur_proc, done, rate, upd_off, ex_type, ex_note = row
        ex_note = ex_note.replace("{T+4}", cal(T, 4).isoformat())
        prog_rows.append([
            order_id, cur_proc, done, 6, rate, cal(T, upd_off), ex_type, ex_note,
        ])
    wb_p, ws_p = new_workbook_with_sheet("进度")
    write_table(ws_p, prog_headers, prog_rows, date_cols=(5,))
    p = d / "进度跟踪.xlsx"
    wb_p.save(p)
    summary.append((p, len(prog_rows)))


# ============================================================================
# 03 · 设备点检
# ============================================================================

# 设备编号, 设备名称, 型号, 车间, 点检周期, 上次点检(偏移函数标记), 下次应检
EQUIP = [
    ("CK-01", "数控车床", "CK6150", "机加车间", "每日", ("wd", 1), ("t", 0)),
    ("CK-02", "数控车床", "CK6150", "机加车间", "每日", ("wd", 1), ("t", 0)),
    ("CK-03", "数控车床", "CK6150", "机加车间", "每日", ("wd", 1), ("t", 0)),
    ("CK-04", "数控车床", "CK6150", "机加车间", "每日", ("wd", 1), ("t", 0)),
    ("CK-05", "数控车床", "CK6150", "机加车间", "每日", ("wd", 1), ("t", 0)),
    ("SK-01", "深孔钻床", "TBM6212", "机加车间", "每日", ("wd", 1), ("t", 0)),
    ("SK-02", "深孔钻床", "TBM6212", "机加车间", "每日", ("wd", 1), ("t", 0)),
    ("XK-01", "数控铣床", "XK714", "机加车间", "每周", ("cal", -7), ("t", 0)),
    ("RC-01", "热处理炉", "RJJ-60-9", "热处理车间", "每周", ("cal", -7), ("t", 0)),
    ("RC-02", "热处理炉", "RJJ-60-9", "热处理车间", "每周", ("cal", -4), ("cal", 3)),
    ("SY-01", "试压台", "SYT-350", "装配车间", "每周", ("cal", -2), ("cal", 5)),
    ("ZJ-01", "震击测试台", "ZJT-100", "装配车间", "每周", ("cal", -10), ("cal", -3)),
    ("KY-01", "空压机", "LGY-22/8", "动力站", "每月", ("cal", -30), ("t", 0)),
    ("XC-01", "行车", "LD5t", "装配车间", "每月", ("cal", -12), ("cal", 18)),
]

DAILY_DEVICES = ["CK-01", "CK-02", "CK-03", "CK-04", "CK-05", "SK-01", "SK-02"]
WEEKLY_MONTHLY_DEVICES = ["XK-01", "RC-01", "RC-02", "SY-01", "ZJ-01", "KY-01", "XC-01"]

INSPECT_ITEMS = {
    "CK-01": "主轴、导轨、润滑、液压站、冷却",
    "CK-02": "主轴、导轨、润滑、液压站、冷却",
    "CK-03": "主轴、导轨、润滑、液压站、冷却",
    "CK-04": "主轴、导轨、润滑、液压站、冷却",
    "CK-05": "主轴、导轨、润滑、液压站、冷却",
    "SK-01": "主轴、给进、冷却、排屑、液压站",
    "SK-02": "主轴、给进、冷却、排屑、液压站",
    "XK-01": "主轴、导轨、分度头、冷却、气路",
    "RC-01": "炉温、控温仪表、密封、传动、安全联锁",
    "RC-02": "炉温、控温仪表、密封、传动、安全联锁",
    "SY-01": "试压泵、压力表、管路密封、安全阀",
    "ZJ-01": "震击台面、传感器、液压系统、紧固螺栓",
    "KY-01": "润滑油位、皮带张紧、安全阀、排水",
    "XC-01": "钢丝绳、限位开关、制动器、吊钩",
}

DEVICE_NAME = {e[0]: e[1] for e in EQUIP}

# wd(-k) 异常行：设备, 说明, 处理状态
DAILY_ABNORMAL = {
    (8, "CK-01"): ("冷却液浓度偏低", "已闭环"),
    (6, "CK-05"): ("导轨润滑不足", "已闭环"),
    (5, "SK-02"): ("主轴振动偏大", "已闭环"),
    (3, "CK-03"): ("液压站渗油", "未闭环"),
    (1, "SK-02"): ("主轴异响", "未闭环"),
}
RC02_ABNORMAL = ("温控仪表读数漂移", "未闭环")

MAINT_NAMED = [
    # wd(-k), 设备, 故障现象, 原因, 处理措施, 停机时长, 状态
    (28, "SK-02", "主轴异响", "主轴轴承磨损", "更换主轴轴承", 6, "已修复"),
    (14, "SK-02", "主轴异响、振动", "轴承预紧力不足", "调整轴承预紧", 4, "已修复"),
    (4, "SK-02", "主轴振动偏大", "原因待查", "临时调整，安排进一步检查", 3, "观察中"),
    (2, "CK-03", "液压站渗油", "液压站密封件老化", "密封件已申购，到货后更换", 0, "待备件"),
]
MAINT_FILLER = [
    # wd(-k), 设备
    (38, "CK-01"), (34, "CK-02"), (30, "CK-04"), (24, "CK-05"),
    (20, "SK-01"), (16, "XK-01"), (11, "RC-01"), (6, "KY-01"),
]
MAINT_FILLER_TEXT = {
    "CK-01": ("导轨爬行", "导轨镶条间隙偏大", "调整镶条间隙，补充导轨油"),
    "CK-02": ("刀塔换刀不到位", "刀塔定位销磨损", "更换定位销，重新标定刀位"),
    "CK-04": ("尾座顶紧力不足", "尾座液压缸内漏", "更换密封件，重新调试压力"),
    "CK-05": ("排屑器卡滞", "铁屑缠绕堵塞", "清理排屑器，加装防缠挡板"),
    "SK-01": ("深孔钻冷却压力不足", "高压泵滤芯堵塞", "更换滤芯，恢复额定压力"),
    "XK-01": ("分度头定位误差", "蜗轮蜗杆磨损间隙超差", "调整蜗轮蜗杆间隙"),
    "RC-01": ("炉温均匀性超差", "热电偶老化漂移", "更换热电偶，重新标定"),
    "KY-01": ("排气量不足", "进气滤芯堵塞，活塞环磨损", "更换滤芯，检查活塞环"),
}


def build_03(out_root: Path, T: date, summary: list) -> None:
    d = out_root / "03-设备点检"

    def resolve(tag):
        kind, n = tag
        if kind == "wd":
            return wd(T, n)
        if kind == "cal":
            return cal(T, n)
        if kind == "t":
            return T
        raise ValueError(tag)

    equip_headers = ["设备编号", "设备名称", "型号", "车间", "点检周期", "上次点检日期",
                      "下次应检日期", "责任岗位"]
    equip_rows = []
    for code, name, model, ws_name, cycle, last_tag, next_tag in EQUIP:
        equip_rows.append([
            code, name, model, ws_name, cycle, resolve(last_tag), resolve(next_tag),
            "设备动力组",
        ])
    wb_e, ws_e = new_workbook_with_sheet("设备台账")
    write_table(ws_e, equip_headers, equip_rows, date_cols=(5, 6))
    p = d / "设备台账.xlsx"
    wb_e.save(p)
    summary.append((p, len(equip_rows)))

    rng_c = random.Random(20260301)
    check_headers = ["日期", "设备编号", "设备名称", "点检项目", "结果", "异常说明", "处理状态"]
    check_rows = []
    for k in range(10, 0, -1):  # wd(-10) ... wd(-1)，日期偏移升序
        for device in DAILY_DEVICES:
            item = INSPECT_ITEMS[device]
            if (k, device) in DAILY_ABNORMAL:
                note, status = DAILY_ABNORMAL[(k, device)]
                result = "异常"
            else:
                note, status, result = "", "", "正常"
            check_rows.append([
                wd(T, k), device, DEVICE_NAME[device], item, result, note, status,
            ])
    # 每周/每月设备，按设备编号升序，统一排在最后
    last_check_tag = {e[0]: e[5] for e in EQUIP}
    for device in sorted(WEEKLY_MONTHLY_DEVICES):
        item = INSPECT_ITEMS[device]
        d_date = resolve(last_check_tag[device])
        if device == "RC-02":
            note, status, result = RC02_ABNORMAL[0], RC02_ABNORMAL[1], "异常"
        else:
            note, status, result = "", "", "正常"
        check_rows.append([d_date, device, DEVICE_NAME[device], item, result, note, status])
    wb_c, ws_c = new_workbook_with_sheet("点检记录")
    write_table(ws_c, check_headers, check_rows, date_cols=(0,))
    p = d / "点检表.xlsx"
    wb_c.save(p)
    summary.append((p, len(check_rows)))

    maint_headers = ["日期", "设备编号", "设备名称", "故障现象", "原因", "处理措施",
                      "停机时长(h)", "状态"]
    maint_entries = []
    for k, device, phenom, cause, action, hours, status in MAINT_NAMED:
        maint_entries.append((k, device, phenom, cause, action, hours, status))
    for k, device in MAINT_FILLER:
        phenom, cause, action = MAINT_FILLER_TEXT[device]
        maint_entries.append((k, device, phenom, cause, action, 2, "已修复"))
    maint_entries.sort(key=lambda r: -r[0])  # 偏移升序 = -k 升序 = k 降序
    maint_rows = []
    for k, device, phenom, cause, action, hours, status in maint_entries:
        maint_rows.append([
            wd(T, k), device, DEVICE_NAME[device], phenom, cause, action, hours, status,
        ])
    wb_m, ws_m = new_workbook_with_sheet("维修记录")
    write_table(ws_m, maint_headers, maint_rows, date_cols=(0,))
    p = d / "维修记录.xlsx"
    wb_m.save(p)
    summary.append((p, len(maint_rows)))


# ============================================================================
# 04 · 备件盘点
# ============================================================================

@dataclass
class Part:
    code: str
    name: str
    spec: str
    unit: str
    current: int
    safety: int
    location: str
    fits: str


PARTS_NAMED = {
    "BJ-007": ("O形圈（氟橡胶）", "Φ120×5.3", "件", 12, 40, "库房A-02-03", "通用密封"),
    "BJ-011": ("深孔钻刀片", "Φ60 机夹式", "片", 20, 20, "库房B-01-05", "SK 系列深孔钻床"),
    "BJ-015": ("抗磨液压油", "46#，200 L/桶", "桶", 2, 6, "库房C-仓储区", "液压系统通用"),
    "BJ-021": ("螺纹车刀片", "16ER AG60", "片", 15, 30, "库房B-02-02", "CK 系列数控车床"),
}

# 其余 26 种自拟：编码 -> (名称, 规格型号, 单位, 库位, 适用设备/工具)
PARTS_OTHER = {
    "BJ-001": ("O形圈（丁腈）", "Φ50×3.55", "件", "库房A-01-01", "通用密封"),
    "BJ-002": ("O形圈（丁腈）", "Φ80×5.3", "件", "库房A-01-02", "通用密封"),
    "BJ-003": ("O形圈（氟橡胶）", "Φ100×5.3", "件", "库房A-01-03", "通用密封"),
    "BJ-004": ("密封件套件", "6-1/2 液压震击器用", "套", "库房A-02-01", "液压震击器"),
    "BJ-005": ("密封件套件", "4-3/4 液压震击器用", "套", "库房A-02-02", "液压震击器"),
    "BJ-006": ("密封件套件", "8 减震器用", "套", "库房A-02-04", "减震器"),
    "BJ-008": ("深孔钻刀片", "Φ50 机夹式", "片", "库房B-01-04", "SK 系列深孔钻床"),
    "BJ-009": ("车刀片", "CNMG120408", "片", "库房B-02-01", "CK 系列数控车床"),
    "BJ-010": ("车刀片", "DNMG150608", "片", "库房B-02-03", "CK 系列数控车床"),
    "BJ-012": ("深孔钻刀片", "Φ73 机夹式", "片", "库房B-01-06", "SK 系列深孔钻床"),
    "BJ-013": ("螺纹车刀片", "11ER AG60", "片", "库房B-02-04", "CK 系列数控车床"),
    "BJ-014": ("深沟球轴承", "6210", "个", "库房D-01-01", "SK-01/SK-02 主轴"),
    "BJ-016": ("抗磨液压油", "32#，200 L/桶", "桶", "库房C-仓储区", "液压系统通用"),
    "BJ-017": ("滤芯", "液压回油滤芯", "个", "库房D-02-01", "液压系统通用"),
    "BJ-018": ("滤芯", "空压机油气分离滤芯", "个", "库房D-02-02", "KY-01 空压机"),
    "BJ-019": ("碟簧", "100×51×4", "片", "库房D-03-01", "震击器总成"),
    "BJ-020": ("卡瓦", "6-1/2 通用型", "套", "库房D-03-03", "钻具卡瓦"),
    "BJ-022": ("皮带", "A 型三角带", "根", "库房D-04-01", "KY-01 空压机"),
    "BJ-023": ("冷却液", "乳化型浓缩液，20 L/桶", "桶", "库房C-仓储区", "机加车间"),
    "BJ-024": ("深沟球轴承", "6212", "个", "库房D-01-02", "SK-01/SK-02 主轴"),
    "BJ-025": ("深孔钻刀片", "Φ89 机夹式", "片", "库房B-01-07", "SK 系列深孔钻床"),
    "BJ-026": ("卡瓦", "8 通用型", "套", "库房D-03-04", "钻具卡瓦"),
    "BJ-027": ("O形圈（丁腈）", "Φ150×5.3", "件", "库房A-01-04", "通用密封"),
    "BJ-028": ("滤芯", "冷却液过滤网", "个", "库房D-02-03", "机加车间"),
    "BJ-029": ("密封件套件", "6-1/4 机械震击器用", "套", "库房A-02-05", "机械震击器"),
    "BJ-030": ("皮带", "B 型三角带", "根", "库房D-04-02", "KY-01 空压机"),
}

ACTIVE_CODES = ["BJ-002", "BJ-004", "BJ-009", "BJ-013", "BJ-018", "BJ-024"]
# code -> (出库_last10, 出库_first30, 入库件数)
# 这里的出库笔数，加上 BJ-021 的 12 笔、EXTRA_OUTBOUND_ONLY 的 12 笔和入库笔数，
# 合计是出入库记录的 80 行。
ACTIVE_PLAN = {
    "BJ-002": (2, 8, 2),
    "BJ-004": (2, 6, 2),
    "BJ-009": (2, 6, 2),
    "BJ-013": (2, 5, 1),
    "BJ-018": (2, 5, 1),
    "BJ-024": (1, 5, 2),
}

# BJ-007、BJ-015：只出库、不入库，消耗平稳。
# code -> (出库_last10, 出库_first30, 每次出库数量)
EXTRA_OUTBOUND_ONLY = {
    "BJ-007": (1, 5, 6),
    "BJ-015": (1, 5, 1),
}

BJ021_OUT_K = [40, 35, 30, 25, 20, 15, 10, 8, 6, 4, 2, 1]

# 备件品类 -> (领用部门, 用途候选)。按名称关键字分类，覆盖台账里全部 30 种备件。
CATEGORY_DEPT_PURPOSE = [
    (("刀片",), "机加车间", ["刀片磨损更换", "换型备刀"]),
    (("密封件套件", "O形圈", "碟簧", "卡瓦"), "装配车间", ["装配领用", "返修更换"]),
    (("轴承", "滤芯", "皮带", "液压油", "冷却液"), "设备动力组", ["设备保养", "设备维修领用"]),
]


def _category_for(name: str) -> tuple[str, list[str]]:
    for keywords, dept, purposes in CATEGORY_DEPT_PURPOSE:
        if any(kw in name for kw in keywords):
            return dept, purposes
    raise ValueError(f"未分类的备件名称: {name}")


def _even_ks(n: int, lo: int, hi: int) -> list[int]:
    if n <= 0:
        return []
    if n == 1:
        return [round((lo + hi) / 2)]
    used = set()
    out = []
    for i in range(n):
        k = round(lo + i * (hi - lo) / (n - 1))
        while k in used and k < hi:
            k += 1
        used.add(k)
        out.append(k)
    return sorted(set(out)) if len(set(out)) == n else out


def build_04(out_root: Path, T: date, summary: list) -> None:
    d = out_root / "04-备件盘点"
    rng_stock = random.Random(20260401)
    rng_txn = random.Random(20260402)
    rng_txt = random.Random(20260403)

    # ---- 出入库记录（先算，因为台账要用最近入库日期） ----
    txn_headers = ["日期", "单据号", "备件编码", "备件名称", "类型", "数量", "领用部门", "用途"]
    txn_entries = []  # (k, code, type, qty, dept, purpose)

    for k in BJ021_OUT_K:
        purpose = "螺纹加工" if k >= 15 else "CK-03 螺纹加工"
        txn_entries.append((k, "BJ-021", "出库", 5, "机加车间", purpose))

    def part_name(code: str) -> str:
        return PARTS_NAMED[code][0] if code in PARTS_NAMED else PARTS_OTHER[code][0]

    active_in_ks: dict[str, list[int]] = {}
    for code in ACTIVE_CODES:
        last10, first30, in_n = ACTIVE_PLAN[code]
        ks_last10 = _even_ks(last10, 1, 10)
        ks_first30 = _even_ks(first30, 11, 40)
        # 同一种备件每次出库数量固定，这样「最近10个工作日 vs 此前平均」的比例只由
        # 笔数决定，不会因为随机数量抽样而意外超过 1.5 倍的门槛。
        qty_out = rng_txn.randint(2, 5)
        # 领用部门/用途由备件品类决定，不是随机配的。
        dept, purposes = _category_for(part_name(code))
        for k in ks_last10 + ks_first30:
            purpose = rng_txn.choice(purposes)
            txn_entries.append((k, code, "出库", qty_out, dept, purpose))
        ks_in = _even_ks(in_n, 15, 39)
        active_in_ks[code] = ks_in
        for k in ks_in:
            txn_entries.append((k, code, "入库", rng_txn.randint(20, 60), "仓储物流组", "采购入库"))

    # BJ-007、BJ-015：只出库、不入库，消耗平稳（各自在 wd(-40)…wd(-1) 里大致均匀分布）。
    for code, (last10, first30, qty_out) in EXTRA_OUTBOUND_ONLY.items():
        ks_last10 = _even_ks(last10, 1, 10)
        ks_first30 = _even_ks(first30, 11, 40)
        dept, purposes = _category_for(part_name(code))
        for k in ks_last10 + ks_first30:
            purpose = rng_txn.choice(purposes)
            txn_entries.append((k, code, "出库", qty_out, dept, purpose))

    # 排序：日期偏移升序（k 降序），同日按备件编码、类型（入库先）
    def txn_sort_key(e):
        k, code, typ, qty, dept, purpose = e
        return (-k, code, 0 if typ == "入库" else 1)

    txn_entries.sort(key=txn_sort_key)

    rk_seq = 0
    ck_seq = 0
    txn_rows = []
    for k, code, typ, qty, dept, purpose in txn_entries:
        name = PARTS_NAMED[code][0] if code in PARTS_NAMED else PARTS_OTHER[code][0]
        if typ == "入库":
            rk_seq += 1
            doc_no = f"RK-{rk_seq:04d}"
        else:
            ck_seq += 1
            doc_no = f"CK-{ck_seq:04d}"
        txn_rows.append([wd(T, k), doc_no, code, name, typ, qty, dept, purpose])

    wb_t, ws_t = new_workbook_with_sheet("出入库")
    write_table(ws_t, txn_headers, txn_rows, date_cols=(0,))
    p = d / "出入库记录.xlsx"
    wb_t.save(p)
    summary.append((p, len(txn_rows)))

    # ---- 备件台账 ----
    codes = sorted(set(PARTS_NAMED) | set(PARTS_OTHER))
    stock_headers = ["备件编码", "备件名称", "规格型号", "单位", "当前库存", "安全库存",
                      "库位", "适用设备/工具", "最近入库日期"]
    stock_rows = []
    no_stock_default_k = {c: 45 + (i % 15) for i, c in enumerate(codes)}
    for code in codes:
        if code in PARTS_NAMED:
            name, spec, unit, current, safety, loc, fits = PARTS_NAMED[code]
        else:
            name, spec, unit, loc, fits = PARTS_OTHER[code]
            safety = rng_stock.randint(15, 35)
            current = round(safety * rng_stock.uniform(1.25, 1.6))
        if code in active_in_ks and active_in_ks[code]:
            last_in_k = min(active_in_ks[code])
            recent_date = wd(T, last_in_k)
        else:
            recent_date = wd(T, no_stock_default_k[code])
        stock_rows.append([code, name, spec, unit, current, safety, loc, fits, recent_date])
    wb_s, ws_s = new_workbook_with_sheet("备件台账")
    write_table(ws_s, stock_headers, stock_rows, date_cols=(8,))
    p = d / "备件台账.xlsx"
    wb_s.save(p)
    summary.append((p, len(stock_rows)))


# ============================================================================
# 05 · 返厂维修（固定日期，与 T 无关）
# ============================================================================

FIVE_START = date(2026, 1, 5)
FIVE_END = date(2026, 9, 18)

# (工具类型, 故障部位, 行数)
GH_GROUPS = [
    ("液压震击器", "密封失效", 6),
    ("液压震击器", "密封件损坏", 4),
    ("液压震击器", "密封圈老化", 3),
    ("螺杆钻具", "定子橡胶脱胶", 8),
    ("螺杆钻具", "传动轴轴承磨损", 5),
    ("减震器", "花键磨损", 5),
    ("液压震击器", "芯轴拉伤", 4),
    ("液压震击器", "心轴拉伤", 2),
    ("机械震击器", "卡瓦磨损", 3),
    ("减震器", "密封失效", 3),
    ("螺杆钻具", "螺纹粘扣", 3),
    ("液压震击器", "丝扣损伤", 2),
]
TJ_GROUPS = [
    ("液压震击器", "密封漏油", 5),
    ("液压震击器", "O形圈损坏", 3),
    ("液压震击器", "密封失效", 2),
    ("螺杆钻具", "定子掉胶", 4),
    ("螺杆钻具", "定子橡胶脱胶", 2),
    ("螺杆钻具", "万向轴断裂", 3),
    ("减震器", "花键磨损", 4),
    ("减震器", "碟簧断裂", 3),
    ("机械震击器", "卡瓦磨损", 3),
    ("液压震击器", "芯轴拉伤", 3),
    ("加速器", "密封失效", 2),
    ("螺杆钻具", "螺纹粘扣", 2),
]
XJ_GROUPS = [
    ("液压震击器", "密封失效", 4),
    ("液压震击器", "密封件损坏", 3),
    ("螺杆钻具", "定子橡胶脱胶", 5),
    ("螺杆钻具", "传动轴轴承磨损", 4),
    ("减震器", "花键磨损", 2),
    ("液压震击器", "心轴拉伤", 3),
    ("机械震击器", "卡瓦磨损", 2),
    ("螺杆钻具", "丝扣损伤", 3),
    ("减震器", "密封圈老化", 2),
    ("加速器", "密封失效", 2),
]

#  fault -> ([2-3 条故障描述/现象], [2 条处理措施])，每行从各自列表里轮换选一条；
#  文字不能和「故障部位」列一字不差。三个文件文风不同：广汉细、天津中、新疆最简。
GH_TEXT = {
    "密封失效": (["下井作业后检出内漏，密封失效", "起出后试压不保压，密封部位失效"],
                 ["更换整套密封件，试压合格后入库", "更换密封件并复核配合面，试压合格后入库"]),
    "密封件损坏": (["返修检查发现密封件表面划伤破损", "拆检时发现密封件局部挤出损坏"],
                  ["更换损坏密封件，检查配合面后重新装配", "更换密封件并打磨配合面毛刺"]),
    "密封圈老化": (["密封圈老化变硬，弹性不足", "密封圈表面龟裂，失去弹性"],
                  ["更换全部密封圈", "按周期更换全部密封圈并记录批次"]),
    "定子橡胶脱胶": (["定子橡胶层局部脱胶起泡", "定子内壁橡胶与钢套脱层"],
                   ["更换定子总成，报废旧定子", "定子总成整体更换并报废"]),
    "传动轴轴承磨损": (["传动轴轴承磨损，间隙超差", "传动轴轴承滚道点蚀，径向间隙偏大"],
                     ["更换传动轴轴承，重新装配校验", "更换轴承并复测同心度"]),
    "花键磨损": (["花键磨损，配合间隙增大", "花键齿面磨损，传动有明显冲击"],
                ["更换花键套，检查配合公差", "更换花键套并复核扭矩传递"]),
    "芯轴拉伤": (["芯轴表面拉伤，有明显划痕", "芯轴外圆多处拉毛，密封面受损"],
                ["芯轴表面修复处理，抛光后检测", "芯轴局部堆焊修复后磨削抛光"]),
    "心轴拉伤": (["心轴表面拉伤，有明显划痕", "心轴外圆多处拉毛，密封面受损"],
                ["心轴表面修复处理，抛光后检测", "心轴局部堆焊修复后磨削抛光"]),
    "卡瓦磨损": (["卡瓦牙型磨损，咬合力下降", "卡瓦牙尖磨圆，抓持力不足"],
                ["更换卡瓦，检查配合间隙", "更换卡瓦并复核抓持力"]),
    "螺纹粘扣": (["上卸扣时螺纹粘扣，扣型受损", "螺纹上扣阻力异常，拆检发现损伤"],
                ["螺纹修复或更换接头，做通止规检查", "更换接头并做通止规复检"]),
    "丝扣损伤": (["丝扣局部损伤，密封面受损", "丝扣根部有磕碰损伤"],
                ["修复丝扣或更换零件，做通止规检查", "更换零件并做通止规复检"]),
}
TJ_TEXT = {
    "密封漏油": (["密封处渗油，压力保不住", "密封位漏油明显"], ["更换密封件，试压确认", "更换密封件后复压"]),
    "O形圈损坏": (["O形圈挤压损坏，出现裂纹", "O形圈老化开裂"], ["更换O形圈", "更换O形圈并检查沟槽"]),
    "密封失效": (["密封失效，内漏", "密封不保压"], ["更换密封件", "更换密封件并试压"]),
    "定子掉胶": (["定子橡胶局部掉胶", "定子胶层起皮脱落"], ["更换定子总成", "定子整体更换"]),
    "定子橡胶脱胶": (["定子橡胶脱胶起层", "定子胶层与钢套分离"], ["更换定子总成", "定子整体更换"]),
    "万向轴断裂": (["万向轴疲劳断裂", "万向轴根部断裂"], ["更换万向轴总成", "更换万向轴并复检同心度"]),
    "花键磨损": (["花键磨损超差", "花键齿侧磨损明显"], ["更换花键套", "更换花键套并复检间隙"]),
    "碟簧断裂": (["碟簧组断裂", "碟簧局部裂纹断裂"], ["更换碟簧组", "更换碟簧组并复检预紧力"]),
    "卡瓦磨损": (["卡瓦牙型磨损", "卡瓦抓持力下降"], ["更换卡瓦", "更换卡瓦并测试抓持力"]),
    "芯轴拉伤": (["芯轴表面拉伤", "芯轴外圆有划痕"], ["芯轴修复抛光", "芯轴抛光后复测圆度"]),
    "螺纹粘扣": (["螺纹粘扣损伤", "上卸扣时粘扣"], ["修复螺纹或更换接头", "更换接头并做通止规检查"]),
}
XJ_TEXT = {
    "密封失效": (["密封失效，内漏", "内漏保不住压"], ["更换密封件", "换密封"]),
    "密封件损坏": (["密封件破损，渗油", "密封处漏油"], ["更换密封件", "换密封件"]),
    "定子橡胶脱胶": (["定子脱胶", "定子掉胶"], ["更换定子", "换定子总成"]),
    "传动轴轴承磨损": (["轴承磨损", "传动轴间隙偏大"], ["更换轴承", "换轴承"]),
    "花键磨损": (["花键齿磨损", "花键间隙偏大"], ["更换花键套", "换花键套"]),
    "心轴拉伤": (["心轴表面拉伤", "心轴有划痕"], ["心轴抛光", "心轴修复"]),
    "卡瓦磨损": (["卡瓦牙磨损", "卡瓦打滑"], ["更换卡瓦", "换卡瓦"]),
    "丝扣损伤": (["丝扣局部损伤", "丝扣碰伤"], ["修复丝扣", "换接头"]),
    "密封圈老化": (["密封圈老化变硬", "密封圈发脆"], ["更换密封圈", "换密封圈"]),
}

CUSTOMERS = list("ABCDEFGH")


def _well_no(rng: random.Random, cust_letter: str) -> str:
    n = rng.randint(1, 199)
    suffix = "H" if rng.random() < 0.5 else ""
    if rng.random() < 0.3:
        return f"{cust_letter}-{rng.randint(1,20)}-{rng.randint(1,9)}"
    return f"{cust_letter}-{n}{suffix}"


# 三个文件用互不相交的号段，且一份文件里不重复使用同一个编号（同一支工具在同一份文件
# 里只出现一次，规格也就随编号唯一确定，不会再出现「同一编号两种规格」）。
FILE_NUMBER_RANGE = {
    "广汉": (1, 299),
    "天津": (300, 599),
    "新疆": (600, 899),
}
# 07 文件夹点名过的编号（不含前缀的数字部分），三个文件都要避开，避免和 07 的工具混淆。
NAMED_07_NUMBERS = {388, 301, 122, 215, 344, 45, 108, 366, 251, 230, 131, 402, 244, 115, 410, 51, 140}


def _draw_distinct_numbers(rng: random.Random, lo: int, hi: int, n: int) -> list[int]:
    forbidden = RESERVED_NUMBERS | NAMED_07_NUMBERS
    pool = [x for x in range(lo, hi + 1) if x not in forbidden]
    rng.shuffle(pool)
    if len(pool) < n:
        raise ValueError(f"号段 {lo}-{hi} 里可用编号不够抽 {n} 个")
    return pool[:n]


def _gen_group_rows(rng: random.Random, groups, text_map, numbers: list[int]):
    """按 (类型, 故障部位, 行数) 生成行，返回 [(date, tool, type, spec, cust, well, fault,
    desc, action, hours)]。`numbers` 是这份文件里互不相同的编号，按生成顺序消耗一个用一个，
    所以同一支工具（同一个编号）在这份文件里只会出现一次。"""
    numbers_iter = iter(numbers)
    rows = []
    span_days = (FIVE_END - FIVE_START).days
    for type_, fault, count in groups:
        desc_list, action_list = text_map[fault]
        for _ in range(count):
            num = next(numbers_iter)
            tool = f"{TYPE_PREFIX[type_]}-{num:04d}"
            spec = rng.choice(SPEC_POOL)
            cust_letter = rng.choice(CUSTOMERS)
            well = _well_no(rng, cust_letter)
            offset = rng.randint(0, span_days)
            d = FIVE_START + timedelta(days=offset)
            hours = rng.randint(3, 26)
            desc = rng.choice(desc_list)
            action = rng.choice(action_list)
            rows.append([d, tool, type_, spec, f"客户{cust_letter}", well, fault,
                         desc, action, hours])
    rows.sort(key=lambda r: r[0])
    return rows


def build_05(out_root: Path, T: date, summary: list) -> None:
    d = out_root / "05-返厂维修"

    rng_gh = random.Random(20260501)
    rng_tj = random.Random(20260502)
    rng_xj = random.Random(20260503)
    rng_gh_no = random.Random(20260511)
    rng_tj_no = random.Random(20260512)
    rng_xj_no = random.Random(20260513)

    gh_n = sum(c for _, _, c in GH_GROUPS)
    tj_n = sum(c for _, _, c in TJ_GROUPS)
    xj_n = sum(c for _, _, c in XJ_GROUPS)
    gh_numbers = _draw_distinct_numbers(rng_gh_no, *FILE_NUMBER_RANGE["广汉"], gh_n)
    tj_numbers = _draw_distinct_numbers(rng_tj_no, *FILE_NUMBER_RANGE["天津"], tj_n)
    xj_numbers = _draw_distinct_numbers(rng_xj_no, *FILE_NUMBER_RANGE["新疆"], xj_n)

    gh_rows = _gen_group_rows(rng_gh, GH_GROUPS, GH_TEXT, gh_numbers)
    tj_rows = _gen_group_rows(rng_tj, TJ_GROUPS, TJ_TEXT, tj_numbers)
    xj_rows = _gen_group_rows(rng_xj, XJ_GROUPS, XJ_TEXT, xj_numbers)

    # ---- 广汉-2026.xlsx ----
    wb_gh = Workbook()
    ws_main = wb_gh.active
    ws_main.title = "维修记录"
    style_workbook(wb_gh)
    gh_headers = ["序号", "返厂日期", "工具编号", "工具类型", "规格(OD,in)", "客户", "井号",
                  "故障部位", "故障描述", "处理措施", "维修工时(h)"]
    gh_table = []
    for i, r in enumerate(gh_rows, start=1):
        d_, tool, type_, spec, cust, well, fault, desc, action, hours = r
        gh_table.append([i, d_, tool, type_, spec, cust, well, fault, desc, action, hours])
    write_table(ws_main, gh_headers, gh_table, date_cols=(1,))
    ws_note = wb_gh.create_sheet("填写说明")
    note_lines = [
        "返厂日期按实际到厂日填写。",
        "工具编号、工具类型、规格(OD,in) 三项须与出库记录一致。",
        "故障部位填主要失效部位，一支工具多处失效的填最严重的一处。",
        "故障描述、处理措施用简明文字记录，便于统计归类。",
        "维修工时(h) 按实际投入工时填写，不含等待备件的时间。",
        "本表为培训用示例数据，不对应任何真实工具或客户。",
    ]
    for i, line in enumerate(note_lines, start=1):
        ws_note.cell(row=i, column=1, value=line)
    ws_note.column_dimensions["A"].width = 50
    wb_gh.active = 0  # 维修记录为活动工作表
    p = d / "广汉-2026.xlsx"
    wb_gh.save(p)
    summary.append((p, len(gh_table)))

    # ---- 天津-2026.xlsx ----
    wb_tj, ws_tj = new_workbook_with_sheet("Sheet1")
    tj_headers = ["日期", "工具号", "工具类别", "外径(in)", "客户", "井号", "失效部位",
                  "失效现象", "处理", "工时"]
    tj_table = []
    for r in tj_rows:
        d_, tool, type_, spec, cust, well, fault, desc, action, hours = r
        tj_table.append([d_, tool, type_, spec, cust, well, fault, desc, action, hours])
    write_table(ws_tj, tj_headers, tj_table, date_cols=(0,))
    total_hours = sum(r[-1] for r in tj_table)
    sum_row = len(tj_table) + 2
    ws_tj.cell(row=sum_row, column=1, value="合计")
    ws_tj.cell(row=sum_row, column=2, value=f"{len(tj_table)} 条")
    ws_tj.cell(row=sum_row, column=10, value=total_hours)
    p = d / "天津-2026.xlsx"
    wb_tj.save(p)
    summary.append((p, len(tj_table)))

    # ---- 新疆-2026.xlsx ----
    wb_xj = Workbook()
    ws_xj = wb_xj.active
    ws_xj.title = "返厂登记"
    style_workbook(wb_xj)
    xj_headers = ["登记日期", "工具编号", "类型", "规格", "甲方", "井号", "故障位置",
                  "情况说明", "维修措施", "工时(h)"]
    xj_table = []
    for r in xj_rows:
        d_, tool, type_, spec, cust, well, fault, desc, action, hours = r
        xj_table.append([d_, tool, type_, spec, cust, well, fault, desc, action, hours])
    title_cell = ws_xj.cell(row=1, column=1, value="新疆服务中心 2026 年返厂维修登记表（示例数据）")
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_xj.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(xj_headers))
    write_table(ws_xj, xj_headers, xj_table, date_cols=(0,), header_row=2, freeze="A3")
    p = d / "新疆-2026.xlsx"
    wb_xj.save(p)
    summary.append((p, len(xj_table)))


# ============================================================================
# 07 · 租赁台账
# ============================================================================

RATE_TABLE = {
    ("液压震击器", "6-1/2"): 850,
    ("液压震击器", "4-3/4"): 650,
    ("液压震击器", "8"): 980,
    ("减震器", "8"): 620,
    ("减震器", "6-1/2"): 520,
    ("螺杆钻具", "6-3/4"): 1200,
    ("螺杆钻具", "4-3/4"): 900,
    ("机械震击器", "6-1/4"): 680,
    ("机械震击器", "4-3/4"): 560,
    ("加速器", "6-1/2"): 480,
    ("加速器", "8"): 560,
}

# 工具编号, 类型, 规格, 客户, 井号, 出库偏移, 到期日("off"数字 或 None 或 "待定"), 备注
OUT_NAMED = [
    ("YZJ-0388", "液压震击器", "6-1/2", "客户C", "C-18", -95, -50, ""),
    ("YZJ-0301", "液压震击器", "6-1/2", "客户A", "A-101H", -81, -21, ""),
    ("JZQ-0122", "减震器", "8", "客户B", "B-7-2", -72, -12, ""),
    ("LG-0215", "螺杆钻具", "6-3/4", "客户C", "C-23H", -52, -7, ""),
    ("YZJ-0344", "液压震击器", "4-3/4", "客户D", "D-5", -48, -3, ""),
    ("JSQ-0045", "加速器", "6-1/2", "客户B", "B-9-1", -44, None, ""),
    ("JZJ-0108", "机械震击器", "6-1/4", "客户A", "A-102H", -31, -1, ""),
    ("YZJ-0366", "液压震击器", "6-1/2", "客户E", "E-11H", -30, 0, ""),
    ("LG-0251", "螺杆钻具", "4-3/4", "客户B", "B-10-3", -29, 1, ""),
    ("LG-0230", "螺杆钻具", "4-3/4", "客户F", "F-3H", -26, "待定", "合同续签中"),
    ("JZQ-0131", "减震器", "6-1/2", "客户D", "D-6", -25, 5, ""),
    ("YZJ-0388", "液压震击器", "6-1/2", "客户G", "G-2H", -20, 10, ""),
    ("YZJ-0402", "液压震击器", "8", "客户H", "H-14H", -18, 12, ""),
    ("LG-0244", "螺杆钻具", "6-3/4", "客户A", "A-103H", -15, 15, ""),
    ("JZJ-0115", "机械震击器", "4-3/4", "客户E", "E-12H", -12, 18, ""),
    ("YZJ-0410", "液压震击器", "6-1/2", "客户F", "F-4H", -10, 20, ""),
    ("JSQ-0051", "加速器", "8", "客户G", "G-3H", -8, 22, ""),
    ("JZQ-0140", "减震器", "8", "客户H", "H-15H", -5, 25, ""),
]

NAMED_OUT_OFFSETS = {r[5] for r in OUT_NAMED}

OTHER_OUT_OFFSETS = [
    -170, -160, -150, -142, -135, -128, -120, -112, -105,
    -98, -90, -85, -78, -68, -62, -55, -42, -38,
]
assert len(OTHER_OUT_OFFSETS) == 18
assert not (set(OTHER_OUT_OFFSETS) & NAMED_OUT_OFFSETS)

# 维修中的备注要和工具类型对得上，不能随手轮换。
REPAIR_REMARK_BY_TYPE = {
    "液压震击器": ["更换密封件和液压油"],
    "机械震击器": ["更换卡瓦"],
    "减震器": ["更换碟簧", "更换花键套"],
    "螺杆钻具": ["更换传动轴轴承"],
    "加速器": ["更换密封件"],
}


def build_07(out_root: Path, T: date, summary: list) -> None:
    d = out_root / "07-租赁台账"
    rng = random.Random(20260701)

    used_ids = {r[0] for r in OUT_NAMED}
    rate_combos = list(RATE_TABLE.keys())

    other_rows = []
    zx_idx = {3, 7, 11, 15}  # 18 行里这 4 个下标是"维修中"
    for i, out_off in enumerate(OTHER_OUT_OFFSETS):
        type_, spec = rate_combos[i % len(rate_combos)]
        prefix = TYPE_PREFIX[type_]
        while True:
            num = rng.randint(1, 550)
            if num in RESERVED_NUMBERS:
                continue
            tid = f"{prefix}-{num:04d}"
            if tid not in used_ids and tid not in RESERVED_TOOL_IDS:
                used_ids.add(tid)
                break
        cust_letter = rng.choice(CUSTOMERS)
        well = _well_no(rng, cust_letter)
        margin = min(120, abs(out_off) - 15)
        duration = rng.randint(20, max(21, margin))
        due_off = out_off + duration
        if due_off > -1:
            due_off = -1
        ret_off = due_off + rng.randint(-8, 0)
        if ret_off < out_off + 3:
            ret_off = out_off + 3
        if ret_off > -1:
            ret_off = -1
        status = "维修中" if i in zx_idx else "已归还"
        remark = rng.choice(REPAIR_REMARK_BY_TYPE[type_]) if i in zx_idx else ""
        other_rows.append({
            "tool": tid, "type": type_, "spec": spec, "cust": f"客户{cust_letter}",
            "well": well, "out_off": out_off, "due_off": due_off, "ret_off": ret_off,
            "status": status, "remark": remark,
        })

    all_rows = []
    for r in OUT_NAMED:
        tool, type_, spec, cust, well, out_off, due, remark = r
        all_rows.append({
            "tool": tool, "type": type_, "spec": spec, "cust": cust, "well": well,
            "out_off": out_off, "due": due, "ret_off": None, "status": "在外",
            "remark": remark,
        })
    for r in other_rows:
        all_rows.append({
            "tool": r["tool"], "type": r["type"], "spec": r["spec"], "cust": r["cust"],
            "well": r["well"], "out_off": r["out_off"], "due": r["due_off"],
            "ret_off": r["ret_off"], "status": r["status"], "remark": r["remark"],
        })

    all_rows.sort(key=lambda r: (r["out_off"], r["tool"]))

    headers = ["序号", "工具编号", "工具类型", "规格(OD,in)", "客户", "井号", "出库日期",
               "合同到期日", "归还日期", "状态", "日租金(元)", "备注"]
    table = []
    for i, r in enumerate(all_rows, start=1):
        out_date = cal(T, r["out_off"])
        due = r["due"]
        if due is None:
            due_val = None
        elif due == "待定":
            due_val = "待定"
        else:
            due_val = cal(T, due)
        ret_val = cal(T, r["ret_off"]) if r["ret_off"] is not None else None
        rate = RATE_TABLE[(r["type"], r["spec"])]
        table.append([
            i, r["tool"], r["type"], r["spec"], r["cust"], r["well"], out_date,
            due_val, ret_val, r["status"], rate, r["remark"],
        ])

    wb, ws = new_workbook_with_sheet("租赁台账")
    write_table(ws, headers, table, date_cols=(6, 7, 8))
    p = d / "租赁台账.xlsx"
    wb.save(p)
    summary.append((p, len(table)))


# ============================================================================
# 回读 + 断言 + 生成答案文档
# ============================================================================


class Checker:
    def __init__(self):
        self.failures: list[str] = []

    def check(self, cond: bool, msg: str) -> None:
        if not cond:
            self.failures.append(msg)

    def finish_or_exit(self) -> None:
        if self.failures:
            sys.stderr.write("自检失败，共 %d 条：\n" % len(self.failures))
            for m in self.failures:
                sys.stderr.write(f"  - {m}\n")
            sys.exit(1)


def _norm(v: Any) -> Any:
    """openpyxl 读回日期单元格时给的是 datetime.datetime；统一收敛成 date 以便比较。"""
    if isinstance(v, datetime):
        return v.date()
    return v


def _rows(ws: Worksheet, header_row: int, first: int, last: int) -> list[tuple]:
    out = []
    for r in range(first, last + 1):
        out.append(tuple(_norm(ws.cell(row=r, column=c).value) for c in range(1, ws.max_column + 1)))
    return out


def _no_reserved_ids(ck: Checker, path: Path, wb) -> None:
    reserved_suffixes = [f"-{n:04d}" for n in RESERVED_NUMBERS]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for rid in RESERVED_TOOL_IDS:
                        ck.check(rid not in cell.value, f"{path.name} 出现了保留编号 {rid}（{cell.coordinate}）")
                    for suf in reserved_suffixes:
                        ck.check(suf not in cell.value,
                                  f"{path.name} 出现了保留编号数字 {suf}（{cell.coordinate}：{cell.value}）")


def _no_formula(ck: Checker, path: Path, wb) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    ck.check(False, f"{path.name} {cell.coordinate} 是公式而不是纯值：{cell.value}")


def verify_01(out_root: Path, T: date, ck: Checker) -> dict:
    d = out_root / "01-质量周报"
    F = friday_on_or_before(T)
    wk = iso_week(F)
    stat: dict[str, Any] = {"wk": wk, "F": F}

    files = {
        "机加车间": "检验记录-机加车间.xlsx",
        "热处理车间": "检验记录-热处理车间.xlsx",
        "装配车间": "检验记录-装配车间.xlsx",
    }
    totals = {}
    rows_by_ws: dict[str, list] = {}
    for ws_name, fname in files.items():
        p = d / fname
        wb = load_workbook(p, data_only=True)
        ws = wb["检验记录"]
        _no_formula(ck, p, wb)
        rows = _rows(ws, 1, 2, ws.max_row)
        expect_n = {"机加车间": 25, "热处理车间": 10, "装配车间": 15}[ws_name]
        ck.check(len(rows) == expect_n, f"{fname} 行数应为 {expect_n}，实际 {len(rows)}")
        bad_total = sum(r[8] for r in rows)
        repair = sum(r[8] for r in rows if r[10] == "返修")
        scrap = sum(r[8] for r in rows if r[10] == "报废")
        totals[ws_name] = (bad_total, repair, scrap, len(rows))
        rows_by_ws[ws_name] = rows

    ck.check(totals["机加车间"][:3] == (16, 15, 1), f"机加车间合计应为不良16/返修15/报废1，实际 {totals['机加车间'][:3]}")
    ck.check(totals["热处理车间"][:3] == (8, 4, 4), f"热处理车间合计应为不良8/返修4/报废4，实际 {totals['热处理车间'][:3]}")
    ck.check(totals["装配车间"][:3] == (3, 3, 0), f"装配车间合计应为不良3/返修3/报废0，实际 {totals['装配车间'][:3]}")
    grand_bad = sum(v[0] for v in totals.values())
    grand_repair = sum(v[1] for v in totals.values())
    grand_scrap = sum(v[2] for v in totals.values())
    ck.check((grand_bad, grand_repair, grand_scrap) == (27, 22, 5),
             f"三车间合计应为不良27/返修22/报废5，实际 {(grand_bad, grand_repair, grand_scrap)}")

    # 螺纹加工只出现在 CK-03
    p = d / files["机加车间"]
    wb = load_workbook(p, data_only=True)
    ws = wb["检验记录"]
    rows = _rows(ws, 1, 2, ws.max_row)
    thread_devices = {r[5] for r in rows if r[4] == "螺纹加工"}
    ck.check(thread_devices == {"CK-03"}, f"螺纹加工应只出现在 CK-03，实际设备集合 {thread_devices}")

    # 热处理硬度不合格逐日 0/1/1/2/4
    p2 = d / files["热处理车间"]
    wb2 = load_workbook(p2, data_only=True)
    ws2 = wb2["检验记录"]
    rows2 = _rows(ws2, 1, 2, ws2.max_row)
    by_day: dict[date, int] = {}
    for r in rows2:
        by_day.setdefault(r[0], 0)
        if r[9] == "硬度不合格":
            by_day[r[0]] += r[8]
    days_sorted = sorted(by_day)
    seq = [by_day[dd] for dd in days_sorted]
    ck.check(seq == [0, 1, 1, 2, 4], f"热处理硬度不合格逐日应为 0/1/1/2/4，实际 {seq}")

    # 不良品处理单
    p3 = d / "不良品处理单.xlsx"
    wb3 = load_workbook(p3, data_only=True)
    ws3 = wb3["不良品处理单"]
    _no_formula(ck, p3, wb3)
    ncr_rows = _rows(ws3, 1, 2, ws3.max_row)
    ck.check(len(ncr_rows) == 8, f"不良品处理单行数应为 8，实际 {len(ncr_rows)}")
    ck.check(ncr_rows[0][0] == f"NCR-{wk:02d}-01", f"NCR 单号前缀应含 wk={wk:02d}，实际 {ncr_rows[0][0]}")

    # docx
    docx_path = d / "上周质量周报.docx"
    doc = Document(docx_path)
    heading_text = doc.paragraphs[0].text if doc.paragraphs else ""
    wk_prev = wk - 1
    ck.check(f"第 {wk_prev} 周" in heading_text, f"docx 标题应含“第 {wk_prev} 周”，实际“{heading_text}”")

    # 埋的问题：每个不良数 > 0 的行，文件名 + 行号 + 现象
    problems = []
    for ws_name, fname in files.items():
        for i, r in enumerate(rows_by_ws[ws_name]):
            if r[8] and r[8] > 0:
                excel_row = i + 2
                problems.append((fname, excel_row, r[5], r[4], r[3], r[8], r[9], r[10]))

    stat["totals"] = totals
    stat["grand"] = (grand_bad, grand_repair, grand_scrap)
    stat["rows_by_ws"] = rows_by_ws
    stat["ncr_rows"] = ncr_rows
    stat["problems"] = problems
    stat["files_listing"] = [
        (files["机加车间"], "检验记录", totals["机加车间"][3]),
        (files["热处理车间"], "检验记录", totals["热处理车间"][3]),
        (files["装配车间"], "检验记录", totals["装配车间"][3]),
        ("不良品处理单.xlsx", "不良品处理单", len(ncr_rows)),
        ("上周质量周报.docx", "（docx，无工作表）", None),
    ]
    stat["spot_checks"] = [
        ("检验记录-机加车间.xlsx", 2, rows_by_ws["机加车间"][0]),
        ("不良品处理单.xlsx", 2, ncr_rows[0]),
        ("不良品处理单.xlsx", 9, ncr_rows[7]),
    ]
    stat["wk_prev"] = wk_prev
    return stat


def verify_02(out_root: Path, T: date, ck: Checker) -> dict:
    d = out_root / "02-交期风险"
    stat: dict[str, Any] = {}

    p1 = d / "订单清单.xlsx"
    wb1 = load_workbook(p1, data_only=True)
    ws1 = wb1["订单"]
    _no_formula(ck, p1, wb1)
    orders = _rows(ws1, 1, 2, ws1.max_row)
    ck.check(len(orders) == 12, f"订单清单行数应为 12，实际 {len(orders)}")
    order_dates = {r[0]: (r[5], r[6]) for r in orders}

    p2 = d / "排产计划.xlsx"
    wb2 = load_workbook(p2, data_only=True)
    ws2 = wb2["排产"]
    _no_formula(ck, p2, wb2)
    sched = _rows(ws2, 1, 2, ws2.max_row)
    ck.check(len(sched) == 72, f"排产计划行数应为 72，实际 {len(sched)}")
    # 每张订单第一道工序开始 = 下单+3，最后一道工序完成 = 交期-2
    by_order: dict[str, list] = {}
    for r in sched:
        by_order.setdefault(r[0], []).append(r)
    for order_id, rs in by_order.items():
        rs.sort(key=lambda r: r[1])
        ck.check(len(rs) == 6, f"{order_id} 排产工序数应为 6，实际 {len(rs)}")
        order_date, due_date = order_dates[order_id]
        ck.check(rs[0][3] == order_date + timedelta(days=3),
                 f"{order_id} 第一道工序开始日期应为下单+3")
        ck.check(rs[-1][4] == due_date - timedelta(days=2),
                 f"{order_id} 最后一道工序完成日期应为交期-2")
        thread_rows = [r for r in rs if r[2] == "螺纹加工"]
        ck.check(len(thread_rows) == 1 and thread_rows[0][5] == "CK-03",
                 f"{order_id} 螺纹加工应固定在 CK-03")
    sk2_101 = [r for r in sched if r[0] == "SO-26-101" and r[1] == 4]
    ck.check(sk2_101 and sk2_101[0][5] == "SK-02", "SO-26-101 第 4 道工序应固定在 SK-02")

    p3 = d / "进度跟踪.xlsx"
    wb3 = load_workbook(p3, data_only=True)
    ws3 = wb3["进度"]
    _no_formula(ck, p3, wb3)
    prog = _rows(ws3, 1, 2, ws3.max_row)
    ck.check(len(prog) == 12, f"进度跟踪行数应为 12，实际 {len(prog)}")
    ck.check(all(r[3] == 6 for r in prog), "进度跟踪总工序数应全部为 6")
    ck.check(all(isinstance(r[4], int) for r in prog), "进度跟踪完成率应为整数")
    row103 = next(r for r in prog if r[0] == "SO-26-103")
    t4_text = cal(T, 4).isoformat()
    ck.check(t4_text in row103[7], f"SO-26-103 异常说明应含 {t4_text}，实际“{row103[7]}”")

    stat["order_count"] = len(orders)
    stat["sched_count"] = len(sched)
    stat["prog_count"] = len(prog)
    stat["spot_checks"] = [
        ("订单清单.xlsx", 2, orders[0]),
        ("排产计划.xlsx", 2, sched[0]),
        ("进度跟踪.xlsx", 4, row103),
    ]
    return stat


def verify_03(out_root: Path, T: date, ck: Checker) -> dict:
    d = out_root / "03-设备点检"
    stat: dict[str, Any] = {}

    p1 = d / "设备台账.xlsx"
    wb1 = load_workbook(p1, data_only=True)
    ws1 = wb1["设备台账"]
    _no_formula(ck, p1, wb1)
    equip = _rows(ws1, 1, 2, ws1.max_row)
    ck.check(len(equip) == 14, f"设备台账行数应为 14，实际 {len(equip)}")
    due_today = [r for r in equip if r[6] == T]
    overdue = [r for r in equip if isinstance(r[6], date) and r[6] < T]
    ck.check(len(due_today) == 10, f"今天应检应为 10 台，实际 {len(due_today)}")
    ck.check(len(overdue) == 1 and overdue[0][0] == "ZJ-01",
             f"逾期未检应为 1 台且为 ZJ-01，实际 {[r[0] for r in overdue]}")
    if overdue:
        overdue_days = (T - overdue[0][6]).days
        ck.check(overdue_days == 3, f"ZJ-01 逾期天数应为 3，实际 {overdue_days}")

    p2 = d / "点检表.xlsx"
    wb2 = load_workbook(p2, data_only=True)
    ws2 = wb2["点检记录"]
    _no_formula(ck, p2, wb2)
    checks = _rows(ws2, 1, 2, ws2.max_row)
    ck.check(len(checks) == 77, f"点检表行数应为 77，实际 {len(checks)}")
    abnormal = [r for r in checks if r[4] == "异常"]
    ck.check(len(abnormal) == 6, f"异常行数应为 6，实际 {len(abnormal)}")
    normal = [r for r in checks if r[4] == "正常"]
    ck.check(all((r[5] in (None, "") and r[6] in (None, "")) for r in normal),
              "正常行的异常说明/处理状态应留空")
    expect_abn = {
        ("CK-01", "冷却液浓度偏低", "已闭环"),
        ("CK-05", "导轨润滑不足", "已闭环"),
        ("SK-02", "主轴振动偏大", "已闭环"),
        ("CK-03", "液压站渗油", "未闭环"),
        ("SK-02", "主轴异响", "未闭环"),
        ("RC-02", "温控仪表读数漂移", "未闭环"),
    }
    got_abn = {(r[1], r[5], r[6]) for r in abnormal}
    ck.check(got_abn == expect_abn, f"异常行内容不符，实际 {got_abn}")

    p3 = d / "维修记录.xlsx"
    wb3 = load_workbook(p3, data_only=True)
    ws3 = wb3["维修记录"]
    _no_formula(ck, p3, wb3)
    maint = _rows(ws3, 1, 2, ws3.max_row)
    ck.check(len(maint) == 12, f"维修记录行数应为 12，实际 {len(maint)}")
    from collections import Counter
    dev_counts = Counter(r[1] for r in maint)
    over = {k: v for k, v in dev_counts.items() if v > 1 and k != "SK-02"}
    ck.check(not over, f"除 SK-02 外不应有设备出现两次以上，实际 {over}")
    ck.check(dev_counts.get("SK-02") == 3, f"SK-02 应出现 3 次，实际 {dev_counts.get('SK-02')}")

    stat["due_today"] = len(due_today)
    stat["overdue"] = overdue[0][0] if overdue else None
    stat["abnormal_count"] = len(abnormal)
    stat["spot_checks"] = [
        ("设备台账.xlsx", 2, equip[0]),
        ("点检表.xlsx", 2, checks[0]),
        ("维修记录.xlsx", 2, maint[0]),
    ]
    return stat


def verify_04(out_root: Path, T: date, ck: Checker) -> dict:
    d = out_root / "04-备件盘点"
    stat: dict[str, Any] = {}

    p1 = d / "备件台账.xlsx"
    wb1 = load_workbook(p1, data_only=True)
    ws1 = wb1["备件台账"]
    _no_formula(ck, p1, wb1)
    stock = _rows(ws1, 1, 2, ws1.max_row)
    ck.check(len(stock) == 30, f"备件台账行数应为 30，实际 {len(stock)}")
    below = [r for r in stock if r[4] < r[5]]
    equal = [r for r in stock if r[4] == r[5]]
    above = [r for r in stock if r[4] >= r[5] * 1.2]
    ck.check({r[0] for r in below} == {"BJ-007", "BJ-015", "BJ-021"},
             f"低于安全库存应恰好是 BJ-007/015/021，实际 {[r[0] for r in below]}")
    ck.check({r[0] for r in equal} == {"BJ-011"}, f"等于安全库存应恰好是 BJ-011，实际 {[r[0] for r in equal]}")
    ck.check(len(above) == 26, f"≥1.2×安全库存应恰好 26 种，实际 {len(above)}")

    p2 = d / "出入库记录.xlsx"
    wb2 = load_workbook(p2, data_only=True)
    ws2 = wb2["出入库"]
    _no_formula(ck, p2, wb2)
    txns = _rows(ws2, 1, 2, ws2.max_row)
    ck.check(len(txns) == 80, f"出入库记录行数应为 80，实际 {len(txns)}")
    bj021 = [r for r in txns if r[2] == "BJ-021"]
    ck.check(len(bj021) == 12 and all(r[4] == "出库" for r in bj021),
              f"BJ-021 应有 12 行出库、0 行入库，实际 {len(bj021)}")
    ck.check(sum(r[5] for r in bj021) == 60, f"BJ-021 出库合计应为 60 片，实际 {sum(r[5] for r in bj021)}")
    in_count = sum(1 for r in txns if r[4] == "入库")
    out_count = sum(1 for r in txns if r[4] == "出库")
    ck.check(in_count + out_count == 80, "入库+出库应等于总行数")

    last10_cut = T - timedelta(days=1)  # 具体按 wd(k) 比较，下面用偏移重算
    # 用日期反推 wd 偏移做「最近10个工作日 / 此前每10个工作日」判断
    workdays_desc = []
    dd = T - timedelta(days=1)
    while len(workdays_desc) < 40:
        if dd.weekday() < 5:
            workdays_desc.append(dd)
        dd -= timedelta(days=1)
    last10_set = set(workdays_desc[:10])
    first30_set = set(workdays_desc[10:40])

    from collections import defaultdict
    out_by_code = defaultdict(lambda: [0, 0])  # code -> [last10_qty, first30_qty]
    for r in txns:
        if r[4] != "出库":
            continue
        code = r[2]
        dte = r[0]
        if dte in last10_set:
            out_by_code[code][0] += r[5]
        elif dte in first30_set:
            out_by_code[code][1] += r[5]
    ratio_ok = True
    ratio_detail = []
    for code, (recent, prior) in out_by_code.items():
        if code == "BJ-021":
            continue
        prior_avg10 = prior / 3.0
        if prior_avg10 == 0:
            ok = recent == 0
        else:
            ok = recent <= 1.5 * prior_avg10 + 1e-9
        ratio_detail.append((code, recent, prior_avg10, ok))
        if not ok:
            ratio_ok = False
    ck.check(ratio_ok, f"存在超过 1.5 倍的备件：{[d for d in ratio_detail if not d[3]]}")
    bj021_recent = out_by_code.get("BJ-021")
    ck.check(bj021_recent is not None and bj021_recent[0] == 30 and bj021_recent[1] == 30,
              f"BJ-021 最近/此前应各为 30 片，实际 {bj021_recent}")
    bj021_ratio = bj021_recent[0] / (bj021_recent[1] / 3.0)
    ck.check(abs(bj021_ratio - 3.0) < 1e-9, f"BJ-021 最近10个工作日消耗应为此前的 3 倍，实际 {bj021_ratio:.2f}")

    stat["below"] = sorted(r[0] for r in below)
    stat["equal"] = sorted(r[0] for r in equal)
    stat["in_count"] = in_count
    stat["out_count"] = out_count
    stat["bj021_ratio"] = bj021_ratio
    stat["spot_checks"] = [
        ("备件台账.xlsx", 2, stock[0]),
        ("出入库记录.xlsx", 2, txns[0]),
    ]
    return stat


FAULT_MERGE = {
    "密封失效": "密封", "密封件损坏": "密封", "密封圈老化": "密封",
    "密封漏油": "密封", "O形圈损坏": "密封",
    "定子橡胶脱胶": "定子橡胶", "定子掉胶": "定子橡胶",
    "芯轴拉伤": "芯轴", "心轴拉伤": "芯轴",
    "螺纹粘扣": "螺纹", "丝扣损伤": "螺纹",
    "传动轴轴承磨损": "传动轴轴承", "花键磨损": "花键", "卡瓦磨损": "卡瓦",
    "万向轴断裂": "万向轴", "碟簧断裂": "碟簧",
}


def verify_05(out_root: Path, T: date, ck: Checker) -> dict:
    d = out_root / "05-返厂维修"
    stat: dict[str, Any] = {}

    p_gh = d / "广汉-2026.xlsx"
    wb_gh = load_workbook(p_gh, data_only=True)
    ck.check(wb_gh.active.title == "维修记录", f"广汉-2026.xlsx 活动工作表应为“维修记录”，实际 {wb_gh.active.title}")
    ck.check(set(wb_gh.sheetnames) == {"维修记录", "填写说明"}, "广汉-2026.xlsx 应有且只有两张工作表")
    ws_gh = wb_gh["维修记录"]
    _no_formula(ck, p_gh, wb_gh)
    gh_rows = _rows(ws_gh, 1, 2, 49)
    ck.check(ws_gh.max_row == 49, f"广汉维修记录应到第 49 行，实际到第 {ws_gh.max_row} 行")
    ck.check(len(gh_rows) == 48, f"广汉数据行数应为 48，实际 {len(gh_rows)}")

    p_tj = d / "天津-2026.xlsx"
    wb_tj = load_workbook(p_tj, data_only=True)
    ws_tj = wb_tj["Sheet1"]
    _no_formula(ck, p_tj, wb_tj)
    tj_rows = _rows(ws_tj, 1, 2, 37)
    ck.check(len(tj_rows) == 36, f"天津数据行数应为 36，实际 {len(tj_rows)}")
    sum_row = tuple(ws_tj.cell(row=38, column=c).value for c in range(1, 11))
    ck.check(sum_row[0] == "合计", f"天津第 38 行 A 列应为“合计”，实际 {sum_row[0]}")
    ck.check(sum_row[1] == "36 条", f"天津第 38 行工具号列应为“36 条”，实际 {sum_row[1]}")
    recompute_hours = sum(r[9] for r in tj_rows)
    ck.check(sum_row[9] == recompute_hours, f"天津合计工时应为 {recompute_hours}，实际 {sum_row[9]}")

    p_xj = d / "新疆-2026.xlsx"
    wb_xj = load_workbook(p_xj, data_only=True)
    ws_xj = wb_xj["返厂登记"]
    _no_formula(ck, p_xj, wb_xj)
    ck.check(ws_xj["A1"].value == "新疆服务中心 2026 年返厂维修登记表（示例数据）",
              f"新疆 A1 标题文本不符：{ws_xj['A1'].value}")
    xj_headers = tuple(ws_xj.cell(row=2, column=c).value for c in range(1, 11))
    ck.check(xj_headers[0] == "登记日期", f"新疆表头应在第 2 行，实际第 2 行为 {xj_headers}")
    xj_rows = _rows(ws_xj, 2, 3, 32)
    ck.check(len(xj_rows) == 30, f"新疆数据行数应为 30，实际 {len(xj_rows)}")

    total = len(gh_rows) + len(tj_rows) + len(xj_rows)
    ck.check(total == 114, f"三文件合计应为 114，实际 {total}")

    # 分类计数重新核对
    def count_groups(rows, type_idx, fault_idx, groups):
        from collections import Counter
        actual = Counter((r[type_idx], r[fault_idx]) for r in rows)
        for t, f, n in groups:
            got = actual.get((t, f), 0)
            ck.check(got == n, f"({t},{f}) 应有 {n} 行，实际 {got}")

    count_groups(gh_rows, 3, 7, GH_GROUPS)
    count_groups(tj_rows, 2, 6, TJ_GROUPS)
    count_groups(xj_rows, 2, 6, XJ_GROUPS)

    # 合并口径统计（类型, 归并部位）
    from collections import Counter
    merged = Counter()
    for r in gh_rows:
        merged[(r[3], FAULT_MERGE[r[7]])] += 1
    for r in tj_rows:
        merged[(r[2], FAULT_MERGE[r[6]])] += 1
    for r in xj_rows:
        merged[(r[2], FAULT_MERGE[r[6]])] += 1
    top5 = merged.most_common(5)
    expect_top5 = [
        (("液压震击器", "密封"), 30), (("螺杆钻具", "定子橡胶"), 19),
        (("液压震击器", "芯轴"), 12), (("减震器", "花键"), 11),
        (("螺杆钻具", "传动轴轴承"), 9),
    ]
    ck.check(top5 == expect_top5, f"合并后前 5 名应为 {expect_top5}，实际 {top5}")

    gh_unmerged = Counter((r[3], r[7]) for r in gh_rows)
    gh_un_top1 = gh_unmerged.most_common(1)[0]
    ck.check(gh_un_top1 == (("螺杆钻具", "定子橡胶脱胶"), 8),
              f"广汉不合并第一名应为 螺杆钻具·定子橡胶脱胶 8，实际 {gh_un_top1}")
    gh_merged = Counter((r[3], FAULT_MERGE[r[7]]) for r in gh_rows)
    gh_m_top1 = gh_merged.most_common(1)[0]
    ck.check(gh_m_top1 == (("液压震击器", "密封"), 13),
              f"广汉合并后第一名应为 液压震击器·密封 13，实际 {gh_m_top1}")

    for wb, path in ((wb_gh, p_gh), (wb_tj, p_tj), (wb_xj, p_xj)):
        _no_reserved_ids(ck, path, wb)

    stat["counts"] = {"广汉": len(gh_rows), "天津": len(tj_rows), "新疆": len(xj_rows), "合计": total}
    stat["top5"] = top5
    stat["gh_top1_raw"] = gh_un_top1
    stat["gh_top1_merged"] = gh_m_top1
    stat["tj_sum_row"] = 38
    stat["tj_sum_hours"] = sum_row[9]
    stat["spot_checks"] = [
        ("广汉-2026.xlsx", 2, gh_rows[0]),
        ("天津-2026.xlsx", 2, tj_rows[0]),
        ("新疆-2026.xlsx", 3, xj_rows[0]),
    ]
    return stat


def verify_07(out_root: Path, T: date, ck: Checker) -> dict:
    d = out_root / "07-租赁台账"
    p = d / "租赁台账.xlsx"
    wb = load_workbook(p, data_only=True)
    ws = wb["租赁台账"]
    _no_formula(ck, p, wb)
    _no_reserved_ids(ck, p, wb)
    rows = _rows(ws, 1, 2, ws.max_row)
    ck.check(len(rows) == 36, f"租赁台账行数应为 36，实际 {len(rows)}")
    for rid in RESERVED_TOOL_IDS:
        ck.check(all(r[1] != rid for r in rows), f"不应出现保留编号 {rid}")

    outbound = [r for r in rows if r[9] == "在外"]
    returned = [r for r in rows if r[9] == "已归还"]
    repairing = [r for r in rows if r[9] == "维修中"]
    ck.check(len(outbound) == 18, f"在外行数应为 18，实际 {len(outbound)}")
    ck.check(len(returned) == 14, f"已归还行数应为 14，实际 {len(returned)}")
    ck.check(len(repairing) == 4, f"维修中行数应为 4，实际 {len(repairing)}")
    ck.check(all(r[8] is None for r in outbound), "在外行的归还日期应留空")

    uniq_ids = {r[1] for r in outbound}
    ck.check(len(uniq_ids) == 17, f"在外行不重复的工具编号应为 17 个，实际 {len(uniq_ids)}")

    earlier = [r for r in outbound if isinstance(r[7], date) and r[7] < T]
    equal_t = [r for r in outbound if isinstance(r[7], date) and r[7] == T]
    no_date = [r for r in outbound if not isinstance(r[7], date)]
    later = [r for r in outbound if isinstance(r[7], date) and r[7] > T]
    ck.check(len(earlier) == 6, f"到期日早于 T 的应为 6 行，实际 {len(earlier)}")
    ck.check(len(equal_t) == 1, f"到期日等于 T 的应为 1 行，实际 {len(equal_t)}")
    ck.check(len(no_date) == 2, f"没有可用到期日的应为 2 行，实际 {len(no_date)}")
    ck.check(len(later) == 9, f"到期日晚于 T 的应为 9 行，实际 {len(later)}")

    dup_ids = [tid for tid, c in __import__("collections").Counter(r[1] for r in outbound).items() if c > 1]
    ck.check(dup_ids == ["YZJ-0388"], f"在外行里重复的编号应只有 YZJ-0388，实际 {dup_ids}")
    dup_rows = [i + 2 for i, r in enumerate(rows) if r[1] == "YZJ-0388"]

    rate_map: dict[tuple, set] = {}
    for r in rows:
        key = (r[2], r[3])
        rate_map.setdefault(key, set()).add(r[10])
    bad_rate = {k: v for k, v in rate_map.items() if len(v) > 1}
    ck.check(not bad_rate, f"同类型同规格日租金应一致，实际不一致：{bad_rate}")

    for r in returned + repairing:
        ck.check(isinstance(r[6], date) and r[6] < T, f"{r[1]} 出库日期应早于 T")
        ck.check(isinstance(r[7], date) and r[7] < T, f"{r[1]} 合同到期日应早于 T 且已填写")
        ck.check(isinstance(r[8], date) and r[8] < T, f"{r[1]} 归还日期应早于 T 且已填写")
    for r in repairing:
        ck.check(bool(r[11]), f"{r[1]} 维修中应有备注说明在修什么")

    other_ids = {r[1] for r in (returned + repairing)}
    ck.check(len(other_ids & uniq_ids) == 0, "已归还/维修中的编号不应与在外编号重复")
    ck.check(len(other_ids) == len(returned) + len(repairing), "已归还/维修中的编号彼此不应重复")

    outbound_list = []
    for i, r in enumerate(rows):
        if r[9] != "在外":
            continue
        excel_row = i + 2
        out_days = (T - r[6]).days
        if isinstance(r[7], date):
            overdue = r[7] < T
            due_status = "早于T" if r[7] < T else ("等于T" if r[7] == T else "晚于T")
        else:
            overdue = None
            due_status = "无日期" if r[7] is None else "文本待定"
        outbound_list.append({
            "row": excel_row, "tool": r[1], "cust": r[4], "out": r[6], "due": r[7],
            "out_days": out_days, "overdue": overdue, "due_status": due_status,
        })

    stat = {
        "outbound_list": outbound_list,
        "dup_rows": dup_rows,
        "earlier": len(earlier), "equal_t": len(equal_t), "no_date": len(no_date),
        "later": len(later),
        "spot_checks": [
            ("租赁台账.xlsx", dup_rows[0], rows[dup_rows[0] - 2]),
            ("租赁台账.xlsx", 2, rows[0]),
            ("租赁台账.xlsx", rows.index([r for r in rows if r[9]=="维修中"][0]) + 2, [r for r in rows if r[9]=="维修中"][0]),
        ],
    }
    return stat


# ============================================================================
# 答案文档（样例数据说明.md）
#
# 这一节不用生成时的任何中间结果：它把写好的文件重新读一遍，文档里的数字和行号
# 都从读到的内容里算。
# ============================================================================

WEEKDAY_CN = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]


def weekday_cn(d: date) -> str:
    return WEEKDAY_CN[d.weekday()]


def fmt_d(v: Any) -> str:
    """单元格的值转成文档里的写法：空单元格写「（空）」，日期写 ISO。"""
    if v is None or v == "":
        return "（空）"
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


def fmt_day(d: date) -> str:
    return f"{d.isoformat()}（{weekday_cn(d)}）"


def md_table(headers: list[str], rows: list[list[Any]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for r in rows:
        lines.append("| " + " | ".join("" if v is None else str(v) for v in r) + " |")
    return "\n".join(lines)


def read_sheet(path: Path, sheet: Optional[str] = None, header_row: int = 1) -> list[dict]:
    """把一张工作表读成字典列表，键是表头；每行另带 `_row`（Excel 行号）。整行为空的跳过。"""
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    headers = [c.value for c in ws[header_row]]
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        vals = [_norm(ws.cell(row=r, column=c).value) for c in range(1, len(headers) + 1)]
        if all(v is None for v in vals):
            continue
        row = dict(zip(headers, vals))
        row["_row"] = r
        out.append(row)
    return out


def rows_text(rows: list[dict]) -> str:
    """一组行的行号，写成「第 2、7、17 行」。"""
    return "第 " + "、".join(str(r["_row"]) for r in rows) + " 行"


def days_from_today(d: date, T: date) -> str:
    n = (d - T).days
    if n == 0:
        return "就是今天"
    return f"还有 {n} 天" if n > 0 else f"已过 {-n} 天"


def read_inspection_records(html_path: Path) -> list[dict]:
    """从《检验记录表.html》里取三条记录的内容（照片就是用它渲染的）。"""
    import re

    text = html_path.read_text(encoding="utf-8")
    records = []
    for m in re.finditer(r'"(\d)":\s*\{(.*?)\n\s*\}', text, flags=re.S):
        fields = dict(re.findall(r'(\w+):\s*"([^"]*)"', m.group(2)))
        if "toolId" in fields:
            records.append(fields)
    return records


def doc_header(T: date) -> list[str]:
    return [
        "# 样例数据说明（讲师用）",
        "",
        "> **只给讲师。** 不要发给学员，也不要放进 `D:\\OpenWorker培训\\`——放进去会被它当成资料读到，答案就漏了。",
        ">",
        "> 这份文档由 `样例数据制作/生成样例数据.py` 生成，每个数字和行号都是从生成好的文件里重新读出来的。",
        "> 要改内容请改脚本再重新生成，不要直接改这份文档。",
        "",
        f"- 数据里的「今天」是 **{fmt_day(T)}**。上课不在这一天，先按 [README 的样例数据一节](./README.md#样例数据)把日期对齐。",
        "- 「行号」一律指 Excel 左边的行号：表头在第 1 行时，第一条数据是第 2 行。",
        "- 数据全是编的：客户叫「客户A」，检验员只有工号，井号是虚构的。",
        "",
        "## 几个文件夹讲的是同一个厂",
        "",
        "三台设备的事贯穿了好几个文件夹。学员追问「为什么」时，可以引导他们换个文件夹接着找。",
        "",
        md_table(
            ["设备", "文件夹", "在那里看到什么"],
            [
                ["CK-03 数控车床", "01-质量周报", "螺纹加工五天里四天出「螺纹中径超差」，换过刀片还复发"],
                ["", "02-交期风险", "所有订单的螺纹加工都排在这一台上"],
                ["", "03-设备点检", "液压站渗油没闭环，在等备件"],
                ["", "04-备件盘点", "它用的螺纹车刀片，最近的消耗是平时的 3 倍"],
                ["SK-02 深孔钻床", "02-交期风险", "SO-26-101 被它停机拖过了交期"],
                ["", "03-设备点检", "一个多月里修了三次，都是主轴的毛病"],
                ["RC-02 热处理炉", "01-质量周报", "「硬度不合格」逐日上升，周五整批报废"],
                ["", "03-设备点检", "点检发现温控仪表读数漂移，没闭环"],
            ],
        ),
        "",
    ]


def doc_01(out_root: Path, T: date) -> list[str]:
    d = out_root / "01-质量周报"
    F = friday_on_or_before(T)
    days = [F - timedelta(days=4 - i) for i in range(5)]
    shops = ["机加车间", "热处理车间", "装配车间"]
    data = {s: read_sheet(d / f"检验记录-{s}.xlsx") for s in shops}
    ncr = read_sheet(d / "不良品处理单.xlsx")

    out = ["## 01-质量周报（实务一）", ""]
    out.append(f"本周是第 {iso_week(F)} 周：{fmt_day(days[0])}至 {fmt_day(days[4])}。")
    out.append("")
    out.append(md_table(
        ["文件", "工作表", "数据行数", "内容"],
        [[f"检验记录-{s}.xlsx", "检验记录", len(data[s]), "每行一个送检批次"] for s in shops]
        + [["不良品处理单.xlsx", "不良品处理单", len(ncr), "本周开出的处理单"],
           ["上周质量周报.docx", "—", "—", "上一周的周报，当格式参考"]],
    ))
    out += ["", "### 参考答案：本周合计", ""]
    total_rows = []
    grand = [0, 0, 0, 0]
    for s in shops:
        sent = sum(r["送检数"] for r in data[s])
        bad = sum(r["不良数"] for r in data[s])
        rework = sum(r["不良数"] for r in data[s] if r["处置"] == "返修")
        scrap = sum(r["不良数"] for r in data[s] if r["处置"] == "报废")
        total_rows.append([s, sent, bad, rework, scrap])
        grand = [a + b for a, b in zip(grand, [sent, bad, rework, scrap])]
    total_rows.append(["**合计**", *[f"**{v}**" for v in grand]])
    out.append(md_table(["车间", "送检", "不良", "返修", "报废"], total_rows))
    rate = (grand[0] - grand[1]) / grand[0] * 100
    out += ["", f"合格率 {rate:.1f}%（{grand[0] - grand[1]} / {grand[0]}）。", ""]

    thread = [r for r in data["机加车间"] if r["不良类型"] == "螺纹中径超差"]
    out += ["### 藏的问题一：重复出现（CK-03 的螺纹中径超差）", ""]
    out.append(
        f"`检验记录-机加车间.xlsx` 里，「螺纹中径超差」出现 {len(thread)} 次、共 "
        f"{sum(r['不良数'] for r in thread)} 件，全部出在 "
        f"{'、'.join(sorted({r['设备编号'] for r in thread}))} 的螺纹加工上："
    )
    out.append("")
    ck3 = [r for r in data["机加车间"] if r["设备编号"] == "CK-03" and r["工序"] == "螺纹加工"]
    out.append(md_table(
        ["行号", "日期", "零件", "送检", "不良", "处置"],
        [[r["_row"], fmt_day(r["日期"]), r["零件名称"], r["送检数"], r["不良数"], r["处置"] or "—"] for r in ck3],
    ))
    clean_days = [r for r in ck3 if r["不良数"] == 0]
    if clean_days:
        out += ["", f"只有 {fmt_day(clean_days[0]['日期'])}那一批是干净的，第二天又复发。"]
    out.append("")

    hard = [r for r in data["热处理车间"] if r["不良类型"] == "硬度不合格"]
    by_day = [sum(r["不良数"] for r in hard if r["日期"] == dd) for dd in days]
    out += ["### 藏的问题二：趋势（热处理的硬度不合格）", ""]
    out.append(f"「硬度不合格」周一到周五逐日 {' → '.join(str(n) for n in by_day)}：")
    out.append("")
    out.append(md_table(
        ["行号", "日期", "设备", "零件", "送检", "不良", "处置"],
        [[r["_row"], fmt_day(r["日期"]), r["设备编号"], r["零件名称"], r["送检数"], r["不良数"], r["处置"]]
         for r in hard],
    ))
    out.append("")

    others = []
    for s in shops:
        for r in data[s]:
            if r["不良数"] and r["不良类型"] not in ("螺纹中径超差", "硬度不合格"):
                others.append([f"检验记录-{s}.xlsx", r["_row"], fmt_day(r["日期"]), r["设备编号"],
                               r["不良类型"], r["不良数"], r["处置"]])
    out += ["### 其余零散的不良", ""]
    out.append(md_table(["文件", "行号", "日期", "设备", "不良类型", "不良", "处置"], others))
    out.append("")

    closed = sum(1 for r in ncr if r["状态"] == "已关闭")
    out += ["### 不良品处理单", ""]
    out.append(f"共 {len(ncr)} 张：已关闭 {closed} 张，处理中 {len(ncr) - closed} 张。"
               "CK-03 那几张的原因分析一张比一张深，可以让学员追问「这几张处理单说明了什么」。")
    out.append("")
    out.append(md_table(
        ["行号", "单号", "日期", "设备", "不良描述", "数量", "原因分析", "状态"],
        [[r["_row"], r["单号"], fmt_d(r["日期"]), r["设备编号"], r["不良描述"], r["数量"], r["原因分析"], r["状态"]]
         for r in ncr],
    ))
    out.append("")

    doc = Document(d / "上周质量周报.docx")
    paras = [p.text for p in doc.paragraphs if p.text.strip()]
    quote = next((p for p in paras if "螺纹中径超差" in p), None)
    out += ["### 上周周报", ""]
    out.append(f"`上周质量周报.docx` 的标题是「{paras[0]}」，分四节：本周概况、不良分布、重复出现的问题、下周关注点。")
    if quote:
        out.append(f"里面有一句「{quote}」——和本周的 {sum(r['不良数'] for r in thread)} 件放在一起看，"
                   "就是「换了刀片也没解决」。")
    out.append("")
    return out


def doc_02(out_root: Path, T: date) -> list[str]:
    d = out_root / "02-交期风险"
    orders = read_sheet(d / "订单清单.xlsx")
    plan = read_sheet(d / "排产计划.xlsx")
    prog = {r["订单号"]: r for r in read_sheet(d / "进度跟踪.xlsx")}

    out = ["## 02-交期风险（实务二）", ""]
    out.append(md_table(
        ["文件", "工作表", "数据行数", "内容"],
        [["订单清单.xlsx", "订单", len(orders), "每行一张订单"],
         ["排产计划.xlsx", "排产", len(plan), "每张订单 6 道工序的计划起止"],
         ["进度跟踪.xlsx", "进度", len(prog), "每张订单当前做到哪、有什么异常"]],
    ))
    out += ["", "三张表靠「订单号」对得上。把订单清单和进度跟踪并到一起看：", ""]
    table = []
    for o in orders:
        p = prog[o["订单号"]]
        table.append([
            o["_row"], o["订单号"], o["产品"], fmt_d(o["合同交期"]), days_from_today(o["合同交期"], T),
            o["订单状态"], f"{p['完成率(%)']}%", p["异常类型"], p["异常说明"] or "",
        ])
    out.append(md_table(
        ["行号", "订单号", "产品", "合同交期", "距今天", "订单状态", "完成率", "异常类型", "异常说明"], table))
    out += ["", "（订单清单和进度跟踪的行号相同，都按订单号排。）", ""]

    open_orders = [o for o in orders if o["订单状态"] != "已发货"]
    late = [o["订单号"] for o in open_orders if o["合同交期"] < T]
    soon = [o["订单号"] for o in open_orders if 0 <= (o["合同交期"] - T).days <= 14]
    by_kind: dict[str, list[str]] = {}
    for o in open_orders:
        kind = prog[o["订单号"]]["异常类型"]
        if kind != "无":
            by_kind.setdefault(kind, []).append(o["订单号"])

    out += ["### 参考答案", ""]
    out.append(f"- 已过交期还没发货：{'、'.join(late)}。")
    out.append(f"- 两周内要交货（交期在今天到 14 天后之间，还没发货）：{'、'.join(soon)}，共 {len(soon)} 张。")
    for kind in ("缺料", "工序拖期", "设备故障"):
        out.append(f"- 异常类型是「{kind}」的：{'、'.join(by_kind.get(kind, [])) or '无'}。")
    out += [
        "",
        "风险排在最前面的应该是 SO-26-101、SO-26-103、SO-26-104 这三张（先后顺序可以讨论）。",
        "",
        "**SO-26-102 是故意放的反例**：在没发货的订单里它交期最近，但进度正常、没有异常。"
        "它要是只看交期、把这张排到最前，就是没做「交叉比对」，让学员追问一句「进度正常的别算」。",
        "",
    ]
    return out


def doc_03(out_root: Path, T: date) -> list[str]:
    d = out_root / "03-设备点检"
    equip = read_sheet(d / "设备台账.xlsx")
    checks = read_sheet(d / "点检表.xlsx")
    maint = read_sheet(d / "维修记录.xlsx")

    out = ["## 03-设备点检（实务三）", ""]
    out.append(md_table(
        ["文件", "工作表", "数据行数", "内容"],
        [["设备台账.xlsx", "设备台账", len(equip), "每行一台设备，有点检周期和下次应检日期"],
         ["点检表.xlsx", "点检记录", len(checks), "每日点检的设备近十个工作日的记录，加上其余设备最近一次的记录"],
         ["维修记录.xlsx", "维修记录", len(maint), "近两个月的维修"]],
    ))
    due = [r for r in equip if r["下次应检日期"] == T]
    overdue = [r for r in equip if r["下次应检日期"] < T]
    out += ["", f"### 参考答案：今天应检 {len(due)} 台，逾期 {len(overdue)} 台", ""]
    out.append(md_table(
        ["行号", "设备编号", "设备名称", "点检周期", "上次点检", "下次应检", "状态"],
        [[r["_row"], r["设备编号"], r["设备名称"], r["点检周期"], fmt_d(r["上次点检日期"]),
          fmt_d(r["下次应检日期"]), "今天应检"] for r in due]
        + [[r["_row"], r["设备编号"], r["设备名称"], r["点检周期"], fmt_d(r["上次点检日期"]),
            fmt_d(r["下次应检日期"]), f"**逾期 {(T - r['下次应检日期']).days} 天**"] for r in overdue],
    ))
    abnormal = [r for r in checks if r["结果"] == "异常"]
    open_items = [r for r in abnormal if r["处理状态"] == "未闭环"]
    out += ["", f"### 藏的问题一：{len(open_items)} 条没闭环", ""]
    out.append(f"`点检表.xlsx` 里结果为「异常」的共 {len(abnormal)} 行，其中 {len(open_items)} 行还没闭环：")
    out.append("")
    out.append(md_table(
        ["行号", "日期", "设备", "异常说明", "处理状态"],
        [[r["_row"], fmt_day(r["日期"]), r["设备编号"], r["异常说明"], r["处理状态"]]
         for r in sorted(abnormal, key=lambda r: r["日期"])],
    ))
    counts: dict[str, int] = {}
    for r in maint:
        counts[r["设备编号"]] = counts.get(r["设备编号"], 0) + 1
    repeat = [k for k, v in counts.items() if v > 1]
    out += ["", "### 藏的问题二：反复坏的设备", ""]
    out.append(f"`维修记录.xlsx` 里只有 {'、'.join(repeat)} 出现了不止一次，其余设备各一次：")
    out.append("")
    notable = [r for r in maint if r["设备编号"] in repeat or r["状态"] != "已修复"]
    out.append(md_table(
        ["行号", "日期", "设备", "故障现象", "处理措施", "停机(h)", "状态"],
        [[r["_row"], fmt_day(r["日期"]), r["设备编号"], r["故障现象"], r["处理措施"], r["停机时长(h)"], r["状态"]]
         for r in notable],
    ))
    out.append("")
    return out


def doc_04(out_root: Path, T: date) -> list[str]:
    d = out_root / "04-备件盘点"
    parts = read_sheet(d / "备件台账.xlsx")
    io = read_sheet(d / "出入库记录.xlsx")

    out = ["## 04-备件盘点（实务四）", ""]
    n_in = sum(1 for r in io if r["类型"] == "入库")
    out.append(md_table(
        ["文件", "工作表", "数据行数", "内容"],
        [["备件台账.xlsx", "备件台账", len(parts), "每行一种备件，有当前库存和安全库存"],
         ["出入库记录.xlsx", "出入库", len(io), f"近 40 个工作日的流水：出库 {len(io) - n_in} 行，入库 {n_in} 行"]],
    ))
    below = [p for p in parts if p["当前库存"] < p["安全库存"]]
    equal = [p for p in parts if p["当前库存"] == p["安全库存"]]
    out += ["", f"### 藏的问题一：低于安全库存的 {len(below)} 种", ""]
    out.append(md_table(
        ["行号", "备件编码", "备件名称", "规格型号", "当前库存", "安全库存", "缺口"],
        [[p["_row"], p["备件编码"], p["备件名称"], p["规格型号"], f"{p['当前库存']} {p['单位']}",
          f"{p['安全库存']} {p['单位']}", p["安全库存"] - p["当前库存"]] for p in below],
    ))
    out += ["", "### 边界：正好等于安全库存", ""]
    for p in equal:
        out.append(f"- 第 {p['_row']} 行 {p['备件编码']} {p['备件名称']}（{p['规格型号']}）："
                   f"当前库存 {p['当前库存']}，安全库存 {p['安全库存']}。")
    out.append("")
    out.append("算不算预警没有标准答案，让学员自己定规矩，再让它照办。")

    recent_from = wd(T, 10)
    usage = []
    for p in parts:
        outs = [r for r in io if r["备件编码"] == p["备件编码"] and r["类型"] == "出库"]
        if not outs:
            continue
        recent = sum(r["数量"] for r in outs if r["日期"] >= recent_from)
        before = sum(r["数量"] for r in outs if r["日期"] < recent_from)
        avg = before / 3
        usage.append((p, outs, before, recent, (recent / avg) if avg else None))
    out += ["", "### 藏的问题二：消耗异常", ""]
    out.append(f"出入库记录覆盖 40 个工作日。把「最近 10 个工作日」（{recent_from.isoformat()} 起）的出库量，"
               "和「此前 30 个工作日平均每 10 个工作日」的出库量相比：")
    out.append("")
    out.append(md_table(
        ["备件编码", "备件名称", "规格型号", "此前 30 个工作日出库", "最近 10 个工作日出库", "倍数"],
        [[p["备件编码"], p["备件名称"], p["规格型号"], before, recent,
          ("—" if ratio is None else (f"**{ratio:.1f}**" if ratio > 1.5 else f"{ratio:.1f}"))]
         for p, _outs, before, recent, ratio in usage],
    ))
    for p, outs, _before, _recent, ratio in usage:
        if ratio is not None and ratio > 1.5:
            late_rows = [r for r in outs if r["日期"] >= recent_from]
            purposes = "、".join(sorted({r["用途"] for r in late_rows}))
            out += ["", f"{p['备件编码']} {p['备件名称']}最近 10 个工作日的出库在{rows_text(late_rows)}，"
                        f"用途写的是「{purposes}」。它同时也低于安全库存，而且这段时间没有入库。"]
    out.append("")
    return out


# 05 三个文件的列名各不相同：每一项是（这一列的意思，三个文件里可能用到的列名）。
COLUMN_MEANINGS = [
    ("序号", ("序号",)),
    ("日期", ("返厂日期", "日期", "登记日期")),
    ("工具编号", ("工具编号", "工具号")),
    ("工具类型", ("工具类型", "工具类别", "类型")),
    ("规格", ("规格(OD,in)", "外径(in)", "规格")),
    ("客户", ("客户", "甲方")),
    ("井号", ("井号",)),
    ("故障部位", ("故障部位", "失效部位", "故障位置")),
    ("故障现象", ("故障描述", "失效现象", "情况说明")),
    ("处理措施", ("处理措施", "处理", "维修措施")),
    ("工时", ("维修工时(h)", "工时", "工时(h)")),
]


def doc_05(out_root: Path) -> list[str]:
    from collections import Counter

    d = out_root / "05-返厂维修"
    layout = [
        ("广汉-2026.xlsx", "维修记录", 1, "工具类型", "故障部位"),
        ("天津-2026.xlsx", "Sheet1", 1, "工具类别", "失效部位"),
        ("新疆-2026.xlsx", "返厂登记", 2, "类型", "故障位置"),
    ]
    data = {}
    extra = {}
    headers = {}
    sheets = {}
    for fname, sheet, hdr, type_col, _part_col in layout:
        rows = read_sheet(d / fname, sheet, hdr)
        data[fname] = [r for r in rows if r[type_col] in TYPE_PREFIX]
        extra[fname] = [r for r in rows if r[type_col] not in TYPE_PREFIX]
        headers[fname] = [k for k in rows[0] if k != "_row"]
        sheets[fname] = load_workbook(d / fname, read_only=True).sheetnames

    out = ["## 05-返厂维修（实务五）", "", "### 三个文件长什么样", ""]
    out.append(md_table(
        ["", *[f for f, *_ in layout]],
        [
            ["工作表", *["、".join(sheets[f]) for f, *_ in layout]],
            ["表头在第几行", *[h for _f, _s, h, *_ in layout]],
            ["数据在哪几行", *[f"第 {data[f][0]['_row']}–{data[f][-1]['_row']} 行" for f, *_ in layout]],
            ["**记录数**", *[f"**{len(data[f])}**" for f, *_ in layout]],
            ["工具类型那一列叫", *[t for _f, _s, _h, t, _p in layout]],
            ["故障部位那一列叫", *[p for _f, _s, _h, _t, p in layout]],
        ],
    ))
    total = sum(len(v) for v in data.values())
    out += ["", f"三个文件合计 **{' + '.join(str(len(data[f])) for f, *_ in layout)} = {total}** 条记录。", ""]
    out.append("结构上的三个坑：")
    out.append("")
    out.append(f"- `广汉-2026.xlsx` 有 {len(sheets['广汉-2026.xlsx'])} 张工作表，"
               "「填写说明」里只有几行文字，不是数据。")
    for r in extra["天津-2026.xlsx"]:
        vals = [fmt_d(v) for k, v in r.items() if k != "_row" and v is not None]
        out.append(f"- `天津-2026.xlsx` 第 {r['_row']} 行是合计行（{' / '.join(vals)}），不是一条记录。"
                   f"把它算进去，总数就成了 {total + 1}。")
    out.append("- `新疆-2026.xlsx` 第 1 行是标题，表头在第 2 行。")
    out += ["", "完整的列名对照（「—」表示这个文件没有这一列）：", ""]
    matched: set[tuple[str, str]] = set()
    name_rows = []
    for meaning, candidates in COLUMN_MEANINGS:
        cells = []
        for f, *_ in layout:
            hit = next((c for c in candidates if c in headers[f]), None)
            if hit:
                matched.add((f, hit))
            cells.append(hit or "—")
        name_rows.append([meaning, *cells])
    out.append(md_table(["这一列是什么", *[f for f, *_ in layout]], name_rows))
    unmatched = [(f, h) for f, *_ in layout for h in headers[f] if (f, h) not in matched]
    if unmatched:
        out += ["", "上表没有对上的列：" + "、".join(f"`{f}` 的「{h}」" for f, h in unmatched) + "。"]

    out += ["", "### 同一种故障的几种写法", ""]
    out.append("下面是各文件里「工具类型 × 故障部位」原样的写法和行数，以及参考的合并口径。"
               "合并口径没有唯一答案，学员认可的那一版才算数；下面这版是出题时用的。")
    out.append("")
    merged_names: dict[str, list[str]] = {}
    for raw, canon in FAULT_MERGE.items():
        merged_names.setdefault(canon, []).append(raw)
    out.append(md_table(
        ["合并成", "包含的写法"],
        [[canon, "、".join(raws)] for canon, raws in merged_names.items() if len(raws) > 1]
        + [["（各自只有一种写法）",
            "、".join(raws[0] for raws in merged_names.values() if len(raws) == 1)]],
    ))
    out.append("")
    raw_rows = []
    for fname, _s, _h, type_col, part_col in layout:
        c = Counter((r[type_col], r[part_col]) for r in data[fname])
        for (t, p), n in sorted(c.items(), key=lambda kv: (-kv[1], kv[0])):
            raw_rows.append([fname, t, p, n, FAULT_MERGE[p]])
    out.append(md_table(["文件", "工具类型", "故障部位（原样）", "行数", "合并后归到"], raw_rows))

    def top(counter: Counter, n: int = 5) -> list[list[Any]]:
        ranked = sorted(counter.items(), key=lambda kv: (-kv[1], kv[0]))[:n]
        return [[i + 1, f"{k[0]} · {k[1]}", v] for i, (k, v) in enumerate(ranked)]

    gh = data["广汉-2026.xlsx"]
    gh_raw = Counter((r["工具类型"], r["故障部位"]) for r in gh)
    gh_merged = Counter((r["工具类型"], FAULT_MERGE[r["故障部位"]]) for r in gh)
    out += ["", "### 参考答案：只看广汉（第 2 步）", "", "不合并写法：", ""]
    out.append(md_table(["名次", "工具类型 · 故障部位", "次数"], top(gh_raw)))
    fifth = sorted(gh_raw.values(), reverse=True)[4]
    tied = [k for k, v in gh_raw.items() if v == fifth]
    if len(tied) > 1:
        out += ["", f"（第 5 名有并列：{'、'.join(f'{t} · {p}' for t, p in sorted(tied))} 都是 {fifth} 次。）"]
    out += ["", "合并之后：", ""]
    out.append(md_table(["名次", "工具类型 · 故障部位", "次数"], top(gh_merged)))
    first = top(gh_merged, 1)[0]
    first_type, first_part = first[1].split(" · ")
    out += ["", f"抽查合并后的第一名（{first[1]}，{first[2]} 次）：在 Excel 里筛「工具类型」= {first_type}，"
                "再看「故障部位」——", ""]
    for raw in merged_names[first_part]:
        hit = [r for r in gh if r["工具类型"] == first_type and r["故障部位"] == raw]
        if hit:
            out.append(f"- {raw}：{len(hit)} 行，{rows_text(hit)}")

    all_merged: Counter = Counter()
    for fname, _s, _h, type_col, part_col in layout:
        all_merged.update((r[type_col], FAULT_MERGE[r[part_col]]) for r in data[fname])
    out += ["", "### 参考答案：三个文件合并（第 3、4 步）", ""]
    out.append(md_table(["名次", "工具类型 · 故障部位", "次数"], top(all_merged)))
    out += ["", f"全部 {len(all_merged)} 类加起来是 {sum(all_merged.values())}，等于三个文件的记录数之和。", ""]
    out += ["每个文件各给一条抽查用的记录：", ""]
    for fname, _s, _h, type_col, part_col in layout:
        r = data[fname][len(data[fname]) // 2]
        first_col = headers[fname][0] if headers[fname][0] != "序号" else headers[fname][1]
        id_col = next(h for h in headers[fname] if h in ("工具编号", "工具号"))
        out.append(f"- `{fname}` 第 {r['_row']} 行：{fmt_d(r[first_col])}，{r[id_col]}，"
                   f"{r[type_col]}，{part_col}写的是「{r[part_col]}」。")
    out.append("")
    return out


def doc_06(script_dir: Path) -> list[str]:
    out = ["## 06-返厂检验（实务六）", ""]
    html = script_dir / "检验记录表.html"
    if not html.is_file():
        out += ["（没找到 `样例数据制作/检验记录表.html`，这一节没有生成。）", ""]
        return out
    recs = read_inspection_records(html)
    names = [f"返厂检验记录-{r['toolId']}.jpg" for r in recs]
    out.append("三张照片是同一张表单《井下工具返厂检验记录》的三份填写件。逐栏核对时拿下面这张表当原件。")
    out.append("")
    use = ["第 1 步：手动做一份", "第 2 步：缺项的记录", "第 4 步：新会话回归测试"]
    conclusion = {"rent": "可再租", "repair": "需维修", "scrap": "报废"}

    def cell(r: dict, key: str) -> str:
        return r.get(key) or "**（空着）**"

    fields = [
        ("工具编号", "toolId"), ("工具名称", "toolName"), ("规格（OD，in）", "spec"), ("客户", "customer"),
        ("井号", "well"), ("入井日期", "dateIn"), ("出井日期", "dateOut"), ("累计入井时间（h）", "hours"),
        ("外观检查", "appearance"), ("密封与液压油", "seal"), ("震击测试 · 上击释放力（kN）", "forceUp"),
        ("震击测试 · 下击释放力（kN）", "forceDown"), ("震击测试 · 判定", "judge"),
    ]
    rows = [["课上用在", *use[: len(recs)]]]
    rows += [[label, *[cell(r, key) for r in recs]] for label, key in fields]
    rows.append(["结论（打勾）", *[conclusion.get(r.get("conclusion", ""), "") for r in recs]])
    rows += [[label, *[cell(r, key) for r in recs]]
             for label, key in (("处理意见", "disposition"), ("检验员（工号）", "inspector"), ("检验日期", "inspectDate"))]
    out.append(md_table(["栏目", *[f"`{n}`" for n in names]], rows))
    blank = [n for n, r in zip(names, recs) if not r.get("judge")]
    out += [
        "",
        f"`{'`、`'.join(blank)}` 的震击测试三个填写位是空的。成品里这一栏必须写「待补」；"
        "写成「合格」「未测」或者任何数字，都是它自己编的。",
        "",
        "照片用的是固定日期，和「今天」无关，重新生成数据时不用重做。要改表单内容，"
        "改 `样例数据制作/检验记录表.html` 再运行 `样例数据制作/渲染检验记录.cjs`。",
        "",
    ]
    return out


def doc_07(out_root: Path, T: date) -> list[str]:
    from collections import Counter

    rows = read_sheet(out_root / "07-租赁台账" / "租赁台账.xlsx")
    status = Counter(r["状态"] for r in rows)
    outside = [r for r in rows if r["状态"] == "在外"]

    def verdict(r: dict) -> str:
        due = r["合同到期日"]
        if not isinstance(due, date):
            return "算不出来"
        if due < T:
            return f"超期 {(T - due).days} 天"
        return "今天到期" if due == T else "未到期"

    out = ["## 07-租赁台账（实务七）", ""]
    out.append(f"`租赁台账.xlsx` 一张工作表「租赁台账」，{len(rows)} 行："
               + "、".join(f"{k} {v} 行" for k, v in status.most_common()) + "。")
    out += ["", f"### 参考答案：在外的 {len(outside)} 行（第 1 步）", ""]
    out.append(md_table(
        ["行号", "工具编号", "工具类型", "客户", "井号", "出库日期", "合同到期日", "在外天数", "到期情况"],
        [[r["_row"], r["工具编号"], r["工具类型"], r["客户"], r["井号"], fmt_d(r["出库日期"]),
          fmt_d(r["合同到期日"]), (T - r["出库日期"]).days, verdict(r)] for r in outside],
    ))
    overdue = [r for r in outside if isinstance(r["合同到期日"], date) and r["合同到期日"] < T]
    today = [r for r in outside if r["合同到期日"] == T]
    nodate = [r for r in outside if not isinstance(r["合同到期日"], date)]
    later = [r for r in outside if isinstance(r["合同到期日"], date) and r["合同到期日"] > T]
    dup_ids = [k for k, v in Counter(r["工具编号"] for r in outside).items() if v > 1]
    out += ["", "### 三种边界（第 2 步）", ""]
    for r in today:
        out.append(f"- **到期日正好是今天**：第 {r['_row']} 行 {r['工具编号']}（{r['客户']}，{r['井号']}）。")
    for r in nodate:
        note = f"，备注写着「{r['备注']}」" if r["备注"] else ""
        if r["合同到期日"] is None:
            out.append(f"- **到期日没填**：第 {r['_row']} 行 {r['工具编号']}，单元格空着{note}。")
        else:
            out.append(f"- **到期日填的不是日期**：第 {r['_row']} 行 {r['工具编号']}，"
                       f"填的是「{r['合同到期日']}」{note}。")
    for tid in dup_ids:
        pair = [r for r in outside if r["工具编号"] == tid]
        desc = "；".join(f"第 {r['_row']} 行（{r['客户']}，{fmt_d(r['出库日期'])} 出库）" for r in pair)
        out.append(f"- **同一个编号出现两次**：{tid}，{desc}。两行都写着「在外」、都没有归还日期。"
                   "一支工具不可能同时在两口井上，早的那行是工具归还后台账没更新。")
    dup_old = [min((r for r in outside if r["工具编号"] == tid), key=lambda r: r["出库日期"]) for tid in dup_ids]
    dup_old_overdue = [r for r in dup_old if r in overdue]
    out += ["", "### 规矩不同，答案就不同", ""]
    out.append(f"在外 {len(outside)} 行里：到期日早于今天的 {len(overdue)} 行，正好是今天的 {len(today)} 行，"
               f"晚于今天的 {len(later)} 行，算不出来的 {len(nodate)} 行。"
               f"不重复的工具编号是 {len({r['工具编号'] for r in outside})} 个。")
    out.append("")
    out.append(md_table(
        ["学员定的规矩", "超期行数"],
        [["晚于到期日才算超期（今天到期的不算）", len(overdue)],
         ["到期日当天也算超期", len(overdue) + len(today)],
         ["晚于到期日才算，并且剔除重复编号里早的那一行", len(overdue) - len(dup_old_overdue)]],
    ))
    out += ["", "算不出来的那几行怎么处理（单独列出来、还是按超期提醒）也由学员定。"
                "定下来的规矩要原样写进自动化的指令里。", ""]
    return out


def compile_doc(T: date, out_root: Path, script_dir: Path) -> str:
    parts = doc_header(T)
    parts += doc_01(out_root, T)
    parts += doc_02(out_root, T)
    parts += doc_03(out_root, T)
    parts += doc_04(out_root, T)
    parts += doc_05(out_root)
    parts += doc_06(script_dir)
    parts += doc_07(out_root, T)
    return "\n".join(parts).rstrip("\n") + "\n"


# ============================================================================
# CLI / 主流程
# ============================================================================


def main() -> None:
    ap = argparse.ArgumentParser(description="生成 OpenWorker 培训样例数据")
    ap.add_argument("--today", type=str, default=None, help="YYYY-MM-DD，数据里的「今天」")
    ap.add_argument("--out", type=str, default=None, help="输出根目录")
    ap.add_argument("--doc", type=str, default=None, help="答案文档路径")
    args = ap.parse_args()

    script_dir = Path(__file__).resolve().parent
    T = date.fromisoformat(args.today) if args.today else date.today()
    out_root = Path(args.out).resolve() if args.out else (script_dir.parent / "OpenWorker培训").resolve()
    doc_path = Path(args.doc).resolve() if args.doc else (script_dir.parent / "样例数据说明.md").resolve()

    clean_outputs(out_root)

    summary: list = []
    build_01(out_root, T, summary)
    build_02(out_root, T, summary)
    build_03(out_root, T, summary)
    build_04(out_root, T, summary)
    build_05(out_root, T, summary)
    build_07(out_root, T, summary)

    ck = Checker()
    for verify in (verify_01, verify_02, verify_03, verify_04, verify_05, verify_07):
        verify(out_root, T, ck)

    # 全量兜底扫描：任何生成的 .xlsx 里都不应出现保留给 06 的编号
    for p, _n in summary:
        if p.suffix.lower() == ".xlsx":
            _no_reserved_ids(ck, p, load_workbook(p, data_only=True))

    ck.finish_or_exit()

    doc_text = compile_doc(T, out_root, script_dir)
    doc_path.parent.mkdir(parents=True, exist_ok=True)
    with open(doc_path, "w", encoding="utf-8", newline="\n") as f:
        f.write(doc_text)

    print(f"--today {T.isoformat()}　输出目录 {out_root}")
    print("")
    for p, n in summary:
        rel = p.relative_to(out_root) if p.is_relative_to(out_root) else p
        if n is None:
            print(f"  {rel}")
        else:
            print(f"  {rel}\t{n} 行")
    print("")
    print(f"答案文档：{doc_path}")


if __name__ == "__main__":
    main()
