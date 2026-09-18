"""Real Office files, written in-process — `write_spreadsheet` (.xlsx), and the path /
atomic-write layer that `write_document` (.docx, `tools/document.py`) shares with it.

Why this exists: `write_file` encodes UTF-8 text, so it can only ever produce a text file
wearing a spreadsheet's extension (see `tools/files.py`), and the script route it points at
needs Python AND openpyxl on the USER's machine — the packaged backend is a frozen exe and
`run_shell` spawns the user's own shell, so on an ordinary office PC neither is a given.
This tool writes the ZIP container itself, in the backend process, with openpyxl bundled.

Two invariants the rest of the app depends on:

* **The `path` argument IS the file written.** No extension is appended, no name is
  adjusted. The permission engine scopes the write, tracks provenance and registers the
  artifact from that one argument, so anything else would be a lie on the user's screen.
* **Path semantics are `write_file`'s**, deliberately re-implemented rather than borrowed:
  aisuite's `write_file` takes `content: str`, so a binary payload cannot ride through it.
  `resolve_target` mirrors `FileToolkit._resolve` (relative → primary root; absolute/`~`
  → resolved, then must land inside a declared root; read-only root refused) down to the
  error wording, and `tests/test_office_tools.py` cross-checks the two side by side so they
  cannot drift apart.

The names without a leading underscore — `root_entries`, `resolve_target`, `root_label`,
`prepare_target`, `atomic_write`, `replace_denied`, `receipt_head`, `ILLEGAL_XML_CHARS` —
are the shared surface `tools/document.py` imports. Every binary Office writer has to make
the same five decisions (which root, is it writable, write to a sibling `.part`, swap, say
why the swap was refused), and a second copy of them is a second thing to keep in step
with aisuite. `_replace`, `_save_workbook` and `_receipt` keep their underscored spelling
because they are the SEAMS the tests stand in for by name; `atomic_write` looks `_replace`
up as a module global, so one monkeypatch reaches both tools. (`_replace_denied` is only an
alias for the public `replace_denied`, which is what production code calls.)

openpyxl is imported lazily, inside the call: a build without it must still start, and the
import costs ~40ms that every session which never writes a spreadsheet would pay.
"""

from __future__ import annotations

import os
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable, Optional
from uuid import uuid4

import aisuite as ai

# -- limits ---------------------------------------------------------------------
_MAX_SHEETS = 50
_MAX_CELLS = 500_000  # summed used areas; past this a spreadsheet is the wrong container
_MAX_STYLE_CELLS = 200_000  # one `styles[]` range
_MAX_MERGES = 1_000  # per sheet
_XLSX_MAX_COL = 16_384
_XLSX_MAX_ROW = 1_048_576
_SHEET_NAME_MAX = 31

_AUTO_WIDTH_MIN = 6
_AUTO_WIDTH_MAX = 60
_AUTO_WIDTH_PAD = 2
_FORMULA_WIDTH = 12  # a formula's TEXT is not what the reader sees; don't size for it
_EXCEL_WIDTH_MAX = 255  # Excel's own ceiling; a wider value is rejected on open

_HEADER_FILL = "FFDDEBF7"

_SCHEMA = {
    "type": "function",
    "function": {
        "name": "write_spreadsheet",
        "description": (
            "Create a real Excel .xlsx file (works without Python or Excel installed). "
            "Default look: bold shaded header, thin borders on all cells, auto column "
            "widths, frozen header. Overwrites an existing file. Path rules are the same "
            "as write_file."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "path": {
                    "type": "string",
                    "description": "Output path; must end with .xlsx",
                },
                "sheets": {
                    "type": "array",
                    "description": "One object per worksheet",
                    "items": {
                        "type": "object",
                        "properties": {
                            "name": {
                                "type": "string",
                                "description": "Sheet name, max 31 chars",
                            },
                            "rows": {
                                "type": "array",
                                "description": (
                                    "All rows top to bottom; rows[0] is sheet row 1. Each "
                                    "row is an array of strings. Numeric-looking strings "
                                    "become numbers; prefix ' to keep text (e.g. '00123); "
                                    '"=SUM(C4:C9)" is a formula; "" is an empty cell. For '
                                    "a title, put it in the first cell of row 1 and merge "
                                    "that row across all columns."
                                ),
                                "items": {"type": "array", "items": {"type": "string"}},
                            },
                            "header_rows": {
                                "type": "integer",
                                "description": (
                                    "Number of header rows (after any full-width merged "
                                    "title rows). Default 1"
                                ),
                            },
                            "merges": {
                                "type": "array",
                                "description": (
                                    'Ranges to merge, e.g. "A1:F1", "A2:A3"; the value '
                                    "goes in the top-left cell"
                                ),
                                "items": {"type": "string"},
                            },
                            "col_widths": {
                                "type": "array",
                                "description": (
                                    "Column widths in characters, left to right. Default: "
                                    "auto-fit"
                                ),
                                "items": {"type": "number"},
                            },
                            "borders": {
                                "type": "boolean",
                                "description": (
                                    "Thin borders on every used cell. Default true"
                                ),
                            },
                            "styles": {
                                "type": "array",
                                "description": (
                                    "Optional formatting overrides, applied in order"
                                ),
                                "items": {
                                    "type": "object",
                                    "properties": {
                                        "range": {
                                            "type": "string",
                                            "description": 'e.g. "A1:F1" or "C4:C20"',
                                        },
                                        "bold": {"type": "boolean"},
                                        "size": {"type": "number"},
                                        "color": {
                                            "type": "string",
                                            "description": "Font color, hex like #FF0000",
                                        },
                                        "fill": {
                                            "type": "string",
                                            "description": (
                                                "Background color, hex like #DDEBF7"
                                            ),
                                        },
                                        "align": {
                                            "type": "string",
                                            "enum": ["left", "center", "right"],
                                        },
                                        "wrap": {"type": "boolean"},
                                        "format": {
                                            "type": "string",
                                            "description": (
                                                "Number format, e.g. 0.00, #,##0, 0%, "
                                                "yyyy-mm-dd"
                                            ),
                                        },
                                        "height": {
                                            "type": "number",
                                            "description": (
                                                "Row height in points for the rows in range"
                                            ),
                                        },
                                    },
                                    "required": ["range"],
                                },
                            },
                        },
                        "required": ["rows"],
                    },
                },
            },
            "required": ["path", "sheets"],
        },
    },
}


# -- paths ----------------------------------------------------------------------


def root_entries(roots: Optional[list]) -> list[tuple[Path, bool, str]]:
    """(resolved path, writable, label) per root, recomputed on EVERY call — the list is
    shared and mutated in place when the user grants a folder mid-session (see
    `roots.resolved_paths`); snapshotting it is what made read_file blind to a fresh grant.

    Shapes accepted match `FileToolkit._roots`: RootDir-like, dict, or a bare path. A bare
    path has no `writable` flag, and the toolkit reads `allow_write` there — which the
    catalog leaves at its default when it passes `roots=` — so a bare path is read-only.
    """
    out: list[tuple[Path, bool, str]] = []
    for r in roots or []:
        if isinstance(r, dict):
            raw = r.get("path", "")
            writable = bool(r.get("writable", False))
            label = str(r.get("label") or "")
        elif isinstance(r, (str, Path)):
            raw, writable, label = r, False, ""
        else:  # duck-typed RootDir-like
            raw = getattr(r, "path", "")
            writable = bool(getattr(r, "writable", False))
            label = str(getattr(r, "label", "") or "")
        if raw:
            path = Path(str(raw)).expanduser().resolve()
            out.append((path, writable, label or path.name))
    return out


def resolve_target(
    path: str, primary: Path, entries: list[tuple[Path, bool, str]]
) -> Path:
    """`FileToolkit._resolve(path, for_write=True)`, re-implemented for a binary payload.

    Same rules, same order, same exception types and wording — a divergence here would mean
    `write_spreadsheet` reaching a folder `write_file` refuses (or the reverse), which is
    both a security answer and a "why did that work yesterday" bug report.
    """
    p = Path(str(path)).expanduser()
    candidate = p.resolve() if p.is_absolute() else (primary / p).resolve()
    for root_path, writable, _label in entries:  # FIRST match wins, as upstream
        try:
            candidate.relative_to(root_path)
        except ValueError:
            continue
        if not writable:
            raise PermissionError(f"Path is in a read-only directory: {path}")
        return candidate
    raise PermissionError(f"Path escapes allowed roots: {path}")


def root_label(target: Path, entries: list[tuple[Path, bool, str]]) -> str:
    """Deepest matching root wins — same as `files.write_file_tools`, so a folder nested
    inside another is named with the one the user actually granted."""
    label, depth = "", -1
    for root_path, _writable, entry_label in entries:
        if not target.is_relative_to(root_path):
            continue
        if len(root_path.parts) > depth:
            label, depth = entry_label, len(root_path.parts)
    return label


def prepare_target(
    path: str, roots: Optional[list], primary: Path, tool: str
) -> tuple[Path, list[tuple[Path, bool, str]]]:
    """Everything between the `path` argument and an absolute file this session may write.

    `roots` is read HERE rather than at tool-build time, which is what makes a folder the
    user grants mid-turn reachable on that same turn (see `root_entries`). `tool` only
    names the caller in the refusal — the rules themselves are `write_file`'s.
    """
    entries = root_entries(roots) or [(primary, True, primary.name)]
    if not any(writable for _p, writable, _l in entries):
        raise PermissionError(f"{tool} is disabled: this session has no writable directory.")
    target = resolve_target(path, entries[0][0], entries)
    if target.exists() and target.is_dir():
        raise ValueError(f"Path is a directory: {path}")
    return target, entries


def _check_suffix(path: str) -> None:
    suffix = Path(str(path)).suffix.lower()
    if suffix == ".xlsx":
        return
    if suffix == ".xlsm":
        raise ValueError(
            f"write_spreadsheet cannot create macros, so .xlsm is not supported: {path}. "
            "Pass the same path ending in .xlsx instead."
        )
    if suffix == ".xls":
        raise ValueError(
            f".xls is the old binary Excel format, which this tool cannot write: {path}. "
            "Pass the same path ending in .xlsx instead — Excel and WPS both open it."
        )
    raise ValueError(
        f"path must end in .xlsx (got: {path}). The extension is never added for you, "
        'because the artifact panel and the approval prompt name this exact path: pass e.g. '
        '"报告.xlsx" or an absolute path like "C:/Users/me/Desktop/报告.xlsx".'
    )


# -- A1 ranges ------------------------------------------------------------------

_RANGE_RE = re.compile(
    r"^\$?([A-Za-z]{1,3})\$?([0-9]{1,7})(?::\$?([A-Za-z]{1,3})\$?([0-9]{1,7}))?$"
)


def _col_index(letters: str) -> int:
    idx = 0
    for ch in letters.upper():
        idx = idx * 26 + (ord(ch) - 64)
    return idx


def _col_letter(index: int) -> str:
    out = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        out = chr(65 + rem) + out
    return out


def _parse_range(text: Any, what: str) -> tuple[int, int, int, int]:
    """(min_col, min_row, max_col, max_row), 1-based and normalised."""
    raw = str(text).strip()
    m = _RANGE_RE.match(raw)
    if not m:
        raise ValueError(
            f'bad {what} "{raw}": use A1 notation — a single cell like "B2" or a '
            'rectangle like "A1:F1"'
        )
    c1, r1, c2, r2 = m.group(1), int(m.group(2)), m.group(3), m.group(4)
    col1 = _col_index(c1)
    col2 = _col_index(c2) if c2 else col1
    row2 = int(r2) if r2 else r1
    min_col, max_col = min(col1, col2), max(col1, col2)
    min_row, max_row = min(r1, row2), max(r1, row2)
    if min_row < 1 or max_row > _XLSX_MAX_ROW or max_col > _XLSX_MAX_COL:
        raise ValueError(
            f'{what} "{raw}" is outside the sheet: a worksheet has at most '
            f"{_XLSX_MAX_COL} columns (A..XFD) and {_XLSX_MAX_ROW} rows"
        )
    return min_col, min_row, max_col, max_row


def _range_text(box: tuple[int, int, int, int]) -> str:
    min_col, min_row, max_col, max_row = box
    return (
        f"{_col_letter(min_col)}{min_row}:{_col_letter(max_col)}{max_row}"
        if (min_col, min_row) != (max_col, max_row)
        else f"{_col_letter(min_col)}{min_row}"
    )


def _overlaps(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> bool:
    return (
        a[0] <= b[2] and b[0] <= a[2] and a[1] <= b[3] and b[1] <= a[3]
    )  # col ranges and row ranges both intersect


# -- cell values ----------------------------------------------------------------

# The control characters XML 1.0 has no representation for — everything below \x20 except
# tab, LF and CR. openpyxl refuses them outright (IllegalCharacterError) and lxml refuses
# them for .docx; a model that pastes terminal output hits it. Dropping them beats failing
# a 40-sheet report over one stray \x07.
ILLEGAL_XML_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_ILLEGAL_CHARS = ILLEGAL_XML_CHARS

_INT_RE = re.compile(r"^-?(?:0|[1-9][0-9]*)$")
_DEC_RE = re.compile(r"^-?(?:0|[1-9][0-9]*)\.[0-9]+$")
_THOUSANDS_RE = re.compile(r"^-?[1-9][0-9]{0,2}(?:,[0-9]{3})+(?:\.([0-9]+))?$")
_PERCENT_RE = re.compile(r"^-?(?:0|[1-9][0-9]*)(?:\.([0-9]+))?%$")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_DATETIME_RE = re.compile(r"^(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2}(?::\d{2})?)$")

# 2^53 is where float64 stops counting integers exactly; Excel itself keeps 15 significant
# digits. An 18-digit order number, a bank account, a Chinese ID number: all must stay text
# or the user gets 1.23457E+17 back and the data is gone.
_MAX_SIGNIFICANT_DIGITS = 15
_MAX_FORMULA_CHARS = 8192  # Excel's own formula ceiling
_MAX_CELL_CHARS = 32_767  # Excel's own per-cell ceiling; openpyxl truncates in silence
_MAX_NUMBER_FORMAT_CHARS = 255  # Excel's own format-code ceiling (openpyxl writes any)


def _digit_count(text: str) -> int:
    return sum(1 for ch in text if "0" <= ch <= "9")


def _coerce(raw: Any) -> tuple[Any, Optional[str], bool]:
    """(value, number_format, force_text) for one cell. See the `rows` schema text."""
    if raw is None:
        return None, None, False
    if isinstance(raw, bool):  # before int: bool IS an int in python
        return raw, None, False
    if isinstance(raw, int):
        return (
            (str(raw), None, True)
            if _digit_count(str(abs(raw))) > _MAX_SIGNIFICANT_DIGITS
            else (raw, None, False)
        )
    if isinstance(raw, float):
        # inf/nan have no cell representation; openpyxl writes them and Excel shows #NUM!
        finite = raw == raw and abs(raw) != float("inf")
        return (raw, None, False) if finite else (str(raw), None, True)
    if isinstance(raw, (datetime, date)):
        fmt = "yyyy-mm-dd" if not isinstance(raw, datetime) else "yyyy-mm-dd hh:mm:ss"
        return raw, fmt, False
    if isinstance(raw, (list, tuple, dict, set)):
        raise ValueError(
            f"a cell cannot hold {type(raw).__name__}; each row is a flat array of "
            "strings, e.g. [\"部门一\", \"1200\", \"=B2*2\"]"
        )
    text = _ILLEGAL_CHARS.sub("", str(raw))
    if text == "":
        return None, None, False
    if text[0] == "'":  # the escape hatch: keep it text, whatever it looks like
        # A lone "'" (or "'\x07", once the control character is dropped) leaves nothing
        # behind. Writing that as forced text made openpyxl emit `<c t="inlineStr">` with
        # no `<is>` child — one such cell and Excel declares the WHOLE workbook unreadable
        # and offers to repair it. An empty quote means an empty cell.
        return (text[1:], None, True) if text[1:] else (None, None, False)
    if text[0] == "=" and len(text) > 1:
        return text, None, False  # openpyxl types a leading "=" as a formula
    if _INT_RE.match(text) and _digit_count(text) <= _MAX_SIGNIFICANT_DIGITS:
        return int(text), None, False
    if _DEC_RE.match(text) and _digit_count(text) <= _MAX_SIGNIFICANT_DIGITS:
        return float(text), None, False
    m = _THOUSANDS_RE.match(text)
    if m and _digit_count(text) <= _MAX_SIGNIFICANT_DIGITS:
        decimals = m.group(1) or ""
        plain = text.replace(",", "")
        value = float(plain) if decimals else int(plain)
        fmt = f"#,##0.{'0' * len(decimals)}" if decimals else "#,##0"
        return value, fmt, False
    m = _PERCENT_RE.match(text)
    if m and _digit_count(text) <= _MAX_SIGNIFICANT_DIGITS:
        decimals = m.group(1) or ""
        # Excel stores a percentage as the fraction and formats it; 12% typed as 12 would
        # render as 1200%.
        return (
            float(text[:-1]) / 100.0,
            f"0.{'0' * len(decimals)}%" if decimals else "0%",
            False,
        )
    if _DATE_RE.match(text):
        try:
            return datetime.strptime(text, "%Y-%m-%d").date(), "yyyy-mm-dd", False
        except ValueError:
            return text, None, False  # 2026-02-30: keep the user's own string
    m = _DATETIME_RE.match(text)
    if m:
        clock = m.group(2)
        pattern = "%Y-%m-%d %H:%M:%S" if len(clock) == 8 else "%Y-%m-%d %H:%M"
        try:
            value = datetime.strptime(f"{m.group(1)} {clock}", pattern)
        except ValueError:
            return text, None, False
        return value, "yyyy-mm-dd hh:mm:ss" if len(clock) == 8 else "yyyy-mm-dd hh:mm", False
    return text, None, False


def _balanced(text: str, opener: str, closer: str) -> bool:
    """Brackets balanced, ignoring anything inside a double-quoted literal.

    Known limitation, accepted on purpose: a backslash-escaped quote (`0.00\\"`, legal in a
    number format) reads here as the START of a literal and the input is refused. Teaching
    this the escape rules would also teach it to accept `=SUM(B2:B\\"` — and a wrongly
    refused format costs one turn, while a wrongly accepted one costs the whole workbook.
    """
    depth = 0
    in_string = False
    for ch in text:
        if ch == '"':
            in_string = not in_string
            continue
        if in_string:
            continue
        if ch == opener:
            depth += 1
        elif ch == closer:
            depth -= 1
            if depth < 0:
                return False
    return depth == 0 and not in_string


def _check_formula(text: str, where: str) -> None:
    """Structure only — not a parser, and deliberately not a function-name check.

    The failure worth catching is a TRUNCATED formula: a model that runs out of output
    budget mid-cell emits `=SUM(B2:B`, openpyxl copies it into `<f>` verbatim, and Excel
    then refuses the whole file ("unreadable content"). Refusing the call costs one turn
    and the model resends the formula; writing it costs the user the entire workbook. An
    unknown function name, by contrast, is Excel's business: it shows #NAME? in one cell.
    """
    if len(text) > _MAX_FORMULA_CHARS:
        raise ValueError(
            f"{where}: the formula is {len(text)} characters, over Excel's "
            f"{_MAX_FORMULA_CHARS} limit — compute the value instead, or write it as text "
            "with a leading ' apostrophe"
        )
    if not _balanced(text, "(", ")"):
        raise ValueError(
            f'{where}: the formula {text!r} has unbalanced parentheses or quotes — it '
            "looks truncated. Send the whole formula, or prefix ' to write it as text."
        )


def _check_number_format(text: str, where: str) -> None:
    """Same reasoning one layer down: a half-typed format string ("0.00[red") lands in
    styles.xml and condemns the file, and the number format is not something Excel can
    show an error inside a single cell for."""
    if len(text) > _MAX_NUMBER_FORMAT_CHARS:
        raise ValueError(
            f"{where}: the number format is {len(text)} characters, over Excel's "
            f"{_MAX_NUMBER_FORMAT_CHARS} limit — use a short format code like "
            '"#,##0.00"'
        )
    if _ILLEGAL_CHARS.search(text) or "\n" in text or "\r" in text or "\t" in text:
        raise ValueError(
            f"{where}: a number format cannot contain control characters (got: {text!r})"
        )
    if not _balanced(text, "[", "]"):
        raise ValueError(
            f"{where}: the number format {text!r} has unbalanced [] brackets or quotes; "
            'use e.g. "0.00", "#,##0", "0%", "yyyy-mm-dd" or "#,##0.00;[Red]-#,##0.00"'
        )


def _display_width(text: str) -> int:
    """Column width is measured in characters of the default font, so a CJK glyph counts
    double — without this every Chinese header is clipped."""
    widest = 0
    for line in str(text).split("\n"):
        widest = max(
            widest,
            sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in line),
        )
    return widest


# -- sheet names ----------------------------------------------------------------

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
# Excel keeps "History" for a shared workbook's change log and refuses a sheet by that
# name, case regardless. Pre-seeding `taken` routes it through the same dedupe path as a
# duplicate, so it comes out "History (2)" and lands in the receipt's Adjusted list.
_RESERVED_SHEET_NAMES = {"history"}


def _sheet_names(sheets: list[dict]) -> tuple[list[str], list[str]]:
    """(titles, notes). Excel rejects []:*?/\\ , leading/trailing apostrophes, the reserved
    name History, names over 31 chars and case-insensitive duplicates — and rejects the
    whole FILE, silently, after the user double-clicks it. Fix them here, say so in the
    receipt."""
    titles: list[str] = []
    notes: list[str] = []
    taken: set[str] = set(_RESERVED_SHEET_NAMES)
    for i, sheet in enumerate(sheets, 1):
        given = str(sheet.get("name") or "").strip() if isinstance(sheet, dict) else ""
        name = _BAD_SHEET_CHARS.sub("_", given).strip().strip("'").strip()
        if not name:
            name = f"Sheet{i}"
        name = name[:_SHEET_NAME_MAX]
        base, n = name, 2
        while name.casefold() in taken:
            suffix = f" ({n})"
            name = base[: _SHEET_NAME_MAX - len(suffix)] + suffix
            n += 1
        taken.add(name.casefold())
        titles.append(name)
        if given and given != name:
            notes.append(f'sheet name "{given}" became "{name}"')
    return titles, notes


# -- argument checking ----------------------------------------------------------


def _as_bool(value: Any, what: str, default: bool) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    text = str(value).strip().casefold()
    if text in ("true", "1", "yes"):
        return True
    if text in ("false", "0", "no", ""):
        return False
    raise ValueError(f"{what} must be true or false (got: {value!r})")


def _as_number(value: Any, what: str) -> float:
    try:
        out = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{what} must be a number (got: {value!r})") from None
    if out != out or abs(out) == float("inf") or out <= 0:
        raise ValueError(f"{what} must be a positive number (got: {value!r})")
    return out


_HEX_RE = re.compile(r"^#?([0-9A-Fa-f]{6})$")


def _as_argb(value: Any, what: str) -> str:
    m = _HEX_RE.match(str(value).strip())
    if not m:
        raise ValueError(
            f'{what} must be a hex colour like "#DDEBF7" (six hex digits, # optional); '
            f"got: {value!r}"
        )
    return "FF" + m.group(1).upper()


def _sheet_list(sheets: Any) -> list[dict]:
    if isinstance(sheets, dict):  # one sheet, unwrapped
        sheets = [sheets]
    if not isinstance(sheets, list) or not sheets:
        raise ValueError(
            "sheets must be a non-empty array of sheet objects, e.g. "
            'sheets=[{"name": "汇总", "rows": [["项目", "金额"], ["项目A", "1200"]]}]'
        )
    if len(sheets) > _MAX_SHEETS:
        raise ValueError(f"at most {_MAX_SHEETS} sheets per file (got {len(sheets)})")
    for i, sheet in enumerate(sheets, 1):
        if not isinstance(sheet, dict):
            raise ValueError(
                f"sheets[{i - 1}] must be an object with a `rows` array, not "
                f"{type(sheet).__name__}"
            )
        if "rows" not in sheet:
            raise ValueError(f"sheets[{i - 1}] has no `rows`; every sheet needs one")
    return sheets


def _sheet_rows(sheet: dict, where: str) -> list[list]:
    rows = sheet.get("rows")
    if isinstance(rows, dict) or not isinstance(rows, list):
        raise ValueError(
            f"{where}: rows must be an array of arrays — one array per row, e.g. "
            '[["项目", "金额"], ["项目A", "1200"]]'
        )
    out: list[list] = []
    for i, row in enumerate(rows):
        if isinstance(row, (str, bytes)) or not isinstance(row, (list, tuple)):
            raise ValueError(
                f"{where}: rows[{i}] is {type(row).__name__}, not an array — wrap every "
                'row in its own array, e.g. [["项目", "金额"], ["项目A", "1200"]]'
            )
        out.append(list(row))
    return out


# -- writing --------------------------------------------------------------------


def _save_workbook(workbook: Any, path: Path) -> None:
    """Seam: the one call that can fail deep inside openpyxl (see the atomic-write test)."""
    workbook.save(str(path))


def _replace(src: Path, dst: Path) -> None:
    """Seam: the atomic swap. See `_replace_denied` for what its PermissionError means."""
    os.replace(str(src), str(dst))


def replace_denied(target: Path, app: str = "Excel") -> str:
    """Why the swap was refused, in words the user can act on.

    "Close Excel" is the single most likely cause on Windows and the only one the user can
    fix in five seconds — but it is not the only one: a read-only file, a write-protected
    folder, and Defender's Controlled Folder Access all arrive as the same PermissionError.
    Telling someone to close an Excel they never opened sends them looking for a window
    that does not exist, so the file's own flag is checked first and the fallback names
    both possibilities. `app` is the program to name for the guess — Word for a .docx.
    """
    try:
        read_only = target.exists() and not os.access(target, os.W_OK)
    except OSError:  # pragma: no cover - a stat that fails tells us nothing either way
        read_only = False
    if read_only:
        return (
            f"{target} is read-only / write-protected. Clear the read-only flag, or "
            "choose a different file name."
        )
    return (
        f"{target} is open in another program ({app}?) or is read-only / "
        "write-protected. Close it, clear the read-only flag, or choose a different "
        "file name."
    )


def atomic_write(target: Path, write: Callable[[Path], None], app: str = "Excel") -> None:
    """Run `write` against a sibling `.part`, then swap it into place.

    A half-written .xlsx/.docx is a corrupt ZIP the user double-clicks and Office refuses —
    with the old (good) file already gone. The temp file is a SIBLING on purpose: `os.replace`
    is only atomic within one filesystem, and the system temp dir is routinely another one.

    `_replace` is the SEAM: it is looked up as a module global, so one monkeypatch stands in
    for the swap in both tools (see `test_a_locked_target_says_the_file_is_open_elsewhere`
    and its .docx twin). The wording function it pairs with is the public `replace_denied`.

    The two failure modes are told apart because they need different answers. A
    PermissionError is a file the user can unlock — close Word, clear the read-only flag —
    so it stays a PermissionError carrying `replace_denied`'s advice. Any OTHER OSError is
    the operating system rejecting the NAME, and the only one who can fix that is the
    model: `path="a.txt:b.docx"` (an NTFS alternate data stream) reaches this line and
    `os.replace` raises a bare `OSError: [WinError 87] 参数错误`, which tells the model
    nothing about which argument to change.
    """
    target.parent.mkdir(parents=True, exist_ok=True)
    temp = target.with_name(f".{target.name}.{uuid4().hex[:8]}.part")
    try:
        write(temp)
        try:
            _replace(temp, target)
        except PermissionError as exc:
            raise PermissionError(replace_denied(target, app)) from exc
        except OSError as exc:  # must stay AFTER PermissionError, its subclass
            detail = exc.strerror or str(exc)
            raise ValueError(
                f"the operating system refused this file name: {target} ({detail}). Pass a "
                'plain file name — no : * ? " < > | characters, and not a reserved device '
                "name like CON or NUL — keeping the extension it needs."
            ) from exc
    finally:
        try:
            temp.unlink(missing_ok=True)
        except OSError:  # pragma: no cover - best effort; never masks the real error
            pass


def receipt_head(target: Path, label: str) -> str:
    """`write_file`'s own prefix — the agent is told to quote that string verbatim
    (`roots.render_context`), so every tool that writes a file must look the same."""
    return f"Wrote {target}" + (f" (root: {label})" if label else "")


# Kept for the test that calls the wording function directly with one argument. Production
# code goes through the public `replace_denied`.
_replace_denied = replace_denied


def _receipt(
    target: Path, label: str, summaries: list[str], formulas: bool, notes: list[str]
) -> str:
    parts = [receipt_head(target, label) + ". Sheets: " + "; ".join(summaries) + "."]
    if formulas:
        parts.append("Formulas are calculated when the file is opened in Excel/WPS.")
    if notes:
        parts.append("Adjusted: " + "; ".join(notes) + ".")
    return " ".join(parts)


def office_tools(workspace: str, roots: Optional[list] = None) -> list:
    """`write_spreadsheet` and `write_document`, rooted like the file tools: relative paths
    resolve against `workspace` (or the primary root), absolute paths must land in one of
    `roots`, and a read-only root is refused. `roots` is held BY REFERENCE and re-read per
    call.

    `write_document` is built by `tools/document.py`, imported HERE rather than at module
    scope: that module imports this one for the path and atomic-write layer, and a
    top-level import in both directions is a cycle.
    """
    primary = Path(workspace).resolve()

    def write_spreadsheet(path: str, sheets: list) -> str:
        """Create a real Excel .xlsx file at `path`. Each sheet is an object with `rows`
        (an array of arrays, top to bottom) and optional `name`, `header_rows`, `merges`,
        `col_widths`, `borders` and `styles`."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        except ImportError as exc:  # a build that trimmed the dependency
            raise RuntimeError(
                "this build is missing openpyxl, so it cannot write .xlsx files; deliver "
                "a .csv with write_file instead (start it with the BOM U+FEFF so Excel "
                "reads UTF-8) and tell the user which format you actually delivered."
            ) from exc

        _check_suffix(path)
        target, entries = prepare_target(path, roots, primary, "write_spreadsheet")

        sheet_specs = _sheet_list(sheets)
        titles, notes = _sheet_names(sheet_specs)

        # -- shared style objects: one instance reused across every cell. Building an
        # Alignment per cell is what turns a 100k-cell sheet into a 30-second write.
        thin = Side(style="thin", color="FFB0B0B0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        font_title = Font(bold=True, size=14)
        font_header = Font(bold=True)
        fill_header = PatternFill("solid", start_color=_HEADER_FILL, end_color=_HEADER_FILL)
        align_title = Alignment(horizontal="center", vertical="center")
        align_subtitle = Alignment(horizontal="left", vertical="center")
        align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_data = Alignment(vertical="center")
        align_data_wrap = Alignment(vertical="center", wrap_text=True)

        workbook = Workbook()
        workbook.remove(workbook.active)  # the implicit "Sheet" — ours are named
        # No personal data in the file's properties: this is a document the user forwards.
        workbook.properties.creator = "openworker"
        workbook.properties.lastModifiedBy = "openworker"

        summaries: list[str] = []
        formulas_used = False
        total_cells = 0
        truncated = 0

        for index, (spec, title) in enumerate(zip(sheet_specs, titles), 1):
            where = f'sheets[{index - 1}] ("{title}")'
            rows = _sheet_rows(spec, where)

            merge_boxes: list[tuple[int, int, int, int]] = []
            raw_merges = spec.get("merges") or []
            if isinstance(raw_merges, str):
                raw_merges = [raw_merges]
            if not isinstance(raw_merges, list):
                raise ValueError(f"{where}: merges must be an array of A1 ranges")
            if len(raw_merges) > _MAX_MERGES:
                raise ValueError(
                    f"{where}: at most {_MAX_MERGES} merges per sheet "
                    f"(got {len(raw_merges)})"
                )
            for entry in raw_merges:
                box = _parse_range(entry, f"{where}: merge range")
                if box[0] == box[2] and box[1] == box[3]:
                    continue  # a one-cell "merge" is a no-op, not an error
                for other in merge_boxes:
                    if _overlaps(box, other):
                        raise ValueError(
                            f'{where}: merge ranges "{_range_text(other)}" and '
                            f'"{_range_text(box)}" overlap; each cell may belong to only '
                            "one merge"
                        )
                merge_boxes.append(box)

            # The used area covers the merges too: a title merged across six columns when
            # only three carry data still makes the sheet six columns wide, and that is
            # what "spans every used column" below has to measure against.
            max_row = max([len(rows), *(b[3] for b in merge_boxes)])
            max_col = max(
                [max((len(r) for r in rows), default=0), *(b[2] for b in merge_boxes)]
            )
            total_cells += max_row * max_col
            if total_cells > _MAX_CELLS:
                raise ValueError(
                    f"too much data: {total_cells} cells (rows x columns, all sheets) "
                    f"exceeds the {_MAX_CELLS} limit. Deliver a .csv with write_file "
                    "instead, or split the data across files."
                )

            worksheet = workbook.create_sheet(title=title)
            # Per-cell display width, kept row-aligned with `rows`: the auto-fit below has
            # to leave the title rows out of the measurement (a title spans every column,
            # so measuring it would stretch column A to the width of the whole title), and
            # which rows those are is only known after the merges are parsed.
            cell_widths: list[list[int]] = []

            # 1. values. Before the merges: openpyxl replaces every non-anchor cell of a
            # merge with a read-only MergedCell, so a value written after merging is lost.
            for r, row in enumerate(rows, 1):
                widths_row: list[int] = []
                for c, raw in enumerate(row, 1):
                    value, number_format, force_text = _coerce(raw)
                    if value is None or (isinstance(value, str) and not value):
                        widths_row.append(0)
                        continue
                    if isinstance(value, str):
                        if not force_text and value.startswith("="):
                            _check_formula(
                                value, f"{where} cell {_col_letter(c)}{r}"
                            )
                        elif len(value) > _MAX_CELL_CHARS:
                            # Excel's ceiling. openpyxl cuts it without a word, so cut it
                            # here and say so in the receipt — a silently shortened cell
                            # is a claim the agent makes on our behalf.
                            value = value[:_MAX_CELL_CHARS]
                            truncated += 1
                    cell = worksheet.cell(row=r, column=c, value=value)
                    if force_text:
                        # A quoted "'=SUM(A1:A2)" or "'00123" must arrive as TEXT; setting
                        # .value alone re-types a leading "=" as a formula.
                        cell.data_type = "s"
                    elif cell.data_type == "f":
                        formulas_used = True
                    if number_format:
                        cell.number_format = number_format
                    widths_row.append(
                        _FORMULA_WIDTH
                        if cell.data_type == "f"
                        else _display_width(value if isinstance(value, str) else str(raw))
                    )
                cell_widths.append(widths_row)

            # 2. merges, and the anchor map every style lookup goes through.
            anchor_of: dict[tuple[int, int], tuple[int, int]] = {}
            for min_col, min_row, max_c, max_r in merge_boxes:
                worksheet.merge_cells(
                    start_row=min_row,
                    start_column=min_col,
                    end_row=max_r,
                    end_column=max_c,
                )
                for r in range(min_row, max_r + 1):
                    for c in range(min_col, max_c + 1):
                        anchor_of[(r, c)] = (min_row, min_col)

            def anchor(r: int, c: int, _map=anchor_of) -> Optional[tuple[int, int]]:
                """The cell to style, or None for a non-anchor member of a merge — openpyxl
                silently drops a font/fill/alignment set on one of those."""
                home = _map.get((r, c))
                return None if home is not None and home != (r, c) else (r, c)

            # 3. which rows are the title, and which are the header.
            spanned = {
                r
                for (min_col, min_row, max_c, max_r) in merge_boxes
                if min_col == 1 and max_c == max_col and max_col > 0
                for r in range(min_row, max_r + 1)
            }
            title_rows = 0
            while (title_rows + 1) in spanned:
                title_rows += 1
            header_rows = spec.get("header_rows")
            header_rows = 1 if header_rows is None else header_rows
            try:
                header_rows = int(header_rows)
            except (TypeError, ValueError):
                raise ValueError(
                    f"{where}: header_rows must be a whole number (got: "
                    f"{spec.get('header_rows')!r})"
                ) from None
            header_rows = max(0, min(header_rows, max(0, max_row - title_rows)))
            header_end = title_rows + header_rows

            # 4. the default look.
            draw_borders = _as_bool(spec.get("borders"), f"{where}: borders", True)
            for r in range(1, max_row + 1):
                for c in range(1, max_col + 1):
                    home = anchor(r, c)
                    if home is not None:
                        cell = worksheet.cell(row=home[0], column=home[1])
                        if r <= title_rows:
                            if r == 1:
                                cell.font = font_title
                            cell.alignment = align_title if r == 1 else align_subtitle
                        elif r <= header_end:
                            cell.font = font_header
                            cell.fill = fill_header
                            cell.alignment = align_header
                        else:
                            cell.alignment = (
                                align_data_wrap
                                if isinstance(cell.value, str) and "\n" in cell.value
                                else align_data
                            )
                    if draw_borders:
                        # Every cell of a merged box, all four sides: Excel keeps only the
                        # outward-facing ones on save, which is how a merged title ends up
                        # framed instead of missing three edges.
                        worksheet.cell(row=r, column=c).border = border

            if 1 <= header_end < max_row:
                worksheet.freeze_panes = f"A{header_end + 1}"

            # 5. column widths.
            given_widths = spec.get("col_widths") or []
            if not isinstance(given_widths, list):
                raise ValueError(f"{where}: col_widths must be an array of numbers")
            for c in range(1, max_col + 1):
                if c <= len(given_widths):
                    width = min(
                        _as_number(given_widths[c - 1], f"{where}: col_widths[{c - 1}]"),
                        _EXCEL_WIDTH_MAX,
                    )
                else:
                    measured = max(
                        (
                            widths_row[c - 1]
                            for widths_row in cell_widths[title_rows:]
                            if c - 1 < len(widths_row)
                        ),
                        default=0,
                    )
                    width = min(
                        max(measured + _AUTO_WIDTH_PAD, _AUTO_WIDTH_MIN), _AUTO_WIDTH_MAX
                    )
                worksheet.column_dimensions[_col_letter(c)].width = width

            # 6. explicit overrides, in the order given.
            raw_styles = spec.get("styles") or []
            if isinstance(raw_styles, dict):
                raw_styles = [raw_styles]
            if not isinstance(raw_styles, list):
                raise ValueError(f"{where}: styles must be an array of style objects")
            for s_index, style in enumerate(raw_styles):
                if not isinstance(style, dict) or "range" not in style:
                    raise ValueError(
                        f"{where}: styles[{s_index}] must be an object with a `range`, "
                        'e.g. {"range": "A1:F1", "bold": true}'
                    )
                box = _parse_range(style["range"], f"{where}: styles[{s_index}] range")
                span = (box[2] - box[0] + 1) * (box[3] - box[1] + 1)
                if span > _MAX_STYLE_CELLS:
                    raise ValueError(
                        f'{where}: styles[{s_index}] range "{_range_text(box)}" covers '
                        f"{span} cells, over the {_MAX_STYLE_CELLS} limit"
                    )
                # Clipped to the used area: styling past it would grow the sheet with
                # empty-but-formatted cells the user then has to delete.
                lo_col, lo_row = max(1, box[0]), max(1, box[1])
                hi_col, hi_row = min(box[2], max_col), min(box[3], max_row)
                colour = (
                    _as_argb(style["color"], f"{where}: styles[{s_index}] color")
                    if style.get("color") is not None
                    else None
                )
                fill = (
                    _as_argb(style["fill"], f"{where}: styles[{s_index}] fill")
                    if style.get("fill") is not None
                    else None
                )
                align = style.get("align")
                if align is not None and str(align) not in ("left", "center", "right"):
                    raise ValueError(
                        f"{where}: styles[{s_index}] align must be left, center or right "
                        f"(got: {align!r})"
                    )
                size = (
                    _as_number(style["size"], f"{where}: styles[{s_index}] size")
                    if style.get("size") is not None
                    else None
                )
                bold = (
                    _as_bool(style["bold"], f"{where}: styles[{s_index}] bold", False)
                    if style.get("bold") is not None
                    else None
                )
                wrap = (
                    _as_bool(style["wrap"], f"{where}: styles[{s_index}] wrap", False)
                    if style.get("wrap") is not None
                    else None
                )
                height = (
                    _as_number(style["height"], f"{where}: styles[{s_index}] height")
                    if style.get("height") is not None
                    else None
                )
                number_format = style.get("format")
                if number_format:
                    # No str() rescue here: `["0.00"]` would become the literal format
                    # code `['0.00']`, which passes every structural check and shows the
                    # user a column of `['0.00']`. A wrong type is a model mistake worth
                    # one turn.
                    if not isinstance(number_format, str):
                        raise ValueError(
                            f"{where}: styles[{s_index}] format must be a string, not "
                            f"{type(number_format).__name__} (got: {number_format!r}); "
                            'e.g. "#,##0.00"'
                        )
                    _check_number_format(
                        number_format, f"{where}: styles[{s_index}] format"
                    )
                if height is not None:
                    for r in range(lo_row, hi_row + 1):
                        worksheet.row_dimensions[r].height = height
                for r in range(lo_row, hi_row + 1):
                    for c in range(lo_col, hi_col + 1):
                        home = anchor(r, c)
                        if home is None:
                            continue
                        cell = worksheet.cell(row=home[0], column=home[1])
                        if bold is not None or size is not None or colour is not None:
                            font = Font(
                                name=cell.font.name,
                                size=size if size is not None else cell.font.size,
                                bold=bold if bold is not None else cell.font.bold,
                                italic=cell.font.italic,
                                color=colour if colour is not None else cell.font.color,
                            )
                            cell.font = font
                        if fill is not None:
                            cell.fill = PatternFill(
                                "solid", start_color=fill, end_color=fill
                            )
                        if align is not None or wrap is not None:
                            now = cell.alignment
                            cell.alignment = Alignment(
                                horizontal=(
                                    str(align) if align is not None else now.horizontal
                                ),
                                vertical=now.vertical or "center",
                                wrap_text=wrap if wrap is not None else now.wrap_text,
                            )
                        if number_format:
                            cell.number_format = number_format  # checked above

            merged = len(merge_boxes)
            summaries.append(
                f"{title} {max_row} rows x {max_col} cols"
                + (f", {merged} merge{'s' if merged > 1 else ''}" if merged else "")
            )

        if truncated:
            notes.append(f"{truncated} cell(s) truncated to 32,767 characters")

        atomic_write(target, lambda p: _save_workbook(workbook, p))

        try:
            return _receipt(
                target, root_label(target, entries), summaries, formulas_used, notes
            )
        except Exception:
            # The file is on disk. Everything past the write is cosmetics, and the engine
            # turns a raised tool into `{"error": …}` — which would send the model back to
            # rewrite the file, or make it tell the user the spreadsheet does not exist.
            return f"Wrote {target}"

    write_spreadsheet.__name__ = "write_spreadsheet"
    write_spreadsheet.__aisuite_tool_metadata__ = ai.ToolMetadata(
        name="write_spreadsheet",
        category="filesystem",
        risk_level="medium",
        capabilities=["write_file"],
        requires_approval=True,
    )
    write_spreadsheet.__coworker_schema__ = _SCHEMA

    from coworker.tools.document import document_tools

    return [write_spreadsheet, *document_tools(workspace, roots)]
